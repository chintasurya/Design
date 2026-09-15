import json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

D = json.load(open('talktrack.json'))
BODY, DISPLAY = 'Trenda IG Text', 'Trenda IG Display'
CYAN, INK, MUTE = RGBColor(0x0D,0x6E,0x8C), RGBColor(0x1A,0x1F,0x2B), RGBColor(0x66,0x70,0x80)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.7)
    s.left_margin = s.right_margin = Inches(0.8)

st = doc.styles['Normal']
st.font.name = BODY; st.font.size = Pt(10.5); st.font.color.rgb = INK
st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.18

def para(text='', *, size=10.5, bold=False, color=INK, font=BODY,
         before=0, after=0, caps=False, spacing=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.name = font; r.font.size = Pt(size); r.bold = bold; r.font.color.rgb = color
        r.font.all_caps = caps
        r.element.rPr.rFonts.set(qn('w:eastAsia'), font)
        if spacing is not None:
            sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(int(spacing*20)))
            r.element.rPr.append(sp)
    return p

def rule(space_before=6, space_after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'6')
    bot.set(qn('w:space'),'1'); bot.set(qn('w:color'),'D5DAE2')
    b.append(bot); pPr.append(b)

# ---- cover ----
para('INSIGHT GLOBAL', size=9, bold=True, color=CYAN, font=DISPLAY, caps=True, spacing=1.6, after=4)
para('Presenter talk track', size=26, bold=True, font=DISPLAY, after=6)
para('Intelligent Case Management for McKesson Extended Care', size=13, color=MUTE, after=14)
para(f'{len(D["slides"])} slides  ·  about {D["totalMinutes"]} minutes of speaking, before questions',
     size=10, color=MUTE, after=10)
para('The same text sits in the speaker notes of each slide, so you can present from PowerPoint '
     'Presenter View instead of this document if you prefer. Do not read it out. It is the argument '
     'in the right order, phrased the way it is meant to land. Say it in your own words.',
     size=10, color=MUTE, after=4)
rule(space_before=12, space_after=14)

# ---- running order ----
para('Running order', size=13, bold=True, font=DISPLAY, after=8)
COLS = [Inches(0.4), Inches(5.3), Inches(0.6)]
t = doc.add_table(rows=0, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.LEFT
t.autofit = False
# python-docx only honours widths when the grid and every cell agree.
tblPr = t._tbl.tblPr
layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
for gridCol, w in zip(t._tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol')), COLS):
    gridCol.set(qn('w:w'), str(int(w.inches * 1440)))
for s in D['slides']:
    cells = t.add_row().cells
    for c, w, (txt, bold, col) in zip(cells, COLS, [
        (str(s['n']), True, CYAN),
        (s['title'], False, INK),
        (f"{s['secs']}s", False, MUTE),
    ]):
        c.width = w
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(txt); r.font.name = BODY; r.font.size = Pt(9.5)
        r.bold = bold; r.font.color.rgb = col
        r.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

doc.add_page_break()

# ---- one block per slide ----
for i, s in enumerate(D['slides']):
    para(f'SLIDE {s["n"]}', size=8.5, bold=True, color=CYAN, font=DISPLAY,
         caps=True, spacing=1.6, before=0 if i == 0 else 16, after=2)
    para(s['title'], size=15, bold=True, font=DISPLAY, after=2)
    para(f'{s["purpose"]}   ·   about {s["secs"]} seconds', size=9.5, color=MUTE, after=8)
    para('SAY', size=8.5, bold=True, color=CYAN, font=DISPLAY, caps=True, spacing=1.4, after=3)
    para(s['say'], size=10.5, after=0)
    if s.get('ask'):
        para('WATCH FOR', size=8.5, bold=True, color=CYAN, font=DISPLAY,
             caps=True, spacing=1.4, before=9, after=3)
        para(s['ask'], size=10, color=MUTE, after=0)
    if i < len(D['slides']) - 1:
        rule(space_before=12, space_after=0)

doc.save('McKesson_ICM_Talk_Track.docx')
print('McKesson_ICM_Talk_Track.docx written')
