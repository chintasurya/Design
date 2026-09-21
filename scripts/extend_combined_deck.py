# -*- coding: utf-8 -*-
"""Fill the transition slide and append the platform slide in the combined deck.

Edits the user's merged deck in place so the two architecture posters are
untouched. Same Trenda type system and palette as the rest of the deck.
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = sys.argv[1]
OUT = sys.argv[2]
WIDTHS = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "trenda-widths.json")))

C = dict(
    ink="0E1B2E", ink2="16202C", mint="5FD0C0", white="FFFFFF",
    text="1A1A1A", grey="5A6573", rule="D5D9DE", panel="F4F6F9",
    blue="1B6BC0", teal="00757F", pack="C77400", purple="6A1B9A",
    core="0E8A7D", red="C62828", dim="A9B4C2", dim2="D8DEE6", soft="C9CFCC",
)
F = dict(
    dspBold="Trenda IG Display Bold", dspSemi="Trenda IG Display Semibold",
    txt="Trenda IG Text", txtSemi="Trenda IG Text Semibold",
    txtLight="Trenda IG Text Light",
)

def rgb(h):
    return RGBColor.from_string(h)

def text_w(s, face, pt):
    m = WIDTHS.get(face, WIDTHS["Trenda IG Text"])
    em = sum(m["widths"].get(ch, m["default"]) for ch in s)
    return em * pt / 72.0

def line_h(face, pt, mult=1.0):
    return WIDTHS.get(face, WIDTHS["Trenda IG Text"])["lineHeight"] * pt / 72.0 * mult

def n_lines(s, face, pt, w_in):
    sp = text_w(" ", face, pt)
    n, cur = 1, 0.0
    for word in s.split(" "):
        ww = text_w(word, face, pt)
        if ww > w_in:
            if cur > 0:
                n += 1
            n += int(ww // w_in)
            cur = ww % w_in
        elif cur == 0:
            cur = ww
        elif cur + sp + ww <= w_in:
            cur += sp + ww
        else:
            n += 1
            cur = ww
    return n

def box(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh

def dot(slide, x, y, color, size=0.09):
    return box(slide, x, y, size, size, color)

def txt(slide, s, x, y, w, h, size=10, face=None, color="1A1A1A",
        align=None, anchor=MSO_ANCHOR.TOP, spacing=None, line_mult=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = s.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        if line_mult:
            p.line_spacing = line_mult
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.name = face or F["txt"]
        r.font.color.rgb = rgb(color)
        if spacing:
            r.font._rPr.set("spc", str(int(spacing * 100)))
    return tb

def bullets(slide, items, x, y, w, size=9, face=None, color="1A1A1A",
            gap=4, indent=0.11, line_mult=1.0):
    """Manual hanging bullets: full control of indent at small sizes."""
    face = face or F["txt"]
    cy = y
    for it in items:
        txt(slide, "•", x, cy, 0.12, line_h(face, size) + 0.04,
            size=size, face=face, color=color)
        n = n_lines(it, face, size, w - indent)
        txt(slide, it, x + indent, cy, w - indent, n * line_h(face, size, line_mult) + 0.04,
            size=size, face=face, color=color, line_mult=line_mult)
        cy += n * line_h(face, size, line_mult) + gap / 72.0
    return cy

prs = Presentation(SRC)
W = prs.slide_width / 914400.0

# ---------------------------------------------------------------- slide 4
s4 = prs.slides[3]
# clear the placeholder text box
for sh in list(s4.shapes):
    sh._element.getparent().remove(sh._element)

box(s4, 0, 0, W, 7.5, C["ink"])
# decorative corner blocks, quiet
box(s4, W - 2.0, 0, 2.0, 0.22, C["mint"])
box(s4, 0, 7.28, 2.0, 0.22, C["mint"])

txt(s4, "FROM FRAMEWORK TO PRACTICE", 1.0, 1.32, 8.0, 0.26,
    size=11, face=F["dspSemi"], color=C["mint"], spacing=1.4)
txt(s4, "One framework. Two solutions, for reference.", 1.0, 1.7, 11.0, 0.62,
    size=33, face=F["dspBold"], color=C["white"])
txt(s4, "What follows are two working examples of the same nine layers, built for two unrelated domains. "
        "Same engine, different declarations.", 1.0, 2.62, 8.6, 0.6,
    size=13, face=F["txt"], color=C["soft"], line_mult=1.15)

cards = [
    ("01", "SAMPLE ARCHITECTURE", "Engineering delivery",
     "Requirements, code, tests and releases across six systems of record. "
     "Demand is a requirement. Supply is a component that implements it.", C["purple"]),
    ("02", "SAMPLE ARCHITECTURE", "Clinical network",
     "Providers, facilities and patient demand across seven sources, behind a privacy boundary. "
     "Demand is a patient in a geography. Supply is a credentialed provider.", C["teal"]),
]
for i, (num, label, title, body, col) in enumerate(cards):
    x = 1.0 + i * 5.75
    box(s4, x, 3.62, 5.4, 1.74, C["ink2"])
    box(s4, x, 3.62, 5.4, 0.06, col)
    txt(s4, num, x + 0.3, 3.86, 0.8, 0.4, size=22, face=F["dspBold"], color=col)
    txt(s4, label, x + 1.05, 3.94, 3.0, 0.2, size=8, face=F["txtSemi"],
        color=C["dim"], spacing=1.2)
    txt(s4, title, x + 1.05, 4.14, 4.0, 0.3, size=17, face=F["dspBold"], color=C["white"])
    txt(s4, body, x + 0.3, 4.62, 4.8, 0.84, size=10, face=F["txt"],
        color=C["dim2"], line_mult=1.1)

dot(s4, 1.0, 6.28, C["mint"], 0.1)
txt(s4, "Read them as evidence that the abstraction holds, not as two separate products. "
        "Nine layers, nine layers. Only the five declarations differ.",
    1.22, 6.2, 11.0, 0.3, size=11.5, face=F["txt"], color=C["soft"])

# ---------------------------------------------------------------- slide 7
s7 = prs.slides.add_slide(prs.slide_layouts[0])
box(s7, 0, 0, W, 7.5, C["white"])
box(s7, 0, 0, W, 0.95, C["ink"])
txt(s7, "THE QUESTIONS THAT COME NEXT", 0.5, 0.15, 9.0, 0.24,
    size=10.5, face=F["dspSemi"], color=C["mint"], spacing=1.2)
txt(s7, "Where it runs, what it is written in, and which model", 0.5, 0.4, 10.0, 0.42,
    size=22, face=F["dspBold"], color=C["white"])
txt(s7, "Every one of these\nis a swap, not a rebuild.", 10.85, 0.24, 2.3, 0.5,
    size=9, face=F["txtLight"], color=C["dim"], line_mult=1.08)

cols = [
    (C["blue"], "Where does it run?", [
        "Any cloud, or none: AWS, Azure, GCP, private cloud or on premise",
        "Packaged as containers, so it lands in whatever you already operate",
        "Batch on a schedule to rebuild the graph, an API service to query it",
        "No managed service is assumed anywhere in the design",
    ], "If it runs Python and containers, it runs this."),
    (C["teal"], "Where does the data live?", [
        "Object storage: S3, Azure Blob, Google Cloud Storage",
        "Graph store: Neo4j, TigerGraph, Neptune, Memgraph",
        "Warehouse: Snowflake, BigQuery, Databricks, Synapse",
        "Relational: Postgres, where the graph is small enough to sit there",
    ], "The snapshot is vendor neutral JSON, so the store is a swap, not a rewrite."),
    (C["pack"], "What is it written in?", [
        "Python is the core: ingestion, identity, derived edges and the gap engines",
        "Standard libraries and open packages. No proprietary runtime",
        "Orchestrated by whatever you already run: Airflow, Dagster, Step Functions, cron",
        "Configuration is data: ontology shapes, specifications and thresholds are files, not code",
    ], "A new industry ships as configuration and connectors, not a new codebase."),
    (C["purple"], "Which AI model?", [
        "Claude, Gemini, GPT and Codex, or an open weights model, behind one adapter",
        "The model turns a question into a traversal, and evidence into a sentence",
        "It does not decide the answer. The traversal computes it and the citation proves it",
        "Model choice can differ per environment and per sensitivity rule",
    ], "Swap the model and the answers do not change. That is the test."),
]
CW = 3.1675
for i, (col, title, items, rule) in enumerate(cols):
    x = 0.12 + i * 3.3075
    box(s7, x, 1.08, CW, 3.82, C["panel"], C["rule"])
    box(s7, x, 1.08, CW, 0.42, col)
    txt(s7, title, x + 0.16, 1.08, CW - 0.32, 0.42, size=12.5, face=F["dspBold"],
        color=C["white"], anchor=MSO_ANCHOR.MIDDLE)
    bullets(s7, items, x + 0.16, 1.64, CW - 0.32, size=9.5, color=C["text"], gap=8)
    box(s7, x + 0.16, 3.94, CW - 0.32, 0.8, C["ink"])
    dot(s7, x + 0.3, 4.08, col, 0.08)
    txt(s7, rule, x + 0.3, 4.24, CW - 0.6, 0.46, size=9, face=F["txtSemi"],
        color=C["white"], line_mult=1.02)

box(s7, 0.12, 5.06, W - 0.24, 1.08, C["panel"], C["rule"])
dot(s7, 0.32, 5.28, C["red"], 0.1)
txt(s7, "Three things we do require, whatever you choose", 0.54, 5.19, 6.0, 0.24,
    size=12, face=F["dspBold"], color=C["text"])
reqs = [
    ("A store that can traverse", "or every hop is paid for in hand written SQL"),
    ("An endpoint under the right agreement", "where regulated data is in play, or the model sees aggregates only"),
    ("Reproducibility from a snapshot id", "any answer replayable at a date, whatever the stack underneath"),
]
for i, (a, b) in enumerate(reqs):
    x = 0.54 + i * 4.22
    txt(s7, str(i + 1) + "  " + a, x, 5.54, 4.05, 0.22, size=10, face=F["txtSemi"], color=C["red"])
    txt(s7, b, x, 5.77, 4.05, 0.34, size=9.2, face=F["txt"], color=C["grey"], line_mult=1.0)

box(s7, 0.12, 6.32, W - 0.24, 0.62, C["ink"])
dot(s7, 0.32, 6.53, C["mint"], 0.1)
txt(s7, "The framework is open by construction. Every vendor named here is an example, not a dependency.",
    0.54, 6.32, 12.5, 0.62, size=11, face=F["txt"], color=C["white"], anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
