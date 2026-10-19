"""Governed data foundation for campaigns and ROI: high-level and detailed views.

Brings together the GMDx workflow (marketing lead and clinical sources -> BigQuery -> joined master table ->
Marketing Contribution Margin dashboard), the ROI engine (data matching -> clinical integration -> attribution
-> holistic patient view), the ADDG governance flow, and the campaign path where Marketing Cloud pulls an
approved, tokenized file.

    python3 roi_governance.py [out.pptx]
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from governance_slide import (CYAN_D, GOLD, INK, ORANGE, R, RULE_C, TEMPLATE, lane, node, path,  # noqa: E402
                              tag)
from build_mad_ig_deck import (BLACK, CYAN, CYAN_T, GREEN, GREY, MAGENTA, MUTE, NAVY, VIOLET, VIOLET_T,  # noqa
                               WHITE, YELLOW, YELLOW_T, Canvas, set_notes)
from pptx import Presentation                                    # noqa: E402
from pptx.enum.shapes import MSO_SHAPE                            # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN                   # noqa: E402

OUT = os.path.join(HERE, "Governed Data for Campaigns and ROI.pptx")
MAG_T = "FBE3F2"
TEAL, TEAL_T = "00897B", "DDF2EF"

MARKETING = [("SFMC", "email + SMS engagement", "live"), ("SFDC", "marketing lead records", "live"),
             ("CallRail calls", "", "live"), ("Click to call", "", "wip"), ("Form submissions", "", "live"),
             ("Digitally booked appts", "", "live"), ("Direct mail lists", "new", "new"),
             ("Referrals", "", "live")]
CLINICAL = [("Epic", "appointments, encounters"), ("Strata", "cost and revenue"),
            ("athenahealth", "appointments, encounters")]
DOT = {"live": GREEN, "wip": ORANGE, "new": MAGENTA}
KEYS = ["Name", "Address", "Phone", "DOB", "Email", "EUID"]


def sources(cv, x, y, w, h, compact=False):
    cv.rect(x, y, w, 0.3, BLACK)
    cv.text(x + 0.1, y, w - 0.2, 0.3, [[("SOURCES", {"bold": True, "track": 150}),
                                       ("   land in BigQuery", {"color": "BFBFBF"})]],
            size=7, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    cv.rect(x, y + 0.3, w, h - 0.3, WHITE, MUTE, 0.75)
    yy = y + 0.38
    cv.text(x + 0.1, yy, w - 0.2, 0.18, "MARKETING LEADS", size=6.3, bold=True, color=CYAN_D, track=120)
    yy += 0.2
    step = 0.27 if compact else 0.25
    for n, d, st in MARKETING:
        cv.rect(x + 0.12, yy + 0.07, 0.08, 0.08, DOT[st], None, MSO_SHAPE.OVAL)
        cv.text(x + 0.26, yy, w - 0.34, step, [[(n, {"bold": True})] + ([("  " + d, {"color": MUTE,
                                                                                     "size": 5.6})] if d else [])],
                size=6.6, anchor=MSO_ANCHOR.MIDDLE)
        yy += step
    yy += 0.08
    cv.text(x + 0.1, yy, w - 0.2, 0.18, "CLINICAL", size=6.3, bold=True, color=TEAL, track=120)
    yy += 0.2
    for n, d in CLINICAL:
        cv.rect(x + 0.12, yy + 0.04, w - 0.24, 0.3, TEAL_T, TEAL, 0.75, R, 0.2)
        cv.text(x + 0.2, yy + 0.04, w - 0.36, 0.3, [[(n, {"bold": True}), ("  " + d, {"color": MUTE, "size": 5.6})]],
                size=6.6, anchor=MSO_ANCHOR.MIDDLE)
        yy += 0.36
    leg = [("live", "live"), ("wip", "in progress"), ("new", "new")]
    lx = x + 0.12
    for k, t in leg:
        cv.rect(lx, y + h - 0.17, 0.08, 0.08, DOT[k], None, MSO_SHAPE.OVAL)
        cv.text(lx + 0.11, y + h - 0.22, 0.6, 0.18, t, size=5.5, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)
        lx += 0.62


def boundary(cv, x, y, w, h, sub):
    cv.rect(x, y, w, h, None, NAVY, 1.25, R, 0.02, dash="dash")
    cv.text(x + 0.15, y + 0.03, w - 0.3, 0.24, [[("GOOGLE CLOUD · PHI BOUNDARY", {"bold": True, "track": 150}),
                                                  ("   " + sub, {"color": MUTE})]], size=7, anchor=MSO_ANCHOR.MIDDLE)


def principles(cv, y):
    pr = [("Governed before it is used", "Every column classified, steward-approved and protected."),
          ("One patient, one EUID", "Leads and visits matched inside Google Cloud, never outside."),
          ("PHI never leaves", "Marketing Cloud pulls tokenized IDs only; dashboards are masked."),
          ("Closed loop", "Engagement and booked visits come back to prove ROI.")]
    pw = (12.43 - 3 * 0.15) / 4
    for i, (t, d) in enumerate(pr):
        x = 0.45 + i * (pw + 0.15)
        cv.rect(x, y, pw, 0.56, GREY)
        cv.rect(x, y, 0.06, 0.56, (VIOLET, CYAN_D, NAVY, GREEN)[i])
        cv.text(x + 0.18, y + 0.04, pw - 0.26, 0.5, [[(t, {"bold": True, "size": 8})], [(d, {"color": MUTE})]],
                size=6.6, anchor=MSO_ANCHOR.MIDDLE, after=1)


def sfmc_column(cv, x, y, w, h, detailed):
    cv.rect(x, y, w, 0.3, BLACK)
    cv.text(x + 0.1, y, w - 0.2, 0.3, [[("MARKETING CLOUD", {"bold": True, "track": 150}),
                                       ("   pulls", {"color": "BFBFBF"})]], size=7, color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.rect(x, y + 0.3, w, h - 0.3, WHITE, MUTE, 0.75)
    nx, nw = x + 0.1, w - 0.2
    items = [("Scheduled import", "Automation Studio pulls the approved file, read-only key." if detailed else
              "pulls the approved file, read-only", CYAN_T, CYAN_D),
             ("Campaign data extension", "SubscriberKey (token) · channel · segment · send window · RunID"
              if detailed else "tokenized IDs and send-time fields", None, None),
             ("Journey Builder", "Consent, frequency caps, STOP re-checked at send." if detailed else
              "consent, frequency caps, STOP", WHITE, NAVY),
             ("SMS or email nudge", "No clinical wording, no identifiers in links." if detailed else
              "no clinical wording", WHITE, NAVY)]
    yy, nh, gap = y + 0.42, (h - 0.42 - 0.1 - 3 * 0.14) / 4, 0.14
    tops = []
    for k, (t, d, f, ln) in enumerate(items):
        if f is None:
            cv.frame(nx, yy, nw, nh, WHITE, 0.04)
        else:
            cv.rect(nx, yy, nw, nh, f, ln, 1.25, R, 0.1)
        cv.text(nx + 0.1, yy + 0.04, nw - 0.2, nh - 0.06, [[(t, {"bold": True, "size": 8})],
                                                         [(d, {"color": INK if detailed else MUTE})]],
                size=6.3, anchor=MSO_ANCHOR.MIDDLE, after=1)
        tops.append((yy, nh))
        if k < 3:
            cv.arrow(nx + nw / 2, yy + nh, nx + nw / 2, yy + nh + gap, NAVY, 1.25)
        yy += nh + gap
    return nx, nw, tops


# ================================================================== high-level
HL_NOTES = """SLIDE: Governed data for campaigns and ROI, the high-level view   (about 2 minutes)

KEY MESSAGE
One governed data foundation in Google Cloud serves two jobs: proving marketing ROI and running campaigns. Everything is governed before it is used, patients are matched inside Google Cloud, and Marketing Cloud only pulls an approved, tokenized file.

WALK IT
- Left, sources: marketing leads (SFMC engagement, SFDC lead records, CallRail, click to call, forms, digitally booked appointments, direct mail lists, referrals) and clinical data (Epic, Strata, athenahealth) all land in BigQuery.
- 1 Govern: DLP classifies every column, a data steward approves it in Collibra, and BigQuery masks it or limits raw access.
- 2 Match and join: the data matching engine links leads to patients on name, address, phone, date of birth, email and EUID, into one master table.
- 3a Measure ROI: attribution with lookback windows and contribution percentage, a holistic view by demography, service line, market and ministry, and the Marketing Contribution Margin dashboard.
- 3b Activate: an approved, minimum-necessary audience is published as a tokenized file; Marketing Cloud pulls it into a data extension and sends.
- The loop: engagement and booked visits come back as sources, so every campaign is measured."""


def high_level(prs, layout):
    s = prs.slides.add_slide(layout)
    cv = Canvas(s)
    cv.header("ASCENSION · GOVERNED DATA FOR CAMPAIGNS AND ROI",
              "One governed data foundation: prove marketing ROI, then run the next campaign",
              "Marketing and clinical data are governed and matched inside Google Cloud. The ROI engine measures; "
              "Marketing Cloud pulls only an approved, tokenized audience.")
    top, bot = 1.28, 5.62
    sources(cv, 0.45, top, 2.15, bot - top)
    boundary(cv, 2.75, top, 7.4, bot - top, "governed, matched and measured here")
    y0, ch = 1.66, bot - 1.66 - 0.14

    def stage(x, y, w, h, n, title, sub, col, tint, steps, out):
        cv.rect(x, y, w, h, WHITE, col, 1.5, R, 0.04)
        cv.rect(x, y, w, 0.5, tint, None, shape=R, radius=0.15)
        cv.rect(x + 0.1, y + 0.11, 0.28, 0.28, col, None, MSO_SHAPE.OVAL)
        cv.text(x + 0.1, y + 0.11, 0.28, 0.28, n, size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.45, y + 0.05, w - 0.5, 0.42, [[(title, {"bold": True, "track": 100, "size": 8.5})],
                                                  [(sub, {"color": MUTE, "size": 6.5})]], size=8,
                anchor=MSO_ANCHOR.MIDDLE)
        sy = y + 0.62
        sh = (h - 0.62 - 0.55 - 0.12 * (len(steps) - 1) - 0.1) / len(steps)
        for k, (t, gate) in enumerate(steps):
            cv.rect(x + 0.12, sy, w - 0.24, sh, YELLOW_T if gate else WHITE, GOLD if gate else RULE_C,
                    1.25 if gate else 0.75, R, 0.15, dash="dash" if gate else None)
            cv.text(x + 0.2, sy, w - 0.4, sh, t, size=7.6, bold=gate, anchor=MSO_ANCHOR.MIDDLE)
            if k < len(steps) - 1:
                cv.arrow(x + w / 2, sy + sh, x + w / 2, sy + sh + 0.12, col, 1.0)
            sy += sh + 0.12
        cv.rect(x + 0.12, y + h - 0.5, w - 0.24, 0.4, col, None, shape=R, radius=0.2)
        cv.text(x + 0.16, y + h - 0.5, w - 0.32, 0.4, out, size=6.8, bold=True,
                color=BLACK if col == GOLD else WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    gx, mxx, rx = 2.9, 5.2, 7.5
    sw = 2.15
    stage(gx, y0, sw, ch, "1", "GOVERN", "ADDG", VIOLET, VIOLET_T,
          [("Land in BigQuery, per source", False), ("DLP classifies every column", False),
           ("Data steward approves · Collibra", True), ("Mask or limit raw access in BigQuery", False)],
          "Every sensitive column protected")
    stage(mxx, y0, sw, ch, "2", "MATCH AND JOIN", "data matching engine", CYAN_D, CYAN_T,
          [("Keys: name · address · phone · DOB · email · EUID", False),
           ("Leads matched to patients, with a confidence threshold", False),
           ("Master table: one patient, one EUID", False)],
          "Leads + appointments + encounters")
    hh = (ch - 0.14) / 2
    # 3a measure, 3b activate (compact cards)
    for k, (n, title, sub, col, tint, body, out) in enumerate([
            ("3a", "MEASURE ROI", "attribution", MAGENTA, MAG_T,
             "Lookback windows · contribution % · holistic view by demography, service line, market, ministry",
             "Marketing Contribution Margin dashboard"),
            ("3b", "ACTIVATE", "approve and publish", GOLD, YELLOW_T,
             "Audience mart + quality gate · named approver in Looker · tokenized file to an outbound bucket",
             "Approved, tokenized file per run")]):
        y = y0 + k * (hh + 0.14)
        cv.rect(rx, y, 2.5, hh, WHITE, col, 1.5, R, 0.06)
        cv.rect(rx, y, 2.5, 0.45, tint, None, shape=R, radius=0.2)
        cv.rect(rx + 0.1, y + 0.09, 0.28, 0.28, col, None, MSO_SHAPE.OVAL)
        cv.text(rx + 0.1, y + 0.09, 0.28, 0.28, n, size=7.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(rx + 0.45, y + 0.03, 2.0, 0.4, [[(title, {"bold": True, "track": 100, "size": 8.5})],
                                                [(sub, {"color": MUTE, "size": 6.5})]], size=8,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(rx + 0.14, y + 0.55, 2.22, hh - 1.05, body, size=7.4)
        cv.rect(rx + 0.12, y + hh - 0.48, 2.26, 0.38, col, None, shape=R, radius=0.2)
        cv.text(rx + 0.16, y + hh - 0.48, 2.18, 0.38, out, size=6.8, bold=True,
                color=BLACK if col == GOLD else WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # flow arrows
    cv.arrow(2.6, y0 + ch / 2, gx, y0 + ch / 2, BLACK, 2.0)
    cv.arrow(gx + sw, y0 + ch / 2, mxx, y0 + ch / 2, NAVY, 2.0)
    cv.arrow(mxx + sw, y0 + hh / 2 + 0.3, rx, y0 + hh / 2, NAVY, 2.0)
    cv.arrow(mxx + sw, y0 + ch - hh / 2 - 0.3, rx, y0 + hh + 0.14 + hh / 2, NAVY, 2.0)
    # Marketing Cloud pulls
    mx, mw = 10.3, 2.58
    nx, nw, tops = sfmc_column(cv, mx, top, mw, bot - top, False)
    ya = y0 + hh + 0.14 + hh / 2
    path(cv, [(rx + 2.5, ya), (10.22, ya), (10.22, tops[0][0] + tops[0][1] / 2), (nx, tops[0][0] + tops[0][1] / 2)],
         BLACK, 2.0)
    tag(cv, rx + 2.5 - 0.62, y0 + hh + 0.14 + hh - 0.05 - 0.6, 0.56, "PULL", BLACK)
    # the loop
    ly = 5.78
    path(cv, [(nx + nw / 2, bot), (nx + nw / 2, ly), (1.52, ly), (1.52, bot)], GREEN, 1.75)
    tag(cv, 3.6, ly - 0.085, 4.4, "← ENGAGEMENT AND BOOKED VISITS COME BACK: EVERY CAMPAIGN MEASURED", GREEN)
    principles(cv, 6.12)
    set_notes(s, HL_NOTES)
    return s


# ================================================================== detailed
DT_NOTES = """SLIDE: Governed data for campaigns and ROI, the detailed design   (about 4 minutes)

1 Ingest and govern (ADDG): every source lands in its own raw dataset (CMEK). Cloud DLP proposes an infotype for each column; a data steward approves or rejects it in Collibra; the Data Policy Tagger (Cloud Run, dry-run first) applies SHA-256 masking or raw-access policies in BigQuery; each policy is recorded as a binding and verified against the live schema.
2 Match and join: the data matching engine runs inside the boundary. Deterministic match on EUID first, then name, address, phone, date of birth and email with a fixed confidence threshold; unmatched leads are kept, not guessed. The master table joins leads, appointments and encounters on one EUID.
3 Measure ROI: attribution applies lookback windows fixed before each campaign and computes contribution percentage; the holistic patient view slices by demography, service line, market and ministry; the Marketing Contribution Margin dashboard reads masked, aggregated views only. Lift is measured against a random holdout.
4 Activate: the audience mart holds one table per campaign with the minimum fields; a quality gate blocks untagged columns and stale sources; a named approver signs off in Looker; a Cloud Run job re-checks consent and suppression and writes one tokenized file per run ID to an outbound bucket.
Marketing Cloud pulls that file with an Automation Studio scheduled import (read-only key) into the campaign data extension; Journey Builder sends. SFMC engagement and booked visits come back as sources, closing the loop.

STATUS FROM THE SOURCE SLIDES
Click to call is in progress; direct mail lists are new. In ADDG: DLP and Collibra upstream, reviewed classifications and bindings built, tagger and column policies live in DEV, policy remover in progress."""


def detailed(prs, layout):
    s = prs.slides.add_slide(layout)
    cv = Canvas(s)
    cv.header("ASCENSION · GOVERNED DATA FOR CAMPAIGNS AND ROI",
              "Detailed design: govern, match, measure and activate inside Google Cloud",
              "Each source is governed before it is joined; the ROI engine and every campaign read the same "
              "governed master table.")
    top, bot = 1.27, 6.02
    sources(cv, 0.45, top, 2.1, bot - top, compact=True)
    bx, bw = 2.68, 7.52
    boundary(cv, bx, top, bw, bot - top, "VPC Service Controls · CMEK · audit logs · masked views for people")
    cw, gap = 2.31, 0.15
    cx = [bx + 0.15 + i * (cw + gap) for i in range(3)]
    nh = 0.66

    def lanebox(y, h, fill, line, num, title, sub, right=False):
        cv.rect(bx + 0.08, y, bw - 0.16, h, fill, line, 0.75, R, 0.03)
        cv.rect(bx + 0.15, y + 0.05, 0.19, 0.19, line, None, MSO_SHAPE.OVAL)
        cv.text(bx + 0.15, y + 0.05, 0.19, 0.19, num, size=6.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(bx + 0.4, y + 0.04, bw - 0.6, 0.21, [[(title, {"bold": True, "track": 150}), ("   " + sub,
                                                                                            {"color": MUTE})]],
                size=6.6, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT if right else None)

    L = [1.57, 2.7, 3.83, 4.96]
    lanebox(L[0], 1.03, VIOLET_T, VIOLET, "1", "INGEST AND GOVERN · ADDG", "classify, approve, protect, prove")
    y = L[0] + 0.29
    node(cv, cx[0], y, cw, nh, "Land in BigQuery",
         "One raw dataset per source, CMEK, partitioned by load date.", line=VIOLET)
    node(cv, cx[1], y, cw, nh, "DLP + steward review",
         "Cloud DLP proposes infotypes; a steward approves in Collibra.", "UPSTREAM", fill=YELLOW_T, line=GOLD,
         lw=1.5)
    node(cv, cx[2], y, cw, nh, "Column policies + bindings",
         "Tagger applies SHA-256 masking or raw access; each binding verified.", "LIVE IN DEV", fill=CYAN_T,
         line=CYAN_D, lw=1.25)
    for i in range(2):
        cv.arrow(cx[i] + cw, y + nh / 2, cx[i + 1], y + nh / 2, VIOLET, 1.25)

    lanebox(L[1], 1.03, CYAN_T, CYAN_D, "2", "MATCH AND JOIN", "data matching engine, inside the boundary",
            right=True)
    y2 = L[1] + 0.29
    cv.rect(cx[0], y2, cw, nh, WHITE, CYAN_D, 1.5, R, 0.08)
    cv.text(cx[0] + 0.09, y2 + 0.04, cw - 0.18, 0.2, "Data matching engine", size=7.6, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
    kw = (cw - 0.18 - 5 * 0.04) / 6
    for i, k in enumerate(KEYS):
        kx = cx[0] + 0.09 + i * (kw + 0.04)
        cv.rect(kx, y2 + 0.27, kw, 0.17, CYAN_T, CYAN_D, 0.5, R, 0.4)
        cv.text(kx, y2 + 0.27, kw, 0.17, k, size=5.2, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(cx[0] + 0.09, y2 + 0.46, cw - 0.18, 0.18, "EUID first, then fuzzy keys above a fixed threshold.",
            size=5.9, color=INK)
    node(cv, cx[1], y2, cw, nh, "Master table · one EUID",
         "Leads + appointments + encounters joined per patient; unmatched leads kept.", line=CYAN_D)
    node(cv, cx[2], y2, cw, nh, "Attribution logic",
         "Lookback windows fixed before each campaign; contribution percentage per touch.", fill=MAG_T,
         line=MAGENTA)
    for i in range(2):
        cv.arrow(cx[i] + cw, y2 + nh / 2, cx[i + 1], y2 + nh / 2, CYAN_D, 1.25)
    # governed data flows into matching
    path(cv, [(cx[2] + cw / 2, y + nh), (cx[2] + cw / 2, L[1] - 0.04), (cx[0] + cw / 2, L[1] - 0.04),
              (cx[0] + cw / 2, y2)], VIOLET, 1.5)

    lanebox(L[2], 1.03, MAG_T, MAGENTA, "3", "MEASURE ROI", "masked views only")
    y3 = L[2] + 0.29
    node(cv, cx[0], y3, cw, nh, "Lift vs holdout",
         "Nudged group against a random holdout, so ROI is lift, not raw bookings.", line=MAGENTA)
    node(cv, cx[1], y3, cw, nh, "Contribution Margin dashboard",
         "Looker Studio on masked views: revenue minus campaign cost, per run.", line=MAGENTA, lw=1.5)
    node(cv, cx[2], y3, cw, nh, "Holistic patient view",
         "By demography, service line, market and ministry.", line=MAGENTA)
    path(cv, [(cx[2] + cw / 2, y2 + nh), (cx[2] + cw / 2, y3)], MAGENTA, 1.25)
    cv.arrow(cx[2], y3 + nh / 2, cx[1] + cw, y3 + nh / 2, MAGENTA, 1.25)
    cv.arrow(cx[0] + cw, y3 + nh / 2, cx[1], y3 + nh / 2, MAGENTA, 1.25)

    lanebox(L[3], 1.0, YELLOW_T, GOLD, "4", "ACTIVATE", "approve, then publish for Marketing Cloud to pull",
            right=True)
    y4 = L[3] + 0.29
    node(cv, cx[0], y4, cw, nh, "Audience mart + quality gate",
         "One table per campaign, minimum fields; untagged or stale blocks the run.", fill=WHITE, line=GOLD,
         dash="dash", lw=1.25)
    node(cv, cx[1], y4, cw, nh, "Approve in Looker",
         "Size, segments, exclusions; a named approver signs off.", fill=YELLOW_T, line=GOLD, lw=1.5)
    cv.rect(cx[2], y4, cw, nh, WHITE, NAVY, 1.5, R, 0.08, dash="sysDot")
    cv.text(cx[2] + 0.09, y4 + 0.04, cw - 0.18, nh - 0.06,
            [[("Outbound bucket · tokenized file", {"bold": True, "size": 7.6})],
             [("Consent and suppression re-checked; one file per run ID; files expire.", {})],
             [("never name, DOB, diagnosis or MRN", {"bold": True, "color": MAGENTA})]], size=5.9, after=0.5)
    for i in range(2):
        cv.arrow(cx[i] + cw, y4 + nh / 2, cx[i + 1], y4 + nh / 2, GOLD, 1.25)
    # master table feeds the audience (through the gap between columns 1 and 2)
    gx_ = cx[0] + cw + gap / 2
    path(cv, [(cx[1] + 0.25, y2 + nh), (cx[1] + 0.25, L[2] + 0.02), (gx_, L[2] + 0.02), (gx_, L[3] + 0.02),
              (cx[0] + cw / 2, L[3] + 0.02), (cx[0] + cw / 2, y4)], CYAN_D, 1.5)

    # sources into the boundary
    cv.arrow(2.55, y + nh / 2, cx[0], y + nh / 2, BLACK, 2.0)
    # Marketing Cloud pulls
    mx, mw = 10.33, 2.55
    nx, nw, tops = sfmc_column(cv, mx, top, mw, bot - top, True)
    path(cv, [(cx[2] + cw, y4 + nh / 2), (10.26, y4 + nh / 2), (10.26, tops[0][0] + tops[0][1] / 2),
              (nx, tops[0][0] + tops[0][1] / 2)], BLACK, 2.0)
    # the loop: engagement back as a source
    ly = 6.13
    path(cv, [(nx + nw / 2, bot), (nx + nw / 2, ly), (1.5, ly), (1.5, bot)], GREEN, 1.75)
    tag(cv, 3.6, ly - 0.085, 4.6, "← SFMC ENGAGEMENT AND BOOKED VISITS COME BACK AS SOURCES: THE LOOP CLOSES",
        GREEN)
    principles(cv, 6.32)
    set_notes(s, DT_NOTES)
    return s


def main(out=OUT):
    prs = Presentation(TEMPLATE)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    high_level(prs, layout)
    detailed(prs, layout)
    prs.save(out)
    print("wrote", os.path.abspath(out))


if __name__ == "__main__":
    main(*sys.argv[1:])
