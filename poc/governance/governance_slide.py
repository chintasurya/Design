"""One slide: governed campaign data, from classified BigQuery columns to a Marketing Cloud data extension.

Sources: the ADDG governance flow (DLP -> Collibra steward -> tagger -> column policies) and the campaign
activation deck (campaign mart -> Looker approval -> Cloud Run activator -> Marketing Cloud, outcomes back).

    python3 governance_slide.py [out.pptx]
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "..", "deck")
sys.path.insert(0, DECK)

from pptx import Presentation                                    # noqa: E402
from pptx.enum.shapes import MSO_SHAPE                            # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN                   # noqa: E402

from build_mad_ig_deck import (BLACK, CYAN, CYAN_T, GREEN, GREY, MAGENTA, MUTE, NAVY, VIOLET, VIOLET_T,  # noqa: E402
                               WHITE, YELLOW, YELLOW_T, Canvas, set_notes)

TEMPLATE = os.path.join(DECK, "source", "Ascension Agentic Harness - final.pptx")
OUT = os.path.join(HERE, "Campaign Data Governance.pptx")
R = MSO_SHAPE.ROUNDED_RECTANGLE
CYAN_D, GOLD, INK, ORANGE = "0097AD", "C9A800", "2B2B2B", "FA8B35"
STATUS = {"UPSTREAM": (WHITE, MUTE, MUTE), "BUILT": (GREEN, GREEN, WHITE), "LIVE IN DEV": (CYAN, CYAN, BLACK),
          "IN PROGRESS": (ORANGE, ORANGE, WHITE)}

CX = [0.6 + i * 2.27 for i in range(4)]       # four columns inside the GCP boundary
CW, NH = 2.06, 0.72

NOTES = """SLIDE: Governed data first: how campaign data is protected before Marketing Cloud ever sees it   (about 4 minutes)

KEY MESSAGE
Before any campaign runs, the data is governed in Google Cloud. Every column is classified, approved by a data steward and protected inside BigQuery. Only a minimum-necessary, tokenized audience is approved in Looker and pushed into a Marketing Cloud data extension. PHI never leaves Google Cloud, and every visit booked comes back inside to prove the result.

WALK IT
1 Govern the data (ADDG, top band)
- Cloud DLP scans the BigQuery columns in the production projects and proposes an infotype for each one.
- A data steward approves or rejects each classification in Collibra; nothing changes before that sign-off.
- Outcomes sync into one reviewed-classifications table, the single list every job works from.
- A rollout gate enables production project by project; a person runs each promotion.
- The Data Policy Tagger (Cloud Run job, dry-run first) applies column policies: SHA-256 masking or raw-access only.
- Every applied policy is recorded as a binding and checked against the live schema. If a steward rejects, the policy remover clears it and the binding is retired.

2 Build the audience
- Authorized views apply the campaign criteria on policy-tagged data.
- The campaign mart holds one table per campaign: tokenized member ID, channel, segment, eligibility flags and run ID. No names, dates of birth or diagnoses.
- A quality gate blocks the run if a new column is untagged or a source is stale.
- In Looker, a named approver reviews audience size, segments and exclusions and approves.

3 Activate into Marketing Cloud
- The Cloud Run activator receives the signed approval, re-checks consent and suppression, and pushes the audience into the Marketing Cloud data extension in batches; the run ID makes it idempotent.
- Marketing Cloud sends through Journey Builder with consent, frequency caps and STOP handling. Messages carry no clinical wording and no identifiers in links.
- Booked visits come back into BigQuery outcome tables, inside the boundary, and are measured as lift against a random holdout.

STATUS (from the ADDG flow)
DLP and Collibra are upstream. The reviewed-classifications table, rollout gate and bindings are built. The tagger and the column policies are live in DEV. The policy remover is in progress.

DECISIONS TO CONFIRM
- Which Marketing Cloud products and SMS vendors are under Ascension's BAA.
- Push from Cloud Run (full control) or Data Cloud zero-copy (less code, less control).
- One approver, or a second approver above a set audience size.
- A random holdout to prove lift, or report bookings only."""


def node(cv, x, y, w, h, title, body, status=None, fill=WHITE, line=NAVY, dash=None, lw=0.9):
    cv.rect(x, y, w, h, fill, line, lw, R, 0.08, dash=dash)
    tw = w - 0.16
    if status:
        f, ln, col = STATUS[status]
        sw = 0.14 + 0.055 * len(status)
        cv.rect(x + w - sw - 0.07, y + 0.07, sw, 0.15, f, ln, 0.5, R, 0.5)
        cv.text(x + w - sw - 0.07, y + 0.07, sw, 0.15, status, size=4.8, bold=True, color=col,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, track=60)
        tw -= sw + 0.05
    cv.text(x + 0.09, y + 0.05, tw, 0.2, title, size=7.6, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(x + 0.09, y + 0.26, w - 0.18, h - 0.29, body, size=6.1, color=INK)


def path(cv, pts, color=NAVY, w=1.25, dash=None):
    for k, (p, q) in enumerate(zip(pts, pts[1:])):
        cv.arrow(p[0], p[1], q[0], q[1], color, w, dash=dash, head=(k == len(pts) - 2))


def tag(cv, x, y, w, t, fill=BLACK, color=WHITE):
    cv.rect(x, y, w, 0.17, fill, None, shape=R, radius=0.5)
    cv.text(x, y, w, 0.17, t, size=5.6, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            track=60)


def lane(cv, y, h, fill, line, num, title, sub, x0=0.6):
    cv.rect(0.52, y, 8.98, h, fill, line, 0.75, R, 0.03)
    cv.rect(x0, y + 0.06, 0.2, 0.2, line, None, MSO_SHAPE.OVAL)
    cv.text(x0, y + 0.06, 0.2, 0.2, str(num), size=7, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(x0 + 0.26, y + 0.05, 9.4 - x0 - 0.26, 0.22, [[(title, {"bold": True, "track": 150}), ("   " + sub, {"color": MUTE})]],
            size=7, anchor=MSO_ANCHOR.MIDDLE)


def governance_slide(prs, layout):
    s = prs.slides.add_slide(layout)
    cv = Canvas(s)
    cv.header("ASCENSION · CAMPAIGN DATA GOVERNANCE",
              "Govern the data first: from classified columns to a Marketing Cloud data extension",
              "PHI stays in Google Cloud. Every column is classified, steward-approved and protected before a "
              "campaign uses it; Marketing Cloud gets only a tokenized, approved audience.")

    # ------------------------------------------------------------ the GCP PHI boundary
    cv.rect(0.45, 1.27, 9.12, 4.87, None, NAVY, 1.25, R, 0.02, dash="dash")
    cv.text(0.6, 1.3, 9.0, 0.24, [[("GOOGLE CLOUD · PHI BOUNDARY", {"bold": True, "track": 150}),
                                   ("   VPC Service Controls · CMEK · audit logs to a locked sink · no direct "
                                    "marketing access", {"color": MUTE})]], size=7, anchor=MSO_ANCHOR.MIDDLE)

    # lane 1: govern the data (ADDG)
    lane(cv, 1.58, 2.04, VIOLET_T, VIOLET, 1, "GOVERN THE DATA · ADDG",
         "classify, approve, protect, prove: governance tier asc-ahnat-datarepository")
    r1, r2 = 1.88, 2.82
    node(cv, CX[0], r1, CW, NH, "Source data",
         "17 production projects in BigQuery (asc-datalake): clinical and engagement tables.", line=VIOLET)
    node(cv, CX[1], r1, CW, NH, "DLP classification",
         "Cloud DLP scans every column and proposes an infotype: name, DOB, MRN, phone.", "UPSTREAM", line=VIOLET)
    node(cv, CX[2], r1, CW, NH, "Steward review · Collibra",
         "A data steward approves or rejects each classification. Nothing changes before.", "UPSTREAM",
         fill=YELLOW_T, line=GOLD, lw=1.5)
    node(cv, CX[3], r1, CW, NH, "Reviewed classifications",
         "One governed table, keyed by project, dataset, table and column: the single list.", "BUILT",
         line=VIOLET)
    node(cv, CX[3], r2, CW, NH, "Rollout gate · per project",
         "Production enabled project by project; a person runs every promotion.", "BUILT",
         fill=YELLOW_T, line=GOLD, dash="dash", lw=1.25)
    node(cv, CX[2], r2, CW, NH, "Data Policy Tagger",
         "Cloud Run job, per-environment config, dry-run first. Decides each column's policy.", "LIVE IN DEV",
         line=VIOLET)
    node(cv, CX[1], r2, CW, NH, "Column policies",
         "In BigQuery: SHA-256 masking or raw-access only. If rejected, the policy remover (in progress) clears it.",
         "LIVE IN DEV",
         fill=CYAN_T, line=CYAN_D, lw=1.25)
    node(cv, CX[0], r2, CW, NH, "Bindings + validation",
         "Every policy recorded and read back from the live schema: the audit trail.", "BUILT", line=VIOLET)
    ym1, ym2 = r1 + NH / 2, r2 + NH / 2
    for i in range(3):
        cv.arrow(CX[i] + CW, ym1, CX[i + 1], ym1, VIOLET, 1.25)
    path(cv, [(CX[3] + CW / 2, r1 + NH), (CX[3] + CW / 2, r2)], VIOLET)
    for i in (3, 2, 1):
        cv.arrow(CX[i], ym2, CX[i - 1] + CW, ym2, VIOLET, 1.25)
    tag(cv, CX[2] + 0.3, r1 + NH + 0.04, 1.46, "REJECTED → POLICY REMOVED", WHITE, MAGENTA)

    # lane 2: build the audience
    lane(cv, 3.72, 1.06, CYAN_T, CYAN_D, 2, "BUILD THE AUDIENCE", "minimum necessary, from protected data only",
         x0=CX[2] - 0.5)
    a2 = 3.98
    node(cv, CX[0], a2, CW, NH, "Authorized views",
         "Campaign criteria on policy-tagged data. Masked columns stay masked.", line=CYAN_D)
    node(cv, CX[1], a2, CW, NH, "Campaign mart",
         "One table per campaign; row-level security, CMEK, DLP scan, partitioned by run.", line=CYAN_D)
    node(cv, CX[2], a2, CW, NH, "Quality gate",
         "Blocks the run if a column is untagged or a source is stale (tagger + CI check).",
         fill=YELLOW_T, line=GOLD, dash="dash", lw=1.25)
    node(cv, CX[3], a2, CW, NH, "Approve in Looker",
         "Audience size, segments, exclusions. A named approver clicks Approve.", fill=YELLOW_T, line=GOLD, lw=1.5)
    for i in range(3):
        cv.arrow(CX[i] + CW, a2 + NH / 2, CX[i + 1], a2 + NH / 2, CYAN_D, 1.25)
    # governance feeds the audience
    path(cv, [(CX[0] + CW / 2, r2 + NH), (CX[0] + CW / 2, a2)], VIOLET, 1.75)
    tag(cv, CX[0] + CW / 2 + 0.1, 3.77, 1.95, "ONLY PROTECTED COLUMNS FLOW ON", VIOLET)

    # lane 3: activate and measure
    lane(cv, 4.88, 1.13, GREY, MUTE, 3, "ACTIVATE AND MEASURE", "tokenized IDs out, outcomes back in")
    a3 = 5.14
    node(cv, CX[0], a3, CW, NH, "Lift vs holdout",
         "Nudged group against a random holdout: lift and contribution margin, not raw counts.", line=MUTE)
    node(cv, CX[1], a3, CW, NH, "Outcome tables",
         "Booked visits (Epic, Athena) matched to the run, back inside the boundary.", line=MUTE)
    cv.rect(CX[2], a3, CW, NH, WHITE, NAVY, 0.9, R, 0.08, dash="sysDot")
    cv.text(CX[2] + 0.09, a3 + 0.05, CW - 0.18, NH - 0.08,
            [[("What leaves the boundary", {"bold": True, "size": 7.6})],
             [("member token · channel · segment · eligibility flags · send window · run ID", {"size": 6.1})],
             [("never name, DOB, diagnosis or MRN", {"size": 6.1, "bold": True, "color": MAGENTA})]], size=6.1,
            after=1)
    node(cv, CX[3], a3, CW, NH, "Cloud Run activator",
         "Signed approval; re-checks consent and suppression; pushes in batches, run ID = no duplicates.",
         fill=WHITE, line=NAVY, lw=1.5)
    path(cv, [(CX[3] + CW / 2, a2 + NH), (CX[3] + CW / 2, a3)], GOLD, 1.5)
    cv.arrow(CX[1], a3 + NH / 2, CX[0] + CW, a3 + NH / 2, MUTE, 1.25)

    # ------------------------------------------------------------ Marketing Cloud, outside the boundary
    mx, mw = 9.75, 3.13
    cv.rect(mx, 1.27, mw, 0.3, BLACK)
    cv.text(mx + 0.1, 1.27, mw - 0.2, 0.3, [[("MARKETING CLOUD", {"bold": True, "track": 150}),
                                            ("   outside the boundary", {"color": "BFBFBF"})]],
            size=7, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    cv.rect(mx, 1.57, mw, 4.57, WHITE, MUTE, 0.75)
    nx, nw = mx + 0.1, mw - 0.2
    # the data extension: the hand-off point
    fx, fy, fw, fh = cv.frame(nx, 1.68, nw, 1.62, WHITE, 0.04)
    cv.text(fx + 0.1, fy + 0.05, fw - 0.2, 0.22, "Campaign data extension", size=8.5, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(fx + 0.1, fy + 0.27, fw - 0.2, 0.2, [[("← from the activator: ", {"bold": True, "color": BLACK}),
                                                  ("tokenized IDs only, nothing else", {})]], size=6.2, color=MUTE)
    fields = [("SubscriberKey", "member token"), ("Channel", "SMS / email"), ("Segment", "campaign variant"),
              ("Send window", "send-time fields"), ("RunID", "ties back to the approval")]
    for i, (f_, d_) in enumerate(fields):
        yy = fy + 0.5 + i * 0.19
        cv.text(fx + 0.12, yy, 0.95, 0.18, f_, size=6.3, bold=True, font="Courier New", anchor=MSO_ANCHOR.MIDDLE)
        cv.text(fx + 1.08, yy, fw - 1.2, 0.18, d_, size=6.2, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    node(cv, nx, 3.45, nw, 0.6, "Journey Builder",
         "Timing, frequency caps, variants. Consent and STOP re-checked at send.", line=NAVY)
    node(cv, nx, 4.17, nw, 0.6, "Patient nudge · SMS or email",
         "No clinical wording, no identifiers in links (TCPA consent).", line=NAVY)
    node(cv, nx, 4.89, nw, 0.6, "Visit booked",
         "Online or by phone; AWV billed as G0438 / G0439.", line=NAVY)
    for y1, y2 in ((3.3, 3.45), (4.05, 4.17), (4.77, 4.89)):
        cv.arrow(nx + nw / 2, y1, nx + nw / 2, y2, NAVY, 1.25)

    # activator -> data extension (crosses the boundary)
    path(cv, [(CX[3] + CW, a3 + NH / 2), (9.66, a3 + NH / 2), (9.66, 2.45), (nx, 2.45)], BLACK, 2.0)
    # outcomes back inside
    by = 5.98
    path(cv, [(nx + nw / 2, 5.49), (nx + nw / 2, by), (CX[1] + CW / 2, by), (CX[1] + CW / 2, a3 + NH)], GREEN, 1.75)
    tag(cv, 5.3, by - 0.085, 2.05, "← BOOKED VISITS BACK INSIDE", GREEN)

    # ------------------------------------------------------------ built-in controls
    cy = 6.24
    cv.text(0.45, cy, 1.35, 0.6, [[("BUILT-IN", {"bold": True, "track": 150})], [("CONTROLS", {"bold": True,
                                                                                         "track": 150})]],
            size=7, anchor=MSO_ANCHOR.MIDDLE)
    controls = [("Steward approval first", "no policy before Collibra sign-off"),
                ("Dry run before apply", "every run can report without changing"),
                ("Human-gated promotion", "project by project; never self-promotes"),
                ("Auditable bindings", "each policy recorded and verified"),
                ("Named approver", "Looker approval before any send"),
                ("Minimum necessary", "tokenized fields only to Marketing Cloud")]
    gw = (12.88 - 1.85 - 5 * 0.1) / 6
    for i, (t_, d_) in enumerate(controls):
        x = 1.85 + i * (gw + 0.1)
        cv.rect(x, cy, gw, 0.6, WHITE, GREEN if i < 4 else GOLD, 0.75, R, 0.1)
        cv.rect(x, cy, 0.05, 0.6, GREEN if i < 4 else GOLD)
        cv.text(x + 0.12, cy + 0.05, gw - 0.18, 0.52, [[(t_, {"bold": True, "size": 7})], [(d_, {"color": MUTE})]],
                size=6.1, after=1, anchor=MSO_ANCHOR.MIDDLE)
    set_notes(s, NOTES)
    return s


def main(out=OUT):
    prs = Presentation(TEMPLATE)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    governance_slide(prs, layout)
    prs.save(out)
    print("wrote", os.path.abspath(out))


if __name__ == "__main__":
    main(*sys.argv[1:])
