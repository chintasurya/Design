"""Slide 4 as an interactive walk-through: the full design first, then one step per click.

Each click (Enter, Down arrow, Space or mouse) fades out the previous step and brings in the next:
a spotlight (everything else dimmed), orange frames on the parts being discussed, then that step's
arrows, and a caption that follows the talk track. The last click clears back to the full design.
Built with native PowerPoint animations, so it runs in Slide Show without any add-in.
"""
from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

import review_slides
from build_mad_ig_deck import BLACK, WHITE, Canvas, a, set_notes

ORANGE = "FA8B35"
GREEN_OK = "2E9E1F"
RED = "FF0069"
AREA = (0.4, 1.22, 12.93, 6.95)          # the part of the slide the spotlight dims
R = MSO_SHAPE.ROUNDED_RECTANGLE
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

# gate chips on the bottom strip (same geometry as review_slides.hld_slide)
CHIP = lambda i: (1.55 + i * 0.87, 6.07, 0.82, 0.42)

STEPS = [
    dict(title="Start in Salesforce",
         text="Story NS-1234: change an HFN validation rule. The Jira key, pod and design go in the LWC; Apex "
              "and a Named Credential reach the harness; Codex runs the ADLC skills.",
         spots=[(0.5, 2.15, 2.12, 2.52), (2.6, 3.75, 0.5, 0.3)],
         frames=[(0.55, 2.22, 2.02, 1.02), (0.55, 3.5, 2.02, 1.12)],
         arrows=[[(1.56, 3.24), (1.56, 3.5)], [(2.57, 3.9), (3.05, 3.9)]],
         chips=[], callout=(3.35, 2.25, 4.1)),
    dict(title="Context first",
         text="Before any code, agents ask the knowledge graph what exists, what it touches and who owns it. "
              "Read-only connectors, cited facts, a bounded slice, never an org dump.",
         spots=[(2.9, 1.42, 7.39, 1.38), (10.5, 1.55, 2.38, 1.6)],
         frames=[(2.95, 1.47, 7.29, 1.28), (10.5, 1.6, 2.38, 1.55)],
         arrows=[[(5.47, 3.12), (5.47, 2.65)], [(10.5, 2.2), (10.24, 2.2)]],
         chips=[], callout=(3.0, 3.4, 4.6)),
    dict(title="Orchestrated, step by step",
         text="Cloud Workflows coordinates the story. One agent job per step: requirements, design, build, "
              "review, test, release. Checkpoints mean a failed run resumes where it stopped.",
         spots=[(2.98, 3.02, 4.95, 1.2)],
         frames=[(3.05, 3.12, 1.82, 1.05), (4.97, 3.12, 2.9, 1.05)],
         arrows=[[(4.87, 3.62), (4.97, 3.62)]],
         chips=[], callout=(3.0, 4.35, 4.85)),
    dict(title="Act only through the gateway",
         text="Git, Salesforce and Copado are touched only through the MCP gateway, with the identity allowed "
              "for that step. The Change Guardian checks cross-pod ownership and impact.",
         spots=[(4.92, 3.02, 5.27, 1.9), (10.5, 3.18, 2.38, 2.2)],
         frames=[(7.97, 3.12, 2.17, 0.62), (7.97, 3.82, 2.17, 1.05), (10.5, 3.2, 2.38, 2.15)],
         arrows=[[(7.87, 4.12), (7.97, 4.12)], [(10.14, 4.32), (10.5, 4.32)]],
         chips=[], callout=(0.5, 3.15, 4.3)),
    dict(title="Evals on every step: pass",
         text="Every output is scored before it moves on. A pass goes to the gate, where people see the "
              "evidence and the eval scores, and approve.",
         spots=[(2.98, 3.02, 4.95, 1.9)],
         frames=[(4.97, 4.25, 2.9, 0.62), (3.05, 4.25, 1.82, 0.62)],
         arrows=[[(6.42, 4.17), (6.42, 4.25)], [(4.97, 4.56), (4.87, 4.56)]],
         chips=[(4.45, 4.03, 0.95, "PASS → GATE", GREEN_OK)], callout=(8.1, 3.1, 4.7)),
    dict(title="Fail: a bounded retry",
         text="A failed check goes back to the agent with the reason, at most five times. Then the step stops, "
              "nothing is released, and a named owner takes over. No infinite loops.",
         spots=[(2.98, 3.02, 4.95, 1.9)],
         frames=[(4.97, 4.25, 2.9, 0.62), (4.97, 3.12, 2.9, 1.05)],
         arrows=[[(7.87, 4.56), (8.3, 4.56), (8.3, 3.62), (7.87, 3.62)], [(6.42, 4.87), (6.42, 5.05)]],
         chips=[(8.38, 3.95, 1.3, "FAIL · RETRY ≤ 5", RED),
                (5.35, 5.06, 2.15, "AFTER 5 → NAMED OWNER", RED)],
         callout=(8.1, 2.15, 4.7), arrow_color=RED),
    dict(title="People approve, the owning team deploys",
         text="People approve requirements, architecture, review and UAT. CAB still controls the release, and "
              "the owning team, not an agent, deploys to production.",
         spots=[(1.5, 6.02, 11.38, 0.52), (10.55, 5.3, 2.28, 0.66)],
         frames=[CHIP(i) for i in (1, 3, 6, 8, 10, 11)] + [(10.6, 5.36, 2.18, 0.55)],
         arrows=[[(1.55 + 11 * 0.87 + 0.41, 6.07), (1.55 + 11 * 0.87 + 0.41, 5.91)]],
         chips=[], callout=(2.95, 5.2, 6.9)),
    dict(title="Evidence on every change",
         text="Every step lands in the change passport, and cycle time, defects, rework and cost per change are "
              "reported per pod. Not just automation: a controlled, auditable delivery process.",
         spots=[(5.5, 5.02, 3.15, 0.9)],
         frames=[(5.55, 5.07, 1.6, 0.8), (7.25, 5.07, 1.35, 0.8)],
         arrows=[],
         chips=[], callout=(4.6, 4.3, 5.2)),
]


def _alpha(shape, val):
    clr = shape._element.spPr.find(a("solidFill")).find(a("srgbClr"))
    etree.SubElement(clr, a("alpha"), val=str(val))


def _dim_cells(holes):
    """Rectangles covering AREA minus the holes (grid decomposition, merged along each row)."""
    x0, y0, x1, y1 = AREA
    xs = sorted({x0, x1} | {min(max(h[0], x0), x1) for h in holes} | {min(max(h[0] + h[2], x0), x1) for h in holes})
    ys = sorted({y0, y1} | {min(max(h[1], y0), y1) for h in holes} | {min(max(h[1] + h[3], y0), y1) for h in holes})

    def covered(cx, cy):
        return any(h[0] <= cx <= h[0] + h[2] and h[1] <= cy <= h[1] + h[3] for h in holes)

    out = []
    for j in range(len(ys) - 1):
        run = None
        for i in range(len(xs) - 1):
            cx, cy = (xs[i] + xs[i + 1]) / 2, (ys[j] + ys[j + 1]) / 2
            if covered(cx, cy):
                if run:
                    out.append(run)
                run = None
            elif run:
                run = (run[0], ys[j], xs[i + 1] - run[0], ys[j + 1] - ys[j])
            else:
                run = (xs[i], ys[j], xs[i + 1] - xs[i], ys[j + 1] - ys[j])
        if run:
            out.append(run)
    return out


def _path(cv, pts, color, w=3.0):
    for k, (p, q) in enumerate(zip(pts, pts[1:])):
        cv.arrow(p[0], p[1], q[0], q[1], color, w, head=(k == len(pts) - 2))


def _build_step(slide, n, total, st):
    color = st.get("arrow_color", ORANGE)
    hl = slide.shapes.add_group_shape()
    cv = Canvas(hl)
    for x, y, w, h in _dim_cells(st["spots"]):
        _alpha(cv.rect(x, y, w, h, WHITE), 72000)
    for x, y, w, h in st["frames"]:
        cv.rect(x - 0.04, y - 0.04, w + 0.08, h + 0.08, None, ORANGE, 3.0, R, 0.08)
    cx, cy, cw = st["callout"]
    cv.rect(cx, cy, cw, 0.72, BLACK, ORANGE, 1.0, R, 0.12)
    cv.text(cx + 0.14, cy + 0.07, cw - 0.28, 0.62,
            [[("%d / %d   " % (n, total), {"color": ORANGE}), (st["title"].upper(), {"track": 120})],
             [(st["text"], {"bold": False, "size": 7.5, "color": "E6E6E6"})]],
            size=8, bold=True, color=WHITE, after=2)
    hl.name = "Step %d highlight" % n
    ar = None
    if st["arrows"] or st["chips"]:
        ar = slide.shapes.add_group_shape()
        cva = Canvas(ar)
        for pts in st["arrows"]:
            _path(cva, pts, color)
        for x, y, w, t, c in st["chips"]:
            cva.rect(x, y, w, 0.2, c, None, shape=R, radius=0.5)
            cva.text(x, y, w, 0.2, t, size=6.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                     track=60)
        ar.name = "Step %d arrows" % n
    return hl, ar


# ------------------------------------------------------------------ animation XML
class _Ids:
    def __init__(self):
        self.n = 2

    def __call__(self):
        self.n += 1
        return str(self.n)


def _q(tag):
    return "{%s}%s" % (P_NS, tag)


def _el(parent, tag, **attrs):
    return etree.SubElement(parent, _q(tag), {k: str(v) for k, v in attrs.items()})


def _effect(parent, ids, spid, entr, node, delay=0, filt="fade", dur=350, preset=10, subtype=0):
    par = _el(parent, "par")
    ctn = _el(par, "cTn", id=ids(), presetID=preset, presetClass="entr" if entr else "exit",
              presetSubtype=subtype, fill="hold", grpId=0, nodeType=node)
    _el(_el(ctn, "stCondLst"), "cond", delay=delay)
    kids = _el(ctn, "childTnLst")

    def setvis(val, d):
        s = _el(kids, "set")
        b = _el(s, "cBhvr")
        c = _el(b, "cTn", id=ids(), dur=1, fill="hold")
        _el(_el(c, "stCondLst"), "cond", delay=d)
        _el(_el(b, "tgtEl"), "spTgt", spid=spid)
        _el(_el(b, "attrNameLst"), "attrName").text = "style.visibility"
        _el(_el(s, "to"), "strVal", val=val)

    if entr:
        setvis("visible", 0)
    ae = _el(kids, "animEffect", transition="in" if entr else "out", filter=filt)
    b = _el(ae, "cBhvr")
    _el(b, "cTn", id=ids(), dur=dur)
    _el(_el(b, "tgtEl"), "spTgt", spid=spid)
    if not entr:
        setvis("hidden", dur - 1)


def _click(main_children, ids, exits, enters, arrows):
    par = _el(main_children, "par")
    ctn = _el(par, "cTn", id=ids(), fill="hold")
    _el(_el(ctn, "stCondLst"), "cond", delay="indefinite")
    kids = _el(ctn, "childTnLst")
    # group 1: exits of the previous step and the new highlight, together on the click
    g1 = _el(kids, "par")
    c1 = _el(g1, "cTn", id=ids(), fill="hold")
    _el(_el(c1, "stCondLst"), "cond", delay=0)
    k1 = _el(c1, "childTnLst")
    first = True
    for spid in exits:
        _effect(k1, ids, spid, False, "clickEffect" if first else "withEffect", dur=250)
        first = False
    for spid in enters:
        _effect(k1, ids, spid, True, "clickEffect" if first else "withEffect")
        first = False
    # group 2: the step's arrows wipe in after the highlight
    if arrows:
        g2 = _el(kids, "par")
        c2 = _el(g2, "cTn", id=ids(), fill="hold")
        _el(_el(c2, "stCondLst"), "cond", delay=400)
        k2 = _el(c2, "childTnLst")
        for spid in arrows:
            _effect(k2, ids, spid, True, "afterEffect", filt="wipe(left)", dur=600, preset=22, subtype=8)


def _timing(slide, steps):
    ids = _Ids()
    timing = etree.Element(_q("timing"))
    root = _el(_el(_el(timing, "tnLst"), "par"), "cTn", id=1, dur="indefinite", restart="never", nodeType="tmRoot")
    seq = _el(_el(root, "childTnLst"), "seq", concurrent=1, nextAc="seek")
    main = _el(seq, "cTn", id=2, dur="indefinite", nodeType="mainSeq")
    kids = _el(main, "childTnLst")
    prev = []
    for hl, ar in steps:
        _click(kids, ids, prev, [hl.shape_id], [ar.shape_id] if ar else [])
        prev = [hl.shape_id] + ([ar.shape_id] if ar else [])
    _click(kids, ids, prev, [], [])                      # last click: back to the full design
    pc = _el(_el(seq, "prevCondLst"), "cond", evt="onPrev", delay=0)
    _el(_el(pc, "tgtEl"), "sldTgt")
    nc = _el(_el(seq, "nextCondLst"), "cond", evt="onNext", delay=0)
    _el(_el(nc, "tgtEl"), "sldTgt")
    sld = slide._element
    old = sld.find(_q("timing"))
    if old is not None:
        sld.remove(old)
    anchor = sld.find(_q("transition"))
    if anchor is None:
        anchor = sld.find(_q("clrMapOvr"))
    anchor.addnext(timing)


INTERACTIVE_NOTES = """SLIDE 4  |  Codex to Production: one story through the harness   (about 2 min 45 sec)
MAIN POINT: the complete governed lifecycle of one change.
Advance with Enter, Down arrow, Space or a click. The slide opens on the full design; each click spotlights the next part.

[OPEN, full diagram] This is the main execution slide. Imagine a Network Services story requiring a change to an HFN validation rule.

[CLICK 1  Start in Salesforce] A team member starts from the Salesforce Agentic Harness console, provides the Jira key, selects the pod and adds the design context. Through Apex and a Named Credential, the request reaches the harness, and Codex uses the delivery skills for the run.

[CLICK 2  Context first] Before proposing code, the agents query the knowledge graph to understand existing objects, automations, dependencies and ownership. That graph is populated from read-only connectors, with source references and snapshots. We deliberately provide a bounded context slice rather than a full org dump.

[CLICK 3  Orchestrated] In the center, Cloud Workflows coordinates the story. Agent workers handle requirements, design, development, review, testing and release preparation, with checkpoints and resumable runs.

[CLICK 4  Gateway] Any action against Git, Salesforce or Copado passes through the MCP gateway and the allowed identity for that step. The Change Guardian checks cross-pod ownership and impact.

[CLICK 5  Evals: pass] At every step we evaluate the output. A pass goes to the gate, with the evidence.

[CLICK 6  Evals: fail and retry] If a check fails, the harness sends the result back for correction, within a bounded retry limit: five attempts, then a named owner.

[CLICK 7  People approve] People still approve the requirements, architecture, review and UAT, and CAB still controls the release. The owning team, not an agent, deploys to production.

[CLICK 8  Evidence] We then keep the evidence in a change passport and report cycle time, defects, rework and cost per change.

[CLICK 9  back to the full design] So the end-to-end story is not just automation: it is a controlled, auditable delivery process."""


def hld_interactive_slide(prs, layout, only=None):
    """only=k builds a static storyboard frame: just step k visible (0 = the full design)."""
    s = review_slides.hld_slide(prs, layout)
    if only is None:
        built = [_build_step(s, i + 1, len(STEPS), st) for i, st in enumerate(STEPS)]
        _timing(s, built)
    elif only:
        _build_step(s, only, len(STEPS), STEPS[only - 1])
    set_notes(s, INTERACTIVE_NOTES)
    return s
