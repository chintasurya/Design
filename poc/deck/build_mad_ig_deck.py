"""Put the Ascension MAD Approach deck on the Insight Global template, and add
the architecture and execution-flow slides.

    python3 build_mad_ig_deck.py [source.pptx] [out.pptx]

Defaults: source/Ascension MAD Approach - source.pptx  ->  ../Ascension MAD Approach.pptx

Two jobs, kept apart so the first can never change a word:

1. Restyle. Every existing slide keeps its text, order and notes. What changes
   is the template: white ground, IG palette (black, navy, cyan, yellow,
   magenta; violet, orange and green sparingly), Trenda IG Text headlines in
   bold and large stats in Light, the IG table style, and the ●●●IG mark with a
   page number in the footer of every content slide. Content that ran into the
   footer is scaled about the slide's centre line, fonts and all, until it
   clears it, so nothing re-wraps.

2. Two new slides after "One agent team": the target architecture and the
   execution flow. Differentiators (what an off-the-shelf coding agent does
   not do) carry a border in the three IG logo colours and a DIFFERENTIATOR
   tab; every major block says how mature it is. The client questions these
   slides answer, and where to point for each, live in the speaker notes only.

Palette and type are read from the IG PowerPoint Template User Guide (June
2026): primary FFD700 · 00D6F2 · FF0069 · 000000 · 00273B · FFFFFF,
secondary FA8B35 · D214AA · 4527A0 · 56E13B · F2F2F2.
"""
import copy
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "source", "Ascension MAD Approach - source.pptx")
OUT = os.path.join(HERE, "..", "Ascension MAD Approach.pptx")
MARK_BLACK = os.path.join(HERE, "assets", "ig-mark-black.png")

EMU = 914400
SLIDE_W = 13.333
FOOTER_TOP = 7.03          # where the template puts the ●●●IG mark
CONTENT_FLOOR = 6.93       # content must end above this
NEW_AFTER = 2              # the two new slides follow slide 2

# ---------------------------------------------------------------- palette
YELLOW, CYAN, MAGENTA = "FFD700", "00D6F2", "FF0069"
BLACK, NAVY, WHITE = "000000", "00273B", "FFFFFF"
ORANGE, PURPLE, VIOLET, GREEN, GREY = "FA8B35", "D214AA", "4527A0", "56E13B", "F2F2F2"
MUTE = "595959"
RULE = "D9D9D9"
CYAN_T, YELLOW_T, MAGENTA_T, GREEN_T, VIOLET_T = "E0FAFE", "FFF7CC", "FFE6F0", "EAFBE6", "ECE9F5"

# old colour -> IG colour, by where the colour is used
FILL = {
    "00757F": CYAN, "0E8A7D": CYAN, "1B6BC0": NAVY, "12233B": BLACK, "0E1B2E": BLACK,
    "C2185B": MAGENTA, "C62828": MAGENTA, "C77400": YELLOW, "2E7D32": GREEN,
    "7A4FC0": VIOLET, "4527A0": VIOLET, "9AA9BC": "A6A6A6", "9AA7B5": "A6A6A6", "5C6E85": MUTE,
    "E3F2F4": CYAN_T, "E5F4F1": CYAN_T, "E6EBF2": GREY, "EEF2F7": GREY, "F4F7FC": WHITE,
    "F8FAFD": WHITE, "EAF1FB": GREY, "FBE9F0": MAGENTA_T, "FBF1E3": YELLOW_T,
    "FDF1DE": YELLOW_T, "E9F3EA": GREEN_T, "E8F3E9": GREEN_T, "FDECEC": MAGENTA_T,
    "FBEAEA": MAGENTA_T, "F0EBFA": VIOLET_T,
}
LINE = dict(FILL)
LINE.update({
    "DCE4EE": RULE, "C7D2DE": "BFBFBF", "B9C4D0": "BFBFBF", "EEF2F7": "E6E6E6",
    "E3F2F4": "99EFFA", "FBE9F0": "FFB3D2", "EAF1FB": RULE, "F4F7FC": RULE,
})
TEXT = {
    "12233B": BLACK, "0E1B2E": BLACK, "5C6E85": MUTE, "00757F": NAVY, "0E8A7D": NAVY,
    "1B6BC0": NAVY, "C77400": BLACK, "9A5A00": BLACK, "2E7D32": BLACK, "C2185B": MAGENTA,
    "C62828": MAGENTA, "7A4FC0": VIOLET, "4527A0": VIOLET, "5FD0E8": CYAN, "00D6F1": CYAN,
    "AEBED0": "A6A6A6", "B8BEC6": "A6A6A6",
}
# Slides where red and rose sit side by side as different layers: keep them apart.
FILL_OVERRIDE = {4: {"C62828": ORANGE}, 11: {"C62828": ORANGE}}

THEME_COLORS = {
    "dk1": BLACK, "lt1": WHITE, "dk2": NAVY, "lt2": GREY, "accent1": CYAN, "accent2": YELLOW,
    "accent3": MAGENTA, "accent4": ORANGE, "accent5": VIOLET, "accent6": GREEN,
    "hlink": MAGENTA, "folHlink": PURPLE,
}

TXT, TXT_LIGHT = "Trenda IG Text", "Trenda IG Text Light"
MONO = "Courier New"

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def a(tag):
    return "{%s}%s" % (A, tag)


def p(tag):
    return "{%s}%s" % (P, tag)


def lum(hexv):
    def ch(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (int(hexv[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def contrast(x, y):
    hi, lo = sorted((lum(x), lum(y)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


# ======================================================================
# 1. RESTYLE
# ======================================================================
def xfrm_of(el):
    """The off/ext holder of a top-level shape, or None."""
    for path in ("{%s}spPr/{%s}xfrm" % (P, A), "{%s}grpSpPr/{%s}xfrm" % (P, A), "{%s}xfrm" % P):
        x = el.find(path)
        if x is not None:
            return x
    return None


def box_of(el):
    x = xfrm_of(el)
    if x is None:
        return None
    off, ext = x.find(a("off")), x.find(a("ext"))
    if off is None or ext is None:
        return None
    return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))


def top_shapes(slide):
    tree = slide.shapes._spTree
    return [c for c in tree if c.tag in (p("sp"), p("pic"), p("cxnSp"), p("grpSp"), p("graphicFrame"))]


def scale_slide(slide, k, origin_y):
    """Scale every shape, font, inset and rule by k, about the vertical centre
    line of the slide and the top of the content."""
    cx = SLIDE_W * EMU / 2
    for el in top_shapes(slide):
        x = xfrm_of(el)
        if x is None:
            continue
        off, ext = x.find(a("off")), x.find(a("ext"))
        ox, oy = int(off.get("x")), int(off.get("y"))
        off.set("x", str(int(round(cx + (ox - cx) * k))))
        off.set("y", str(int(round(origin_y + (oy - origin_y) * k))))
        ext.set("cx", str(int(round(int(ext.get("cx")) * k))))
        ext.set("cy", str(int(round(int(ext.get("cy")) * k))))
    root = slide._element
    for tag in ("rPr", "defRPr", "endParaRPr"):
        for r in root.iter(a(tag)):
            for att in ("sz", "spc"):
                if r.get(att) is not None:
                    r.set(att, str(int(round(int(r.get(att)) * k))))
    for el in root.iter(a("spcPts")):
        el.set("val", str(int(round(int(el.get("val")) * k))))
    for tag in ("ln", "lnL", "lnR", "lnT", "lnB"):
        for el in root.iter(a(tag)):
            if el.get("w") is not None:
                el.set("w", str(int(round(int(el.get("w")) * k))))
    for el in root.iter(a("pPr")):
        for att in ("marL", "indent"):
            if el.get(att) is not None:
                el.set(att, str(int(round(int(el.get(att)) * k))))
    for el in root.iter(a("bodyPr")):
        for att, d in (("lIns", 91440), ("rIns", 91440), ("tIns", 45720), ("bIns", 45720)):
            el.set(att, str(int(round(int(el.get(att, d)) * k))))
    for el in root.iter(a("tcPr")):
        for att, d in (("marL", 91440), ("marR", 91440), ("marT", 45720), ("marB", 45720)):
            el.set(att, str(int(round(int(el.get(att, d)) * k))))
    for el in root.iter(a("gridCol")):
        el.set("w", str(int(round(int(el.get("w")) * k))))
    for el in root.iter(a("tr")):
        el.set("h", str(int(round(int(el.get("h")) * k))))


def fit_above_footer(slide):
    boxes = [b for b in (box_of(e) for e in top_shapes(slide)) if b]
    if not boxes:
        return 1.0
    top = min(b[1] for b in boxes)
    bottom = max(b[1] + b[3] for b in boxes)
    floor = CONTENT_FLOOR * EMU
    if bottom <= floor:
        return 1.0
    k = (floor - top) / (bottom - top)
    scale_slide(slide, k, top)
    return k


def context(c):
    for anc in c.iterancestors():
        t = anc.tag
        if t in (a("rPr"), a("defRPr"), a("endParaRPr"), a("buClr")):
            return "text"
        if t in (a("ln"), a("lnL"), a("lnR"), a("lnT"), a("lnB")):
            return "line"
        if t == p("bgPr"):
            return "bg"
    return "fill"


def recolor(slide, num):
    fill = dict(FILL)
    fill.update(FILL_OVERRIDE.get(num, {}))
    line = dict(LINE)
    line.update(FILL_OVERRIDE.get(num, {}))
    for c in slide._element.iter(a("srgbClr")):
        v = c.get("val").upper()
        ctx = context(c)
        m = {"text": TEXT, "line": line, "fill": fill, "bg": {}}[ctx]
        if v in m:
            c.set("val", m[v])


def set_white_background(slide):
    cSld = slide._element.find(p("cSld"))
    bg = cSld.find(p("bg"))
    if bg is not None:
        cSld.remove(bg)
    bg = etree.SubElement(cSld, p("bg"))
    cSld.remove(bg)
    cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, p("bgPr"))
    sf = etree.SubElement(bgPr, a("solidFill"))
    etree.SubElement(sf, a("srgbClr"), val=WHITE)
    etree.SubElement(bgPr, a("effectLst"))


def retype(slide):
    """Headlines Trenda IG Text Bold; large display figures Trenda IG Text Light."""
    for el in top_shapes(slide):
        b = box_of(el)
        is_title_zone = b is not None and b[1] < 0.95 * EMU
        for r in el.iter(a("rPr"), a("endParaRPr"), a("defRPr")):
            faces = [r.find(a(t)) for t in ("latin", "ea", "cs")]
            faces = [f for f in faces if f is not None]
            if not faces:
                continue
            face = faces[0].get("typeface")
            if face == "Trenda IG Display Bold":
                new, bold = TXT, "1"
            elif face == "Trenda IG Display":
                sz = int(r.get("sz", "1800"))
                if is_title_zone and sz >= 1600:
                    new, bold = TXT, "1"
                else:
                    new, bold = TXT_LIGHT, "0"
            else:
                continue
            for f in faces:
                f.set("typeface", new)
            r.set("b", bold)


def fill_of(el):
    sppr = el.find(p("spPr"))
    if sppr is None:
        return None
    sf = sppr.find(a("solidFill"))
    if sf is None:
        return None
    c = sf.find(a("srgbClr"))
    return c.get("val") if c is not None else None


def fix_contrast(slide):
    """White text that now sits on a light IG colour (cyan, yellow, green,
    orange, grey) turns black. The ground is the shape's own fill, or else the
    topmost filled shape under the centre of the text."""
    shapes = top_shapes(slide)
    filled = []
    for i, el in enumerate(shapes):
        b, f = box_of(el), fill_of(el)
        if b and f:
            filled.append((i, b, f))
    for i, el in enumerate(shapes):
        tx = el.find(p("txBody"))
        if tx is None:
            continue
        ground = fill_of(el)
        if ground is None:
            b = box_of(el)
            if b is None:
                continue
            mx, my = b[0] + b[2] / 2, b[1] + b[3] / 2
            ground = WHITE
            for j, fb, ff in filled:
                if j < i and fb[0] <= mx <= fb[0] + fb[2] and fb[1] <= my <= fb[1] + fb[3]:
                    ground = ff
        for c in tx.iter(a("srgbClr")):
            if c.get("val") == WHITE and contrast(WHITE, ground) < 3.0:
                c.set("val", BLACK)


def ln_xml(tag, w=None, color=None, dash=None):
    el = etree.Element(a(tag))
    if color is None:
        etree.SubElement(el, a("noFill"))
        return el
    el.set("w", str(w))
    sf = etree.SubElement(el, a("solidFill"))
    etree.SubElement(sf, a("srgbClr"), val=color)
    if dash:
        etree.SubElement(el, a("prstDash"), val=dash)
    return el


def restyle_tables(slide):
    """IG dense chart style: bold header on white over a black rule, then rows
    alternating F2F2F2 and white on dotted black rules, no verticals."""
    for tbl in slide._element.iter(a("tbl")):
        for ri, tr in enumerate(tbl.findall(a("tr"))):
            for tc in tr.findall(a("tc")):
                tcPr = tc.find(a("tcPr"))
                if tcPr is None:
                    tcPr = etree.SubElement(tc, a("tcPr"))
                for child in list(tcPr):
                    if child.tag in (a("lnL"), a("lnR"), a("lnT"), a("lnB"), a("solidFill"), a("noFill")):
                        tcPr.remove(child)
                bottom = ln_xml("lnB", 15875, BLACK) if ri == 0 else ln_xml("lnB", 9525, BLACK, "sysDot")
                for el in (ln_xml("lnL"), ln_xml("lnR"), ln_xml("lnT"), bottom):
                    tcPr.append(el)
                sf = etree.SubElement(tcPr, a("solidFill"))
                etree.SubElement(sf, a("srgbClr"), val=GREY if ri % 2 == 1 else WHITE)
                if ri == 0:
                    for c in tc.iter(a("srgbClr")):
                        if context(c) == "text":
                            c.set("val", BLACK)


def update_theme(prs):
    master = prs.slide_masters[0]
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            root = etree.fromstring(part.blob)
            scheme = root.find(".//" + a("clrScheme"))
            scheme.set("name", "Insight Global")
            for name, val in THEME_COLORS.items():
                slot = scheme.find(a(name))
                for c in list(slot):
                    slot.remove(c)
                etree.SubElement(slot, a("srgbClr"), val=val)
            fonts = root.find(".//" + a("fontScheme"))
            fonts.set("name", "Insight Global")
            for which in ("majorFont", "minorFont"):
                fonts.find(a(which)).find(a("latin")).set("typeface", TXT)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def add_footer_to_layout(layout):
    """●●●IG mark bottom left and the page number bottom right, where the
    template puts them, on the content layout every content slide uses."""
    tree = layout.shapes._spTree
    for el in list(tree):
        nv = el.find(".//" + p("cNvPr"))
        if nv is not None and nv.get("name", "").startswith("IG Footer"):
            tree.remove(el)
    _, rid = layout.part.get_or_add_image_part(MARK_BLACK)
    w, h = 0.5425, 0.1133
    pic = etree.fromstring(
        '<p:pic xmlns:p="%s" xmlns:a="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:nvPicPr><p:cNvPr id="900" name="IG Footer Mark" descr="Insight Global"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr userDrawn="1"/></p:nvPicPr>'
        '<p:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
        % (P, A, rid, int(0.55 * EMU), int(FOOTER_TOP * EMU), int(w * EMU), int(h * EMU)))
    num = etree.fromstring(
        '<p:sp xmlns:p="%s" xmlns:a="%s"><p:nvSpPr><p:cNvPr id="901" name="IG Footer Page Number"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        '<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/><a:lstStyle/>'
        '<a:p><a:pPr algn="r"/><a:fld id="{5C1D3A57-8D5C-4E4B-9A43-7B1E2F0A9C11}" type="slidenum"><a:rPr lang="en-US" sz="800"><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
        '<a:latin typeface="Trenda IG Text"/></a:rPr><a:t>&#8249;#&#8250;</a:t></a:fld><a:endParaRPr lang="en-US" sz="800"/></a:p></p:txBody></p:sp>'
        % (P, A, int(12.28 * EMU), int(FOOTER_TOP * EMU), int(0.52 * EMU), int(h * EMU)))
    tree.append(pic)
    tree.append(num)


def restyle(prs):
    update_theme(prs)
    content_layout = None
    report = []
    for num, slide in enumerate(prs.slides, 1):
        recolor(slide, num)
        if num == 1:
            continue  # the IG title slide is already the template's own
        content_layout = slide.slide_layout
        set_white_background(slide)
        retype(slide)
        restyle_tables(slide)
        k = fit_above_footer(slide)
        fix_contrast(slide)
        report.append((num, k))
    add_footer_to_layout(content_layout)
    return content_layout, report


# ======================================================================
# 2. NEW SLIDES
# ======================================================================
def rgb(h):
    return RGBColor.from_string(h)


class Canvas:
    def __init__(self, slide):
        self.s = slide
        self.n = 0

    def _name(self, kind):
        self.n += 1
        return "%s %d" % (kind, self.n)

    def rect(self, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE,
             radius=None, dash=None):
        sp = self.s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sp.name = self._name("Shape")
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = radius
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = rgb(fill)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = rgb(line)
            sp.line.width = Pt(lw)
            if dash:
                ln = sp.line._get_or_add_ln()
                etree.SubElement(ln, a("prstDash"), val=dash)
        else:
            sp.line.fill.background()
        sp.shadow.inherit = False
        sp.text_frame.text = ""
        return sp

    def text(self, x, y, w, h, paras, size=7.5, color=BLACK, bold=False, font=TXT,
             align=None, anchor=MSO_ANCHOR.TOP, track=None, line=None, after=None):
        """paras: str | list of (str | list of (text, {overrides}))."""
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.name = self._name("Text")
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        if isinstance(paras, str):
            paras = [paras]
        for i, para in enumerate(paras):
            pg = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if align is not None:
                pg.alignment = align
            if line:
                pg.line_spacing = line
            if after is not None:
                pg.space_after = Pt(after)
            runs = para if isinstance(para, list) else [(para, {})]
            for chunk in runs:
                t, o = chunk if isinstance(chunk, tuple) else (chunk, {})
                r = pg.add_run()
                r.text = t
                f = r.font
                f.name = o.get("font", font)
                f.size = Pt(o.get("size", size))
                f.bold = o.get("bold", bold)
                f.italic = o.get("italic", False)
                f.color.rgb = rgb(o.get("color", color))
                tr = o.get("track", track)
                if tr:
                    r._r.get_or_add_rPr().set("spc", str(tr))
        return tb

    def label(self, x, y, w, t, color=BLACK, size=6.5, align=None):
        return self.text(x, y, w, 0.14, t.upper(), size=size, color=color, bold=True, track=150,
                         align=align, anchor=MSO_ANCHOR.MIDDLE)

    def pill(self, x, y, w, h, t, fill=WHITE, color=BLACK, line=BLACK, size=6, bold=True, track=100):
        self.rect(x, y, w, h, fill, line, 0.75, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5)
        return self.text(x, y, w, h, t, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER,
                         anchor=MSO_ANCHOR.MIDDLE, track=track)

    def frame(self, x, y, w, h, inner=WHITE, t=0.045):
        """The differentiator border: the three IG logo colours, yellow to
        magenta to cyan, around a card. Returns the inner card's box."""
        sp = self.rect(x, y, w, h, WHITE)
        sppr = sp._element.spPr
        for el in list(sppr):
            if el.tag in (a("solidFill"), a("noFill"), a("gradFill")):
                sppr.remove(el)
        grad = etree.fromstring(
            '<a:gradFill xmlns:a="%s" rotWithShape="1"><a:gsLst>'
            '<a:gs pos="0"><a:srgbClr val="%s"/></a:gs><a:gs pos="50000"><a:srgbClr val="%s"/></a:gs>'
            '<a:gs pos="100000"><a:srgbClr val="%s"/></a:gs></a:gsLst><a:lin ang="0" scaled="0"/></a:gradFill>'
            % (A, YELLOW, MAGENTA, CYAN))
        sppr.insert(list(sppr).index(sppr.find(a("prstGeom"))) + 1, grad)
        sp.name = self._name("Differentiator Border")
        self.rect(x + t, y + t, w - 2 * t, h - 2 * t, inner)
        return x + t, y + t, w - 2 * t, h - 2 * t

    def diff_tab(self, right, y, w=0.98):
        """Black DIFFERENTIATOR tab straddling the top edge, right-aligned."""
        self.rect(right - w - 0.12, y - 0.075, w, 0.15, BLACK)
        self.text(right - w - 0.12, y - 0.075, w, 0.15, "DIFFERENTIATOR", size=5.5, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, track=150)

    def tag(self, right, y, t, w=None):
        """Maturity tag, right-aligned at `right`."""
        w = w or (0.16 + 0.066 * len(t))
        self.rect(right - w, y, w, 0.16, BLACK, None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        self.text(right - w, y, w, 0.16, t, size=5.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, track=100)

    def key(self, right, y):
        """One-line key for the differentiator border, on the eyebrow line."""
        w, sw = 4.4, 0.34
        x = right - w
        self.frame(x, y + 0.01, sw, 0.15, WHITE, 0.035)
        self.text(x + sw + 0.08, y, w - sw - 0.08, 0.17,
                  [[("DIFFERENTIATOR", {"bold": True, "track": 150}),
                    ("   what an off-the-shelf coding agent does not do", {"color": MUTE})]],
                  size=6.5, anchor=MSO_ANCHOR.MIDDLE)

    def arrow(self, x1, y1, x2, y2, color=BLACK, w=1.0, dash=None, head=True):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.name = self._name("Arrow")
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(w)
        ln = c.line._get_or_add_ln()
        if dash:
            etree.SubElement(ln, a("prstDash"), val=dash)
        if head:
            etree.SubElement(ln, a("tailEnd"), type="triangle", w="med", len="med")
        return c

    def header(self, eyebrow, title, standfirst):
        self.text(0.45, 0.22, 9.0, 0.2, eyebrow, size=8, bold=True, track=300, anchor=MSO_ANCHOR.MIDDLE)
        self.text(0.45, 0.42, 12.43, 0.45, title, size=22, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        self.text(0.45, 0.88, 12.43, 0.26, standfirst, size=10.5, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)


def architecture_slide(prs, layout):
    s = prs.slides.add_slide(layout)
    cv = Canvas(s)
    cv.header("ASCENSION MAD APPROACH · TARGET ARCHITECTURE",
              "Architecture: agents grounded in one graph, every action governed",
              "Phase 1 builds the Engineering Context Knowledge Graph. Phase 2 runs the MAD agent team on it. "
              "Agents prepare, people approve, and no agent holds a production credential.")
    X0, X1 = 0.45, 12.88
    cv.key(X1, 0.215)

    # --- row 0: the people who own every decision
    y0 = 1.22
    lw0 = 1.5
    cv.rect(X0, y0, lw0, 0.38, BLACK)
    cv.text(X0 + 0.1, y0, lw0 - 0.15, 0.38, "PEOPLE OWN\nEVERY DECISION", size=6.5, bold=True, color=WHITE,
            track=150, anchor=MSO_ANCHOR.MIDDLE)
    owners = [("Pod leads and product owners", "Scope and priority · H1 requirements sign-off"),
              ("Architecture and security", "H2 design approval by risk tier · owns the Guardian rulebook"),
              ("Tech leads and domain approvers", "H3 code, access and security review"),
              ("Business and release owners", "H4 UAT · CAB · production go-live")]
    ow = (X1 - X0 - lw0 - 4 * 0.06) / 4
    for i, (t, d) in enumerate(owners):
        ox = X0 + lw0 + 0.06 + i * (ow + 0.06)
        cv.rect(ox, y0, ow, 0.38, YELLOW_T, YELLOW, 0.75)
        cv.text(ox + 0.1, y0 + 0.04, ow - 0.16, 0.15, t, size=7.5, bold=True)
        cv.text(ox + 0.1, y0 + 0.2, ow - 0.16, 0.14, d, size=6.5, color=MUTE)

    # --- row 1: sources -> phase 1 graph -> phase 2 agents
    y1, h1 = 1.72, 2.75
    sw = 1.75
    cv.rect(X0, y1, sw, h1, WHITE, RULE)
    cv.label(X0 + 0.1, y1 + 0.08, sw - 0.2, "Systems of record")
    srcs = [("Jira", "Stories, acceptance criteria"), ("Confluence", "Designs, business rules"),
            ("Salesforce", "Health, Data and Marketing Cloud"), ("Git", "Apex, LWC, metadata, tests"),
            ("Copado", "Promotions, test runs"), ("Google Drive", "Specs, mappings, evidence")]
    for i, (n, d) in enumerate(srcs):
        sy = y1 + 0.32 + i * 0.36
        cv.rect(X0 + 0.1, sy + 0.04, 0.07, 0.07, NAVY)
        cv.text(X0 + 0.24, sy, sw - 0.3, 0.15, n, size=7.5, bold=True)
        cv.text(X0 + 0.24, sy + 0.15, sw - 0.3, 0.14, d, size=6.5, color=MUTE)
    cv.pill(X0 + 0.1, y1 + h1 - 0.3, sw - 0.2, 0.2, "READ-ONLY · AUTHORITY STAYS", GREY, BLACK, GREY, 5.5, track=50)

    gx = X0 + sw + 0.45
    cv.arrow(X0 + sw + 0.03, y1 + 1.3, gx - 0.03, y1 + 1.3, NAVY, 1.25)
    cv.text(X0 + sw, y1 + 0.94, 0.45, 0.3, "SCHED.\nSYNC", size=5.5, bold=True, color=NAVY,
            align=PP_ALIGN.CENTER, track=100, anchor=MSO_ANCHOR.BOTTOM)

    # phase 1: the knowledge graph, framed as a differentiator
    gw = 3.55
    ix, iy, iw, ih = cv.frame(gx, y1, gw, h1)
    cv.rect(ix, iy, iw, 0.3, VIOLET)
    cv.text(ix + 0.1, iy, iw - 1.3, 0.3, "PHASE 1 · KNOWLEDGE GRAPH", size=6.5, bold=True, color=WHITE,
            track=100, anchor=MSO_ANCHOR.MIDDLE)
    cv.diff_tab(gx + gw, y1)
    cv.text(ix + 0.1, iy + 0.36, iw - 0.2, 0.2, "Engineering Context Knowledge Graph", size=8.5, bold=True,
            anchor=MSO_ANCHOR.MIDDLE)
    for i, t in enumerate(["Connect", "Validate", "Link", "Version"]):
        sx = ix + 0.1 + i * 0.7
        cv.rect(sx, iy + 0.62, 0.56, 0.21, VIOLET_T, VIOLET, 0.5)
        cv.text(sx, iy + 0.62, 0.56, 0.21, t, size=6.5, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 3:
            cv.arrow(sx + 0.57, iy + 0.725, sx + 0.69, iy + 0.725, VIOLET, 0.75)
    cv.pill(ix + 0.1, iy + 0.92, iw - 0.2, 0.18,
            "SANDBOX POC RUNNING · SALESFORCE · 2,177 NODES · 2,786 EDGES", WHITE, BLACK, BLACK, 5.5, track=50)
    facts = ["Typed facts: requirement → design → component → test → release",
             "Every fact carries its source path, SHA-256 and snapshot id",
             "Freshness and coverage checked; unknowns go to the owning team",
             "PHI and secrets redacted before anything is stored"]
    for i, t in enumerate(facts):
        fy = iy + 1.2 + i * 0.21
        cv.rect(ix + 0.12, fy + 0.05, 0.05, 0.05, VIOLET)
        cv.text(ix + 0.24, fy, iw - 0.32, 0.19, t, size=6.5)
    qy = iy + 2.06
    cv.rect(ix + 0.1, qy, iw - 0.2, ih - 2.06 - 0.08, GREY)
    cv.text(ix + 0.18, qy + 0.04, iw - 0.36, 0.16,
            [[("Ask before build. ", {"bold": True}),
              ("Exists, reuse or impact is answered from facts before any code.", {})]], size=6.5)
    cv.text(ix + 0.18, qy + 0.32, iw - 0.36, 0.14, "blast_radius · find_reuse · classify_risk · under 0.01 s",
            size=6, font=MONO)

    ax = gx + gw + 0.45
    cv.arrow(gx + gw + 0.03, y1 + 1.3, ax - 0.03, y1 + 1.3, NAVY, 1.25)
    cv.text(gx + gw, y1 + 0.94, 0.45, 0.3, "CITED\nFACTS", size=5.5, bold=True, color=NAVY,
            align=PP_ALIGN.CENTER, track=100, anchor=MSO_ANCHOR.BOTTOM)

    # phase 2: MAD
    aw = X1 - ax
    cv.rect(ax, y1, aw, h1, WHITE, NAVY, 1.0)
    cv.rect(ax, y1, aw, 0.3, NAVY)
    cv.text(ax + 0.1, y1, aw - 0.2, 0.3, "PHASE 2 · MULTI-AGENT DELIVERY (MAD)", size=6.5, bold=True,
            color=WHITE, track=100, anchor=MSO_ANCHOR.MIDDLE)
    bw = (aw - 0.4) / 3
    ctrl = [("Durable workflow controller", "GCP Workflows + Firestore: run state, checkpoints, retry budgets"),
            ("Codex orchestrator", "Plans the run, routes agents, structured outputs, pauses for people"),
            ("Approval service", "Authenticated decisions bound to run, scope, commit SHA and target")]
    for i, (t, d) in enumerate(ctrl):
        bx = ax + 0.1 + i * (bw + 0.1)
        cv.rect(bx, y1 + 0.4, bw, 0.55, BLACK)
        cv.text(bx + 0.08, y1 + 0.44, bw - 0.16, 0.16, t, size=7.5, bold=True, color=WHITE)
        cv.text(bx + 0.08, y1 + 0.61, bw - 0.16, 0.32, d, size=6.5, color="D9D9D9")
        if i < 2:
            cv.arrow(bx + bw + 0.01, y1 + 0.675, bx + bw + 0.09, y1 + 0.675, CYAN, 1.0)
    agents = [("Orchestrator", "classifies, routes, enforces gates"), ("Requirements", "Jira story → clear brief"),
              ("Architect", "impact, reuse, design from graph"), ("Developer", "builds in DEV + branch only"),
              ("Review", "quality, security, standards"), ("Test", "unit, regression, Zephyr → Copado"),
              ("Release", "Copado to lower environments")]
    cw = (aw - 0.2 - 3 * 0.06) / 4
    for i, (n, d) in enumerate(agents):
        r, c = divmod(i, 4)
        bx = ax + 0.1 + c * (cw + 0.06)
        by = y1 + 1.05 + r * 0.44
        cv.rect(bx, by, cw, 0.38, CYAN)
        cv.text(bx + 0.07, by + 0.04, cw - 0.12, 0.15, n + " Agent", size=7, bold=True)
        cv.text(bx + 0.07, by + 0.2, cw - 0.12, 0.15, d, size=6)
    bx = ax + 0.1 + 3 * (cw + 0.06)
    by = y1 + 1.49
    cv.rect(bx, by, cw, 0.38, WHITE, CYAN, 1.0)
    cv.text(bx + 0.07, by + 0.04, cw - 0.12, 0.15, "Pod skills", size=7, bold=True)
    cv.text(bx + 0.07, by + 0.2, cw - 0.12, 0.15, "Health · Marketing · QA", size=6)
    cv.text(ax + 0.1, y1 + 1.93, aw - 0.2, 0.14,
            "6 of 7 agents are built from the UI, backend and QA agents in use today. Release Agent, Change "
            "Guardian, gateway and control plane are pilot builds.", size=6, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)
    gy = y1 + 2.13
    gh = h1 - 2.13 - 0.08
    gwid = (aw - 0.3) * 0.6
    fx, fy, fw, fh = cv.frame(ax + 0.1, gy, gwid, gh, MAGENTA)
    cv.diff_tab(ax + 0.1 + gwid, gy)
    cv.text(fx + 0.08, fy + 0.05, fw - 1.25, 0.15, "Change Guardian · across every pod", size=7.5, bold=True,
            color=WHITE)
    cv.text(fx + 0.08, fy + 0.22, fw - 0.16, 0.26,
            "Diff + ownership graph + rulebook, at design, commit and promotion: "
            "PASS · NEEDS OWNER APPROVAL · BLOCK", size=6.5, color=WHITE)
    ex = ax + 0.2 + gwid
    ew = aw - 0.3 - gwid
    cv.rect(ex, gy, ew, gh, WHITE, NAVY, 1.0)
    cv.text(ex + 0.08, gy + 0.05, ew - 0.16, 0.15, "Approved model endpoint", size=7.5, bold=True)
    cv.text(ex + 0.08, gy + 0.22, ew - 0.16, 0.3,
            "Allowlisted egress · minimal context · PHI and secrets never sent · model swappable behind MCP",
            size=6.5)

    # --- row 2: tool + policy gateway
    y2 = y1 + h1 + 0.14
    cv.arrow(ax + aw / 2, y1 + h1 + 0.01, ax + aw / 2, y2 - 0.01, BLACK, 1.25)
    cv.text(ax + aw / 2 + 0.06, y1 + h1, 1.8, 0.14, "ONLY APPROVED ACTIONS", size=5.5, bold=True, track=100,
            anchor=MSO_ANCHOR.MIDDLE)
    hg = 0.55
    cv.rect(X0, y2, X1 - X0, hg, BLACK)
    cv.text(X0 + 0.12, y2 + 0.06, 3.6, 0.16, "TOOL + POLICY GATEWAY (MCP)", size=7, bold=True, color=CYAN, track=150)
    cv.text(X0 + 0.12, y2 + 0.24, 3.75, 0.3,
            "Checks identity, target, approval and artifact SHA → runs only an allowed API / CLI action → "
            "records the outcome", size=6.5, color=WHITE)
    ids = [("AI Read Service", "read-only into the graph"), ("AI Development Service", "DEV org + source branch"),
           ("AI Release Service", "Copado, lower environments")]
    iw_ = 1.75
    for i, (t, d) in enumerate(ids):
        ixx = X0 + 4.05 + i * (iw_ + 0.1)
        cv.rect(ixx, y2 + 0.08, iw_, 0.39, None, "8C8C8C", 0.75)
        cv.text(ixx + 0.08, y2 + 0.11, iw_ - 0.12, 0.16, t, size=7, bold=True, color=WHITE)
        cv.text(ixx + 0.08, y2 + 0.27, iw_ - 0.12, 0.16, d, size=6.5, color="D9D9D9")
    nx = X0 + 4.05 + 3 * (iw_ + 0.1) + 0.05
    cv.text(nx, y2 + 0.06, X1 - nx - 0.12, hg - 0.12,
            [[("No production credentials for any agent. ", {"bold": True, "color": YELLOW}),
              ("Owning teams run every production release.", {"color": WHITE})]], size=6.5,
            anchor=MSO_ANCHOR.MIDDLE)

    # --- row 3: delivery paths -> human release gate
    y3 = y2 + hg + 0.14
    paths = [("Salesforce Health Cloud", "Git + Copado · DEV → SIT → UAT"),
             ("Marketing + Data Cloud", "Supported APIs, Data Kits · Test BU"),
             ("GCP + integrations", "Owner-managed CI/CD · DEV → TEST")]
    pw = 2.6
    for i, (t, d) in enumerate(paths):
        px = X0 + i * (pw + 0.12)
        cv.arrow(px + pw / 2, y2 + hg + 0.01, px + pw / 2, y3 - 0.01, BLACK, 1.0)
        cv.rect(px, y3, pw, 0.5, WHITE, RULE)
        cv.text(px + 0.1, y3 + 0.06, pw - 0.2, 0.16, t, size=7.5, bold=True)
        cv.text(px + 0.1, y3 + 0.25, pw - 0.2, 0.2, d, size=6.5, color=MUTE)
    hx = X0 + 3 * (pw + 0.12) + 0.3
    cv.arrow(hx - 0.3, y3 + 0.25, hx - 0.02, y3 + 0.25, BLACK, 1.25)
    cv.rect(hx, y3, X1 - hx, 0.5, YELLOW)
    cv.text(hx + 0.1, y3 + 0.06, 2.0, 0.16, "Human release gate", size=7.5, bold=True)
    cv.text(hx + 0.1, y3 + 0.25, X1 - hx - 0.2, 0.2, "UAT and CAB approve · the owning team deploys to PROD",
            size=6.5)
    cv.tag(X1 - 0.08, y3 + 0.06, "UNCHANGED PROCESS")

    # --- row 4: change passport, learning loop (both differentiators)
    y4 = y3 + 0.62
    h4 = 0.6
    pw4 = 6.15
    fx, fy, fw, fh = cv.frame(X0, y4, pw4, h4)
    cv.diff_tab(X0 + pw4, y4)
    cv.text(fx + 0.1, fy + 0.06, 3.0, 0.16, "CHANGE PASSPORT + AUDIT", size=7, bold=True, track=150)
    cv.text(fx + 0.1, fy + 0.24, fw - 0.2, 0.15,
            "Story → source snapshot → design → approvals → commit SHA → test evidence → deployment → outcome",
            size=6.5)
    cv.text(fx + 0.1, fy + 0.39, fw - 0.2, 0.13,
            "BigQuery audit store · Cloud Logging · CAB evidence is assembled, not collated", size=6.5, color=MUTE)
    rx = X0 + pw4 + 0.13
    rw = X1 - rx
    fx, fy, fw, fh = cv.frame(rx, y4, rw, h4, GREEN_T)
    cv.diff_tab(X1, y4)
    cv.rect(fx + 0.1, fy + 0.1, 0.3, 0.3, GREEN, None, shape=MSO_SHAPE.CIRCULAR_ARROW)
    cv.text(fx + 0.5, fy + 0.06, 3.6, 0.16, "LEARNING LOOP · NO MODEL TRAINED ON OUR DATA", size=7, bold=True,
            track=100)
    cv.text(fx + 0.5, fy + 0.24, fw - 0.6, 0.28,
            "Approved evidence refreshes the graph · each incident becomes a Guardian rule · each rejected output "
            "becomes an evaluation case · a new model must pass the same evaluations", size=6.5)

    # --- row 5: patterns this is built on
    y5 = y4 + h4 + 0.09
    h5 = CONTENT_FLOOR - 0.02 - y5
    cv.rect(X0, y5, X1 - X0, h5, GREY)
    cv.text(X0 + 0.12, y5, X1 - X0 - 0.24, h5,
            [[("BUILT ON PATTERNS ALREADY PROVEN AT SCALE   ", {"bold": True, "track": 100}),
              ("Agent works on its own branch, people approve the PR ", {}),
              ("(GitHub Copilot coding agent)", {"color": MUTE}), ("  ·  Graph-grounded retrieval ", {}),
              ("(Microsoft GraphRAG)", {"color": MUTE}), ("  ·  Open tool protocol ", {}),
              ("(MCP, supported by OpenAI Codex)", {"color": MUTE}), ("  ·  Durable execution ", {}),
              ("(Temporal, Google Cloud Workflows)", {"color": MUTE}), ("  ·  Policy as code, build provenance ", {}),
              ("(OPA, SLSA)", {"color": MUTE})]], size=6.5, anchor=MSO_ANCHOR.MIDDLE)
    set_notes(s, ARCH_NOTES)
    return s


def execution_slide(prs, layout):
    s = prs.slides.add_slide(layout)
    cv = Canvas(s)
    cv.header("ASCENSION MAD APPROACH · EXECUTION FLOW",
              "Execution flow: one story, nine steps, nothing skips a gate",
              "Read left to right, then follow the arrows. Every handoff produces a reviewable artifact bound to a "
              "commit SHA; named people approve at H1 to H4 and CAB.")
    X0, X1 = 0.45, 12.88
    SX1 = 9.15
    CX = 9.4
    cv.key(X1, 0.215)

    # run controller bar
    y0 = 1.22
    cv.rect(X0, y0, X1 - X0, 0.3, BLACK)
    cv.text(X0 + 0.12, y0, 2.4, 0.3, "DURABLE RUN · ONE PER STORY", size=7, bold=True, color=CYAN, track=150,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(X0 + 2.55, y0, X1 - X0 - 2.7, 0.3,
            "Checkpoint after every step  ·  retries within a budget  ·  resumes where it stopped  ·  "
            "cannot skip a step  ·  one Jira story, one auditable run", size=7, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

    steps = [
        ("01", "Understand the request", "REQUIREMENTS AGENT",
         "An approved Jira story starts the run. Drafts scope, acceptance criteria and open questions.",
         "agreed brief linked to source facts", ("H1 APPROVE", "Product owner / business lead")),
        ("02", "Design with evidence", "ARCHITECT AGENT",
         "Asks the graph: exists, reuse or impact? Sets the risk tier and test scope; names affected pods.",
         "design + change plan with cited facts", ("H2 APPROVE", "Architect + affected pod owners")),
        ("03", "Build in an isolated workspace", "DEVELOPER AGENT + TEST AGENT",
         "DEV org and its own branch, scoped tools only. Test cases written from the design in parallel.",
         "pull request + versioned candidate (SHA)", None),
        ("04", "Verify independently", "REVIEW AGENT + TEST AGENT · PARALLEL",
         "Code, access, security and standards; unit, contract and regression tests, bound to the SHA.",
         "findings + passing evidence on the SHA", ("H3 APPROVE", "Tech lead + security / domain")),
        ("05", "Promote and prove in UAT", "RELEASE AGENT + TEST AGENT",
         "Copado promotes the approved SHA to SIT and UAT; integration and cross-pod tests run.",
         "UAT build + test evidence + defects", None),
        ("06", "Accept the business outcome", "HUMAN BUSINESS OWNERS",
         "Acceptance criteria and workflows checked; clinical / legal review where needed. Defects return to this run.",
         "UAT sign-off", ("H4 APPROVE", "Business UAT + domain owners")),
        ("07", "Prepare a release that can recover", "RELEASE AGENT",
         "Packages the approved version, runbook and rollback plan; re-checks source drift and dependencies.",
         "release bundle + CAB evidence", None),
        ("08", "Authorize production release", "CAB + PLATFORM RELEASE OWNER",
         "Approve window, risk and recovery. Owner-controlled CI/CD deploys; campaign activation stays human.",
         "authorized production release", ("CAB APPROVE", "Release owner deploys")),
        ("09", "Observe, learn and close", "SYSTEM OWNERS + RUN CONTROLLER",
         "Smoke checks and business validation; Jira closed; graph, rulebook and evaluations updated.",
         "verified outcome + auditable history", None),
    ]
    tags = {"03": "AI-LED · GUARDIAN CHECKS EVERY COMMIT", "05": "AI-LED · NO PASSPORT, NO PROMOTION",
            "07": "AI-LED · ROLLBACK PLAN REQUIRED", "09": "CLOSES THE LEARNING LOOP"}
    gap = 0.36
    cw = (SX1 - X0 - 2 * gap) / 3
    ch, rg = 1.3, 0.15
    ys = [1.66 + r * (ch + rg) for r in range(3)]
    order = [(0, 0), (0, 1), (0, 2), (1, 2), (1, 1), (1, 0), (2, 0), (2, 1), (2, 2)]
    pos = {}
    for (num, title, agent, body, out, gate), (r, c) in zip(steps, order):
        x, y = X0 + c * (cw + gap), ys[r]
        pos[num] = (x, y)
        human = agent.startswith("HUMAN") or agent.startswith("CAB")
        if num == "02":
            x_, y_, w_, h_ = cv.frame(x, y, cw, ch)
            cv.diff_tab(x + cw, y)
        else:
            cv.rect(x, y, cw, ch, WHITE, RULE)
        cv.rect(x + 0.09, y + 0.1, 0.25, 0.25, YELLOW if human else BLACK, None, shape=MSO_SHAPE.OVAL)
        cv.text(x + 0.09, y + 0.1, 0.25, 0.25, num, size=7, bold=True, color=BLACK if human else WHITE,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.41, y + 0.09, cw - 0.5, 0.27, title, size=9, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.12, y + 0.42, cw - 0.22, 0.13, agent, size=6, bold=True, color=BLACK if human else NAVY,
                track=100, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.12, y + 0.57, cw - 0.22, 0.36, body, size=7)
        cv.rect(x + 0.1, y + 0.94, cw - 0.2, 0.15, CYAN_T)
        cv.text(x + 0.16, y + 0.94, cw - 0.3, 0.15, [[("Output: ", {"bold": True}), (out, {})]], size=6.5,
                anchor=MSO_ANCHOR.MIDDLE)
        if gate:
            g, who = gate
            cv.rect(x + 0.1, y + 1.12, cw - 0.2, 0.14, YELLOW)
            cv.text(x + 0.16, y + 1.12, cw - 0.3, 0.14, [[(g + "  ", {"bold": True, "track": 100}), (who, {})]],
                    size=6, anchor=MSO_ANCHOR.MIDDLE)
        else:
            cv.rect(x + 0.1, y + 1.12, cw - 0.2, 0.14, WHITE, CYAN, 0.75)
            cv.text(x + 0.16, y + 1.12, cw - 0.3, 0.14, tags[num], size=6, bold=True, track=100,
                    anchor=MSO_ANCHOR.MIDDLE)

    def right(a_, b_):
        (x1, y1), (x2, _) = pos[a_], pos[b_]
        cv.arrow(x1 + cw + 0.04, y1 + ch / 2, x2 - 0.04, y1 + ch / 2, NAVY, 1.5)

    def left(a_, b_):
        (x1, y1), (x2, _) = pos[a_], pos[b_]
        cv.arrow(x1 - 0.04, y1 + ch / 2, x2 + cw + 0.04, y1 + ch / 2, NAVY, 1.5)

    def down(a_, b_, dx):
        (x1, y1), (_, y2) = pos[a_], pos[b_]
        cv.arrow(x1 + dx, y1 + ch + 0.01, x1 + dx, y2 - 0.01, NAVY, 1.5)

    right("01", "02"); right("02", "03"); down("03", "04", 0.6)
    left("04", "05"); left("05", "06"); down("06", "07", cw / 2)
    right("07", "08"); right("08", "09")
    fx = pos["03"][0] + cw - 0.32
    cv.arrow(fx, pos["04"][1] - 0.01, fx, pos["03"][1] + ch + 0.01, MAGENTA, 1.25, dash="dash")
    cv.text(fx - 1.75, pos["03"][1] + ch + 0.01, 1.68, 0.13, "FIX LOOP · THEN OWNER", size=5.5, bold=True,
            color=MAGENTA, track=100, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # controls across every step, stacked on the right
    cv.label(CX, 1.6, 3.0, "Controls across every step")
    cy0 = 1.8
    kh = (ys[2] + ch - cy0 - 3 * 0.1) / 4
    kw = X1 - CX
    controls = [
        ("Change Guardian", MAGENTA, WHITE, True, [
            "At design, every commit and promotion: diff + ownership graph + rulebook",
            "PASS · NEEDS OWNER APPROVAL · BLOCK, enforced by CI/CD",
            "Nightly drift watch flags edits made outside the pipeline"]),
        ("Approval service", YELLOW, BLACK, True, [
            "Each decision bound to run, scope, commit SHA and target",
            "Changed scope or artifact → the approval no longer matches",
            "Expired or rejected → the run stops or returns to the owner"]),
        ("Failure and recovery", NAVY, WHITE, False, [
            "Transient error → bounded retry; every step is idempotent",
            "Missing context or failed test → rework, or back to the owner",
            "Timeout or budget breach → pause and escalate to a person"]),
        ("Accuracy guard", CYAN, BLACK, False, [
            "Context only from the graph; every claim cites a node and SHA",
            "Uncited claim or invented dependency → rejected",
            "The agent that builds never approves its own work"]),
    ]
    for i, (t, col, ink, diff, lines) in enumerate(controls):
        ky = cy0 + i * (kh + 0.1)
        if diff:
            kx, kyy, kww, khh = cv.frame(CX, ky, kw, kh)
            cv.diff_tab(X1, ky)
        else:
            kx, kyy, kww, khh = CX, ky, kw, kh
            cv.rect(CX, ky, kw, kh, WHITE, col if col != YELLOW else "E6C200", 1.0)
        cv.rect(kx, kyy, kww, 0.22, col)
        cv.text(kx + 0.1, kyy, kww - 1.3, 0.22, t.upper(), size=6.5, bold=True, color=ink, track=150,
                anchor=MSO_ANCHOR.MIDDLE)
        for j, ln in enumerate(lines):
            ly = kyy + 0.29 + j * 0.2
            cv.rect(kx + 0.1, ly + 0.05, 0.05, 0.05, BLACK)
            cv.text(kx + 0.2, ly, kww - 0.28, 0.18, ln, size=6.5)

    # change passport: the evidence each step writes
    py = ys[2] + ch + 0.14
    ph = CONTENT_FLOOR - 0.02 - py
    ix, iy, iw, ih = cv.frame(X0, py, X1 - X0, ph)
    cv.diff_tab(X1, py)
    cv.text(ix + 0.12, iy + 0.06, 2.0, 0.16, "CHANGE PASSPORT", size=7.5, bold=True, track=150)
    cv.text(ix + 0.12, iy + 0.25, 1.95, ih - 0.3,
            "Written by the run as it moves, kept in BigQuery. CAB evidence without manual collation.",
            size=6.5, color=MUTE)
    items = [("Story", "01"), ("Source snapshot", "02"), ("Design", "02"), ("Approvals", "H1–H4, CAB"),
             ("Commit SHA", "03"), ("Test evidence", "04 · 05"), ("Deployment", "08"), ("Outcome", "09")]
    ga = 0.18
    lx = ix + 2.2
    iw8 = (ix + iw - 0.1 - lx - 7 * ga) / 8
    for i, (n, st) in enumerate(items):
        bx = lx + i * (iw8 + ga)
        cv.rect(bx, iy + 0.1, iw8, ih - 0.2, GREY)
        cv.text(bx + 0.08, iy + ih / 2 - 0.17, iw8 - 0.12, 0.16, n, size=7.5, bold=True)
        cv.text(bx + 0.08, iy + ih / 2 + 0.02, iw8 - 0.12, 0.14, "step " + st if st[0].isdigit() else st, size=6,
                color=MUTE)
        if i < 7:
            cv.arrow(bx + iw8 + 0.02, iy + ih / 2, bx + iw8 + ga - 0.02, iy + ih / 2, BLACK, 1.0)
    set_notes(s, EXEC_NOTES)
    return s


ARCH_NOTES = """SLIDE: Architecture, how MAD is built   (about 4 minutes)

KEY MESSAGE
Two phases, one picture. Phase 1 is the Engineering Context Knowledge Graph: what the agents know. Phase 2 is the MAD agent team: who does the work. Every action they take goes through one gateway, and people own every decision. Read it top to bottom: people, knowledge, agents, gateway, delivery, evidence.

WALK IT
- Top strip: the four groups of people who own the decisions. This is the answer to "who is accountable". Nothing on this slide replaces them.
- Left: the six systems of record. Authority never moves; the graph reads them, read-only.
- Middle (violet): Phase 1. Connect, validate, link, version. Every fact has a source path, a SHA-256 and a snapshot id. PHI and secrets are redacted before anything is stored. The query API answers impact, reuse and risk deterministically, in under 0.01 seconds, without a model.
- Right (navy): Phase 2. The durable controller keeps the state of every run; the Codex orchestrator plans and routes. Seven common agents, six already in use. Pod skills are what change per pod. The Change Guardian sits across every pod. The model is behind an approved, allowlisted endpoint.
- Black bar: the tool and policy gateway. Nothing reaches a system except through it, and it checks identity, target, approval and the artifact SHA first. Three identities: read, development (DEV + branch), release (lower environments). No production credentials for any agent.
- Delivery paths are unchanged: Copado for Health Cloud, supported APIs for Marketing Cloud, owner CI/CD for GCP. UAT and CAB stay human; the owning team deploys to production.
- Bottom: the change passport (the audit trail, assembled automatically) and the learning loop.

LIKELY QUESTIONS, AND WHERE TO POINT ON THE SLIDE (presenter reference, not on the slide)
1. How do the agents work together?  Point at Phase 2. The orchestrator routes work; the durable controller holds state, checkpoints and approvals. Agents never call each other freely: each one hands over an artifact (brief, design, pull request, test report) that a person or the next agent can review.
2. How do you connect to our systems?  Point at the READ ONLY arrow and the gateway. Reads come through read-only connectors (OAuth, JWT bearer, deploy keys, service accounts). Writes go only through the MCP gateway, with one of three scoped identities. MCP is the open tool standard Codex supports, so a new system is a new adapter, not a new build.
3. How secure is the AI?  Point at the model endpoint and the gateway. Three separated identities, least privilege. No production credentials for any agent. The model sees a minimal, redacted context slice; PHI and secrets are never sent. The endpoint is allowlisted; everything runs inside the existing GCP perimeter (VPC Service Controls, CMEK, Secret Manager), with every action logged to BigQuery.
4. What if the AI is wrong or hallucinates?  Point at the Ask before build box in the knowledge graph. Agents reason over cited facts from the graph, not over guesses. Every claim must cite a node id or source reference; a claim without one, or a dependency the graph does not hold, is rejected. Then tests prove behaviour and people approve. The graph answers exist / reuse / impact itself, deterministically, so the most expensive mistakes (duplicates, missed dependencies) never reach the model.
5. How does it keep learning? Does the model get trained?  Point at the learning loop. We do not train or fine-tune any model on Ascension data. The system improves in three governed ways: approved delivery evidence refreshes the graph after every release; every incident or near-miss becomes a Guardian rule; every rejected or reworked output becomes an evaluation case that agent skills (and any new model version) must pass before they are rolled out. It gets smarter every release without patient data ever going into a model.
6. Who controls the SDLC?  Point at the people strip and the controller. Named people own H1 to H4 and CAB. The controller enforces the order and cannot skip a step. Copado, environments and CAB are unchanged.
7. Why are you confident this is achievable?  Point at the maturity tags and the grey strip at the bottom. Six of the seven agents are built from the UI, backend and QA agents already in use. The knowledge graph POC is running in the Ascension sandbox (last build 2,177 nodes and 2,786 edges) and already refuses to duplicate a field that exists. Every other box is a managed GCP service or a tool Ascension already runs. And each pattern is proven elsewhere (bottom strip): what is new is combining them around a graph of the Ascension org.

THE DIFFERENTIATORS (the boxes with the yellow-magenta-cyan border and the DIFFERENTIATOR tab)
- The knowledge graph: agents know the whole Ascension delivery estate (Jira, Confluence, Salesforce, Git, Copado, Drive) with provenance, not just one repository; Ask before build stops duplicates before any code.
- The Change Guardian: cross-pod impact is checked at design, commit and promotion. Off-the-shelf agents judge a change only inside its own repository.
- The change passport: the evidence chain from story to outcome is written by the run itself, so CAB evidence is assembled, not collated.
- The learning loop: it improves every release without training a model on Ascension data.
Everything else on the slide (orchestration, gateways, CI/CD, human gates) is sound engineering that others also do; say so, it builds trust.

OTHER QUESTIONS ARCHITECTS ASK
- Are we locked into OpenAI?  No. The graph is vendor-neutral JSON, tools are behind MCP, and a model change is a configuration change that must pass the same evaluation set.
- What does it cost to run?  The graph answers most questions with no model call at all; the model is called on a bounded slice only when something genuinely new must be designed or built. Cost scales with change volume, not org size.
- What about prompt injection from a Jira ticket or a document?  Source text is data, never instructions. Agents can act only through the gateway, which allows a fixed list of actions per identity and needs an approval bound to the SHA. A poisoned ticket cannot grant itself a permission.
- What if the graph is stale or wrong?  Each sync is versioned with a snapshot id, freshness and coverage are checked, and every answer records which snapshot it came from. Unknown or stale context goes to the owning team instead of being guessed.
- Two pods changing the same object at once?  The controller runs one run per story with concurrency limits; the Guardian sees both changes against the same ownership graph and routes them to the affected owners.
- How is this different from Copilot, Agentforce for Developers or Devin?  Those agents know the repository. None of them knows the Ascension org across Jira, Confluence, Salesforce metadata, Copado and Drive, and none of them governs cross-pod impact. The graph and the Guardian are the difference.

WHAT IS NOT PROVEN YET (say it before they ask)
Reading Flow internals needs the Tooling API credential in the sandbox; that setup is in progress. The Release Agent and the Change Guardian are pilot builds. The pilot exists to prove approvals, recovery and the audit trail on real changes.

INDUSTRY PATTERNS ON THE BOTTOM STRIP
GitHub Copilot coding agent works on its own branch and needs a person to approve its pull request. Microsoft GraphRAG showed that graph-grounded retrieval answers connected questions better than plain search. MCP is the open tool protocol, supported by OpenAI Codex. Temporal and Google Cloud Workflows are durable execution engines. OPA is policy as code; SLSA is build provenance."""

EXEC_NOTES = """SLIDE: Execution flow, one story from request to production   (about 4 minutes)

KEY MESSAGE
This is how the SDLC is controlled. One story is one durable run. It moves through nine steps in a fixed order; agents do the work inside each step, named people approve at H1 to H4 and CAB, and every handoff is an artifact bound to a commit SHA. Nothing skips a gate, and nothing that was approved can silently change.

WALK IT
Row 1, left to right. 01 the Requirements Agent turns an approved Jira story into a brief; the product owner signs (H1). 02 the Architect Agent asks the graph whether this already exists, what can be reused and what it touches, sets the risk tier and names the affected pods; the architect and those pod owners sign (H2). 03 the Developer Agent builds in DEV on its own branch while the Test Agent writes cases from the design.
Row 2, right to left. 04 Review and Test agents verify independently and in parallel; findings go back to the developer on the fix loop, and the tech lead and security sign (H3). 05 Copado promotes that exact SHA to SIT and UAT. 06 business owners accept the outcome (H4); defects return to the same run.
Row 3. 07 the Release Agent prepares a release that can recover: runbook, rollback plan, drift re-check. 08 CAB and the platform release owner authorize, and the owner's own pipeline deploys; no agent touches production. 09 smoke checks and business validation, Jira closed, and the learning loop runs.
Right: four controls that sit across every step. Bottom: the change passport, showing which step writes each piece of evidence.

LIKELY QUESTIONS, AND WHERE TO POINT ON THE SLIDE (presenter reference, not on the slide)
1. How is the SDLC controlled?  Point at the gates. The run cannot advance without the previous step's artifact and, at H1 to H4 and CAB, a named person's decision. The order is enforced by the controller, the verdicts by CI/CD.
2. What happens if a step breaks?  Point at the black run bar and Failure and recovery. Transient errors retry within a budget and every step is idempotent. Missing context or a failing test goes back to rework or to the owner. A timeout or budget breach pauses the run and escalates. Because there is a checkpoint after every step, the run resumes where it stopped and approved work is never redone.
3. What if the output is wrong?  Point at step 04 and the Accuracy guard. The agent that builds never approves its own work. Review and Test are separate agents, the Guardian is separate again, and a person signs H3. Every claim cites a graph node and SHA; an uncited claim or an invented dependency is rejected. Failures go round the fix loop a bounded number of times, then to a person.
4. What if something changes after it was approved?  Point at the Approval service (bordered as a differentiator). Approvals are bound to the run, the scope, the commit SHA and the target. Change any of them and the approval no longer matches, so a new approval is required. Approvals also expire.
5. Can one pod's change break another pod?  Point at the Change Guardian. It checks at design, at every commit and at promotion, against the ownership graph and the rulebook. If another pod is affected, its lead must sign; if a rule is broken, it blocks and the developer agent drafts the fix. A nightly drift watch catches edits made outside the pipeline.
6. What if the production release fails?  Point at step 07 and 08. A release cannot reach CAB without a rollback plan. The owning team's pipeline deploys, and the release owner runs the recovery plan; agents hold no production credentials.
7. Does this really solve our problem?  Point at step 02. The root problem is that nobody can see what already exists before building. Step 02 answers that from the graph before any code: the sandbox POC already refuses to duplicate a field that exists. Then the Guardian stops the second problem: one pod's change breaking another's.

THE DIFFERENTIATORS (the boxes with the yellow-magenta-cyan border)
Step 02 design from graph evidence, the Change Guardian, approvals bound to one SHA, and the change passport. Then say why:
Approval cannot drift: every gate signs one specific SHA. No agent marks its own homework. A broken step resumes from its checkpoint. CAB evidence builds itself as the run moves (the change passport). And every release teaches the system, without training a model on patient data.

OTHER QUESTIONS
- How long does a run take?  Measured with agent-assisted delivery: Health Cloud changes went from 1-2 weeks to 2-4 days. The human gates set the pace, deliberately.
- Can people override an agent?  At every gate. A person can reject, send back with comments or stop the run; the decision is recorded in the passport.
- What does the audit trail hold?  Request, user, timestamp, source snapshot, recommendation, approvals, files changed, diff, tests, environment, deployment and outcome, queryable in BigQuery."""


def set_notes(slide, text):
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        # The deck's notes master has no body placeholder, so add the same
        # "Notes Placeholder" its own notes slides carry.
        ns.shapes._spTree.append(etree.fromstring(
            '<p:sp xmlns:p="%s" xmlns:a="%s"><p:nvSpPr><p:cNvPr id="3" name="Notes Placeholder 2"/>'
            '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="body" idx="1"/></p:nvPr>'
            '</p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>' % (P, A)))
    ns.notes_text_frame.text = text


def move_new_slides(prs, count, after):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    new = ids[-count:]
    for el in new:
        lst.remove(el)
    for i, el in enumerate(new):
        lst.insert(after + i, el)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else SRC
    out = sys.argv[2] if len(sys.argv) > 2 else OUT
    prs = Presentation(src)
    layout, report = restyle(prs)
    for num, k in report:
        if k < 1:
            print("slide %d: scaled to %.3f to clear the footer" % (num, k))
    architecture_slide(prs, layout)
    execution_slide(prs, layout)
    move_new_slides(prs, 2, NEW_AFTER)
    prs.save(out)
    print("wrote", os.path.abspath(out), "-", len(prs.slides), "slides")


if __name__ == "__main__":
    main()
