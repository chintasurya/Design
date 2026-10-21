"""One-slide preview: the Sierra voice agent on governed GCP knowledge, with patient actions through scoped APIs.

Black, muted style (greys, cyan accent, yellow for human gates, the IG gradient only on the hook).
The last slide of the Agentic Harness deck (build_data_harness_deck.py); this script also renders it alone.
Logo: put the official Sierra logo at deck/assets/logos/sierra-white.png (or sierra.png); until then the slide
shows the word "Sierra".

    python3 sierra_voice.py [out.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

import muted_theme as mt
from build_data_harness_deck import LOGOS, R, SRC, card, logo, new_slide, path, pill
from build_mad_ig_deck import CYAN, YELLOW, set_notes
from footer import fix_footers
from story_slides import dark_header

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "build", "sierra_voice.pptx")
CARD, CARD2, EDGE = "141414", "1E1E1E", "333333"
T1, T2, T3 = "F2F2F2", "B3B3B3", "8C8C8C"
GATE = "2B2508"

NOTES = """SLIDE: Sierra voice agent: accurate answers, safe patient actions   (about 2 minutes)

KEY MESSAGE
Sierra answers questions from governed knowledge base articles in Google Cloud, and books, cancels or reschedules only through scoped APIs to the scheduling system of record, after the caller is verified. Patient records never go into the knowledge base.

WALK IT
- Left: the caller talks to the Sierra voice agent.
- Top, knowledge path (no PHI): articles live in GCP with an owner, version, review date and approval; DLP and Collibra scan them for PHI; only approved articles are indexed and synced to Sierra when they change.
- Bottom, patient path: verify the caller first; then a scoped scheduling API returns only that patient's data and enforces the rules (cancellation window, reschedule rules, provider restrictions); bookings are written to the scheduling system of record, never a GCP copy.
- Right, after every call: hand-off to a person when needed, masked transcripts with retention limits, outcomes as aggregates into Gold A, and test calls before every release.
- Bottom: the six safeguards."""


def sierra_logo():
    """Path to the official Sierra logo (white or light version for a dark slide), if it has been added."""
    for name in ("sierra-white.png", "sierra.png"):
        f = os.path.join(LOGOS, name)
        if os.path.exists(f):
            return f
    return None


def gbar(cv, x, y, w, h, ang=5400000):
    sp = cv.rect(x, y, w, h, "FFFFFF", None)
    spPr = sp._element.spPr
    for el in list(spPr):
        if el.tag.endswith("}solidFill"):
            spPr.remove(el)
    spPr.insert(list(spPr).index(spPr.find(mt.q(mt.A, "prstGeom"))) + 1, mt.grad(ang))


def band(cv, x, y, w, h, title, sub):
    cv.rect(x, y, w, h, CARD2, None, shape=R, radius=0.02)
    cv.rect(x, y, w, 0.28, "2B2B2B", None)
    cv.text(x + 0.12, y, w - 0.24, 0.28, [[(title, {"bold": True, "track": 120, "color": T1}),
                                           ("   " + sub, {"color": T2})]], size=6.4, anchor=MSO_ANCHOR.MIDDLE)


def gate(cv, x, y, w, h, title, body):
    cv.rect(x, y, w, h, GATE, YELLOW, 0.9, R, 0.08)
    logo(cv, "people", x + 0.1, y + 0.08, 0.2)
    cv.text(x + 0.36, y + 0.05, w - 0.44, 0.26, title, size=7, bold=True, color=YELLOW, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(x + 0.1, y + 0.36, w - 0.18, h - 0.4, body, size=6, color=T2)


def chips(cv, x, y, w, label, items):
    cv.text(x, y, 1.05, 0.22, label, size=5.4, bold=True, color=T3, track=100, anchor=MSO_ANCHOR.MIDDLE)
    xx = x + 1.08
    for t in items:
        cw = 0.2 + 0.052 * len(t)
        if xx + cw > x + w:
            break
        cv.rect(xx, y, cw, 0.22, CARD, EDGE, 0.6, R, 0.4)
        cv.text(xx, y, cw, 0.22, t, size=5.6, color=T1, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        xx += cw + 0.06


def build(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "ASCENSION · SIERRA VOICE AGENT",
                "Sierra voice agent: governed knowledge for answers, scoped APIs for patient actions",
                "Articles in Google Cloud keep every answer accurate. Booking, cancelling and rescheduling go through "
                "the scheduling system of record, only after the caller is verified.")
    top, bot = 1.3, 5.72
    # ---------------------------------------------------------------- left: caller + Sierra
    lx, lw = 0.45, 2.05
    cv.rect(lx, top, lw, 1.05, CARD, EDGE, 0.75, R, 0.06)
    logo(cv, "people", lx + 0.12, top + 0.1, 0.24)
    cv.text(lx + 0.44, top + 0.08, lw - 0.5, 0.28, "Patient calls", size=8, bold=True, color=T1,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(lx + 0.12, top + 0.42, lw - 0.22, 0.6, "Book, cancel or reschedule, or ask a question: prep, "
            "locations, hours, insurance.", size=6.2, color=T2)
    sy = top + 1.3
    cv.rect(lx, sy, lw, bot - sy, CARD, CYAN, 1.0, R, 0.04)
    mark = sierra_logo()
    if mark:                                    # the official Sierra logo, when the file is in assets/logos
        pic = cv.s.shapes.add_picture(mark, Inches(lx + 0.12), Inches(sy + 0.12), height=Inches(0.26))
        pic.name = "Sierra Logo"
        cv.text(lx + 0.12, sy + 0.42, lw - 0.2, 0.18, "voice agent", size=6.2, color=T2,
                anchor=MSO_ANCHOR.MIDDLE)
    else:
        cv.text(lx + 0.12, sy + 0.1, lw - 0.2, 0.34, [[("Sierra", {})], [("voice agent", {"size": 6.2,
                                                                                         "bold": False,
                                                                                         "color": T2})]],
                size=8.4, bold=True, color=T1, line=0.9)
    jobs = [("Answers questions", "from approved articles only"),
            ("Verifies the caller", "before anything patient-specific"),
            ("Books, cancels, reschedules", "through the scoped API"),
            ("Hands off to a person", "when it should not decide")]
    for k, (t, d) in enumerate(jobs):
        y = sy + 0.6 + k * 0.62
        cv.rect(lx + 0.12, y + 0.07, 0.07, 0.07, CYAN, None, shape=MSO_SHAPE.OVAL)
        cv.text(lx + 0.26, y, lw - 0.34, 0.56, [[(t, {"bold": True, "color": T1})], [(d, {"color": T2,
                                                                                         "size": 5.8})]],
                size=6.6, after=1)
    cv.arrow(lx + lw / 2, top + 1.05, lx + lw / 2, sy, "8C8C8C", 1.0)
    # ---------------------------------------------------------------- middle: two paths
    mx, mw = 2.7, 7.2
    kb_y, kb_h = top, 2.05
    pt_y, pt_h = top + 2.37, bot - top - 2.37
    band(cv, mx, kb_y, mw, kb_h, "1  KNOWLEDGE PATH", "no PHI · keeps every answer accurate")
    band(cv, mx, pt_y, mw, pt_h, "2  PATIENT PATH", "PHI · verified caller only · this patient only")
    cw = (mw - 0.24 - 3 * 0.16) / 4
    cy, ch = kb_y + 0.4, 0.98
    xs = [mx + 0.12 + k * (cw + 0.16) for k in range(4)]
    card(cv, xs[0], cy, cw, ch, "KB articles in GCP", "Source of truth. Each has an owner, version and review "
         "date.", line=EDGE, fill=CARD, icon="googlecloud", tsize=6.6, bsize=5.9,
         tcolor=T1, bcolor=T2)
    card(cv, xs[1], cy, cw, ch, "PHI scan", "DLP and Collibra check every article; no patient example "
         "is published.", line=EDGE, fill=CARD, tsize=6.6, bsize=5.9, tcolor=T1, bcolor=T2)
    gate(cv, xs[2], cy, cw, ch, "Owner approves", "Clinical content gets a clinical reviewer. Conflicting "
         "versions retired.")
    card(cv, xs[3], cy, cw, ch, "Approved KB index", "Synced to Sierra on every change, so no stale answer.",
         line=CYAN, fill=CARD, tsize=6.6, bsize=5.9, tcolor=T1, bcolor=T2)
    for k in range(3):
        cv.arrow(xs[k] + cw, cy + ch / 2, xs[k + 1], cy + ch / 2, "8C8C8C", 1.0)
    rj = kb_y + kb_h - 0.2
    path(cv, [(xs[3] + cw - 0.15, cy + ch), (xs[3] + cw - 0.15, rj), (lx + lw, rj)], CYAN, 1.25)
    cv.text(xs[1] + 0.1, rj - 0.17, 3.0, 0.15, "Sierra reads approved articles only", size=5.6, color=CYAN,
            anchor=MSO_ANCHOR.MIDDLE)
    chips(cv, mx + 0.12, cy + ch + 0.1, mw - 0.24, "ARTICLES",
          ["FAQs", "prep instructions", "locations + hours", "insurance + billing", "visitor rules", "policies"])
    # boundary between the paths
    by = kb_y + kb_h + 0.16
    cv.arrow(mx, by, mx + mw, by, "5C5C5C", 1.0, dash="dash", head=False)
    cv.rect(mx + mw / 2 - 1.6, by - 0.1, 3.2, 0.2, "000000", None)
    pill(cv, mx + mw / 2 - 1.55, by - 0.09, 3.1, "NO PATIENT RECORDS IN THE KNOWLEDGE BASE", CYAN, h=0.18,
         size=5.2)
    # patient path
    py, ph = pt_y + 0.4, 1.0
    pw = [1.75, 2.55, 2.42]
    px = [mx + 0.12]
    for w in pw[:-1]:
        px.append(px[-1] + w + 0.2)
    gate(cv, px[0], py, pw[0], ph, "Verify the caller", "Date of birth plus a one-time code, before any "
         "patient detail.")
    fx, fy, fw, fh = cv.frame(px[1], py, pw[1], ph, CARD, 0.02)
    cv.text(fx + 0.1, fy + 0.04, fw - 0.2, 0.26, "Scoped scheduling API", size=7, bold=True, color=T1,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(fx + 0.1, fy + 0.32, fw - 0.18, fh - 0.34, "Returns only this patient's data and minimum fields. "
            "Enforces the rules; the agent never writes queries.", size=6, color=T2)
    cv.rect(px[2], py, pw[2], ph, CARD, EDGE, 0.75, R, 0.08)
    cv.rect(px[2] + 0.1, py + 0.08, 0.36, 0.2, CARD2, EDGE, 0.6, R, 0.3)
    cv.text(px[2] + 0.1, py + 0.08, 0.36, 0.2, "Epic", size=5.4, bold=True, color=T1, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(px[2] + 0.52, py + 0.05, pw[2] - 0.6, 0.26, "Scheduling system of record", size=7, bold=True, color=T1,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(px[2] + 0.1, py + 0.36, pw[2] - 0.18, ph - 0.4, "Live availability; book, cancel, reschedule written "
            "here. Never a GCP copy, which can lag.", size=6, color=T2)
    cv.arrow(px[0] + pw[0], py + ph / 2, px[1], py + ph / 2, "8C8C8C", 1.0)
    cv.arrow(px[1] + pw[1], py + ph / 2 - 0.12, px[2], py + ph / 2 - 0.12, CYAN, 1.25)
    cv.arrow(px[2], py + ph / 2 + 0.12, px[1] + pw[1], py + ph / 2 + 0.12, CYAN, 1.25)
    cv.arrow(lx + lw, py + ph / 2, px[0], py + ph / 2, CYAN, 1.25)
    chips(cv, mx + 0.12, py + ph + 0.14, mw - 0.24, "RULES IN THE API",
          ["cancellation window", "reschedule rules", "visit types", "provider restrictions", "minimum fields"])
    # ---------------------------------------------------------------- right: after the call
    rx, rw = 10.05, 12.88 - 10.05
    cv.rect(rx, top, rw, bot - top, CARD2, None, shape=R, radius=0.02)
    cv.rect(rx, top, rw, 0.28, "2B2B2B", None)
    cv.text(rx + 0.12, top, rw - 0.24, 0.28, "3  AFTER EVERY CALL", size=6.4, bold=True, color=T1, track=120,
            anchor=MSO_ANCHOR.MIDDLE)
    gate(cv, rx + 0.12, top + 0.4, rw - 0.24, 0.82, "Hand-off to a person", "Identity not verified, upset "
         "caller or a clinical question.")
    after = [("Masked transcripts", "Recordings hold PHI: access control, retention limits, masked before "
              "analytics.", None),
             ("Outcomes to Gold A", "Booked, cancelled, rescheduled, escalated, as aggregates for measurement.",
              "googlebigquery"),
             ("Test calls before go-live", "Wrong caller, outside the window, a question no article covers.",
              None)]
    y = top + 1.36
    for t, d, ic in after:
        card(cv, rx + 0.12, y, rw - 0.24, 0.9, t, d, line=EDGE, fill=CARD, icon=ic, tsize=6.8, bsize=6.0,
             tcolor=T1, bcolor=T2)
        y += 1.0
    # ---------------------------------------------------------------- safeguards
    gy, gh = 5.88, 0.86
    cv.text(0.45, gy, 1.3, gh, [[("SAFEGUARDS", {"bold": True, "color": T1, "track": 120})],
                                [("accurate, and safe for patients", {"color": T3, "size": 5.2})]], size=6.4,
            anchor=MSO_ANCHOR.MIDDLE, after=1)
    guards = [("BAA with Sierra", "calls carry PHI"),
              ("Verify before any PHI", "then this patient only"),
              ("System of record", "never a GCP copy"),
              ("Rules in the API", "not only in an article"),
              ("Governed articles", "owner, version, PHI-scanned"),
              ("Tested before release", "scripted test calls")]
    w = (12.88 - 1.8 - 5 * 0.1) / 6
    for k, (t, d) in enumerate(guards):
        x = 1.8 + k * (w + 0.1)
        cv.rect(x, gy, w, gh, CARD, None)
        gbar(cv, x, gy, 0.05, gh)
        cv.text(x + 0.16, gy, w - 0.22, gh, [[(t, {"bold": True, "size": 7.2, "color": T1})],
                                             [(d, {"color": T2})]], size=6.0, anchor=MSO_ANCHOR.MIDDLE, after=1)
    set_notes(s, NOTES)
    return s


def main(out=OUT):
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    build(prs, layout)
    fix_footers(prs)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print("wrote", os.path.abspath(out))


if __name__ == "__main__":
    main(*sys.argv[1:])
