"""Executive summary deck for the Ascension Agentic Factory: six slides that leadership reads before the full deck
(build_data_harness_deck.py). The same IG template and styles; no pricing.

1 Title
2 The Agentic Factory on one page: three jobs, one platform
3 What changes for Ascension: speed, people, savings
4 How it works: govern the data, agents build and test, people approve, release safely
5 Safe by design
6 Recommended next step: a measured pilot

    python3 exec_deck.py [dark|white]
"""
import os
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

import build_data_harness_deck as bd
from build_data_harness_deck import R, SRC, logo, new_slide
from build_mad_ig_deck import CYAN, MAGENTA, YELLOW, set_notes
from footer import fix_footers
from sierra_voice import CARD, CARD2, EDGE, T1, T2, T3, gbar
from story_slides import dark_header, set_bg

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Ascension Agentic Factory - Executive Summary.pptx")
OUT_WHITE = os.path.join(HERE, "Ascension Agentic Factory - Executive Summary (white).pptx")
TINTS = {CYAN: "06262B", YELLOW: "2B2508", MAGENTA: "2B0716"}
ON = "FFFFFE"           # white text on a brand colour: stays white in the white deck
EYEBROW = "ASCENSION AGENTIC FACTORY · EXECUTIVE SUMMARY"

TITLE_NOTES = """Executive summary of the Ascension Agentic Factory, read before the full deck.
One factory, three jobs: Health Cloud delivery from Codex, Marketing Cloud campaigns on governed Google Cloud data, and the Sierra voice agent on governed knowledge. Five slides: what it is, what changes for Ascension, how it works, why it is safe, and the next step."""


def head(cv, title, sub):
    dark_header(cv, EYEBROW, title, sub)


def card(cv, x, y, w, h, color, title, body, icon=None, tsize=11):
    cv.rect(x, y, w, h, TINTS[color], color, 1.0, R, 0.04)
    tx = x + 0.2
    if icon:
        logo(cv, icon, tx, y + 0.18, 0.3)
        tx += 0.42
    cv.text(tx, y + 0.14, x + w - tx - 0.15, 0.38, title, size=tsize, bold=True, color=T1, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(x + 0.2, y + 0.7, w - 0.4, h - 0.8, [[(b, {})] for b in body], size=11, color=T2, after=8)


# ================================================================== 1 title
def title_slide(prs, white, layout):
    if white:
        bd.white_title(prs, layout, "EXECUTIVE SUMMARY", "Ascension Agentic Factory",
                       "Faster delivery, governed campaigns and a safer voice agent, on the platforms Ascension "
                       "already runs.", TITLE_NOTES)
        return
    s = prs.slides[0]
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        runs = sh.text_frame.paragraphs[0].runs
        if t.lower().startswith("ascension agentic"):
            runs[0].text = "Ascension AGENTIC FACTORY"
        elif t.startswith("An Agentic Harness"):
            runs[0].text = ("Faster delivery, governed campaigns and a safer voice agent, on the platforms Ascension "
                            "already runs.")
        elif t == "AI-DRIVEN SALESFORCE DELIVERY":
            runs[0].text = "EXECUTIVE SUMMARY"
        else:
            continue
        for r in runs[1:]:
            r.text = ""
    set_notes(s, TITLE_NOTES)


# ================================================================== 2 one page
ONE_PAGE_NOTES = """SLIDE: The Agentic Factory on one page   (about 2 minutes)

KEY MESSAGE
One factory does three jobs for Ascension, on one platform: it delivers Health Cloud changes from a Jira story, it plans and builds Marketing Cloud campaigns on governed Google Cloud data, and it keeps the Sierra voice agent's answers accurate and its patient actions safe.

WALK IT
- Delivery: seven delivery agents take a Jira story from requirements to a tested release, the same agents on every pod.
- Campaigns: seven campaign agents plan, design and test on masked data; deterministic services, not agents, release the right patients.
- Sierra voice agent: answers from approved articles; books, cancels and reschedules through scoped APIs after the caller is verified.
- All three share one knowledge graph, one governance layer, one set of evals and approvals, and one change passport.
- It runs on what Ascension already owns: Codex, Salesforce, Google Cloud, Collibra, Copado."""


def one_page(prs, layout):
    s, cv = new_slide(prs, layout)
    head(cv, "The Agentic Factory: one platform, three jobs",
         "Agents do the work, people approve every step, and patient data stays governed, on the platforms Ascension "
         "already runs.")
    y, h = 1.4, 2.7
    w = (12.43 - 2 * 0.2) / 3
    jobs = [(CYAN, "openai-white", "Health Cloud delivery",
             ["Seven delivery agents take a Jira story to a tested, approved release.",
              "The same agents on every pod; each pod adds its own skills.",
              "About 2 weeks to about 2 days per change."]),
            (MAGENTA, "salesforce", "Marketing Cloud campaigns",
             ["Seven campaign agents plan, design and test on masked, governed data.",
              "Data is ready once, not rebuilt for every campaign.",
              "About 90 days to about 10-12 days per campaign."]),
            (YELLOW, None, "Sierra voice agent",
             ["Answers patients from approved, owned, PHI-scanned articles.",
              "Books, cancels and reschedules through scoped APIs, after the caller is verified.",
              "Hands off to a person when it should not decide."])]
    for k, (c, ic, t, b) in enumerate(jobs):
        card(cv, 0.45 + k * (w + 0.2), y, w, h, c, t, b, ic)
    # shared by all three
    sy = y + h + 0.25
    cv.rect(0.45, sy, 12.43, 0.95, CARD, None)
    gbar(cv, 0.45, sy, 0.06, 0.95)
    cv.text(0.7, sy, 2.2, 0.95, [[("SHARED BY ALL THREE", {"bold": True, "color": CYAN, "track": 120})],
                                 [("built once, used by every pod", {"color": T3})]], size=8,
            anchor=MSO_ANCHOR.MIDDLE, after=2)
    shared = [("Knowledge graph", "every fact cited"), ("Governance", "masking, policy, guardrails"),
              ("Evals", "every step scored"), ("Approvals", "people decide at every gate"),
              ("Change passport", "every step recorded")]
    sw = (12.88 - 3.0) / 5
    for k, (t, d) in enumerate(shared):
        cv.text(3.0 + k * sw, sy, sw - 0.1, 0.95, [[(t, {"bold": True, "color": T1, "size": 10})],
                                                   [(d, {"color": T2})]], size=8, anchor=MSO_ANCHOR.MIDDLE, after=2)
    # runs on
    ry = sy + 1.15
    cv.text(0.45, ry, 2.4, 0.4, "RUNS ON WHAT ASCENSION OWNS", size=8, bold=True, color=T2, track=120,
            anchor=MSO_ANCHOR.MIDDLE)
    owns = [("openai-white", "Codex"), ("salesforce", "Salesforce"), ("googlecloud", "Google Cloud"),
            (None, "Collibra"), (None, "Copado"), ("jira", "Jira")]
    x = 3.0
    for ic, t in owns:
        if ic:
            logo(cv, ic, x, ry + 0.08, 0.24)
            x += 0.32
        cv.text(x, ry, 1.4, 0.4, t, size=10, color=T1, anchor=MSO_ANCHOR.MIDDLE)
        x += 0.16 + 0.085 * len(t) + 0.35
    set_notes(s, ONE_PAGE_NOTES)
    return s


# ================================================================== 3 outcomes
OUTCOME_NOTES = """SLIDE: What changes for Ascension   (about 2 minutes)

KEY MESSAGE
Faster delivery, faster campaigns and leaner squads, with people still approving every change.

THE NUMBERS
- A Health Cloud change: about 2 weeks today, from Jira story to UAT sign-off; about 2 days on the factory, story to production. Agent-assisted delivery at Ascension has already measured 1-2 weeks down to 2-4 days.
- A Marketing Cloud campaign: about 90 days today; about 10-12 days, because the data is governed and ready once instead of prepared for every campaign. Legal review becomes the pace setter.
- A squad of 5-7 people becomes about 2.4: a PM / tech lead, a dev / QA directing the agents, and a PM / TL shared across squads.
- About $19.9M saved over three years across the pods, about 50% lower at full run rate.

BE HONEST
The factory figures are pilot targets and planning estimates, confirmed against each pod's Phase 0 baseline."""


def outcomes(prs, layout):
    s, cv = new_slide(prs, layout)
    head(cv, "What changes for Ascension",
         "Faster delivery, faster campaigns and leaner squads, with people still approving every change.")
    y, h = 1.45, 3.55
    w = (12.43 - 3 * 0.18) / 4
    tiles = [(CYAN, "HEALTH CLOUD CHANGE", "~2 weeks", "~2 days", "Jira story to production; 2-4 days already "
              "measured with agent assistance"),
             (MAGENTA, "CAMPAIGN LAUNCH", "~90 days", "~10-12 days", "data governed and ready once, not prepared "
              "for every campaign"),
             (YELLOW, "PEOPLE PER SQUAD", "5-7", "~2.4", "a PM / tech lead and a dev / QA directing the agents; "
              "people approve and own"),
             (CYAN, "SAVED OVER 3 YEARS", "", "~$19.9M", "about 50% lower at full run rate, across the pods")]
    for k, (c, lab, before, after, d) in enumerate(tiles):
        x = 0.45 + k * (w + 0.18)
        cv.rect(x, y, w, h, CARD, None)
        gbar(cv, x, y, w, 0.06, 0)
        cv.text(x + 0.22, y + 0.25, w - 0.4, 0.25, lab, size=8, bold=True, color=c, track=120,
                anchor=MSO_ANCHOR.MIDDLE)
        if before:
            cv.text(x + 0.22, y + 0.7, w - 0.4, 0.4, [[("TODAY  ", {"size": 7.5, "color": T3, "bold": True}),
                                                      (before, {"size": 18, "color": T3})]], anchor=MSO_ANCHOR.MIDDLE)
            cv.text(x + 0.22, y + 1.2, w - 0.4, 0.25, "↓", size=14, color=T3, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.22, y + 1.5, w - 0.4, 0.75, after, size=34, color=c, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.22, y + 2.35, w - 0.4, 1.0, d, size=10.5, color=T2)
    # the line underneath
    by = y + h + 0.25
    fx, fy, fw, fh = cv.frame(0.45, by, 12.43, 0.85, "0A0A0A", 0.02)
    cv.text(fx + 0.25, fy, fw - 0.5, fh, [[("What stays the same: ", {"bold": True, "color": YELLOW}),
                                          ("people approve at every gate, the owning team deploys to production, and "
                                           "agents never see raw patient data.", {"color": T1})]], size=11,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(0.45, by + 0.95, 12.43, 0.22, "Factory figures are pilot targets and planning estimates, confirmed "
            "against each pod's Phase 0 baseline.", size=7.5, color=T3, anchor=MSO_ANCHOR.MIDDLE)
    set_notes(s, OUTCOME_NOTES)
    return s


# ================================================================== 4 how it works
HOW_NOTES = """SLIDE: How it works   (about 2 minutes)

KEY MESSAGE
Four steps, the same for every job: the data is governed first, agents plan, build and test on masked data, people approve, and only then does anything reach production or a patient.

WALK IT
1 Govern the data: in Google Cloud, Bronze, Silver and Gold layers; DLP classifies, a Collibra steward approves, policy tags mask PHI. Agents get aggregates and masked views only.
2 Agents plan, build and test: seven delivery agents and seven campaign agents, with cited context from the knowledge graph. Every step is scored by evals; a failure goes back to the step that caused it, at most five times, then a named owner.
3 People approve: product owner, architect, tech lead and business for delivery (H1-H4, then CAB); marketer, clinical and compliance for campaigns (G1-G3).
4 Release safely: delivery through Copado by the owning team, with no agent credential in production; campaign patients only through the activation gateway, with consent re-checked at send.
Underneath: the Change Guardian checks every change against who owns and uses it, and the change passport records every step."""


def how(prs, layout):
    s, cv = new_slide(prs, layout)
    head(cv, "How it works: four steps, the same for every job",
         "The data is governed first, agents build and test on masked data, people approve, and only then does "
         "anything reach production or a patient.")
    y, h = 1.45, 3.3
    w = (12.43 - 3 * 0.32) / 4
    steps = [(MAGENTA, "1", "Govern the data", "googlecloud",
              ["Google Cloud: Bronze, Silver, Gold.", "DLP classifies; a Collibra steward approves; policy tags "
               "mask PHI.", "Agents get aggregates and masked views only."]),
             (CYAN, "2", "Agents plan, build, test", "harness",
              ["7 delivery + 7 campaign agents.", "Cited context from the knowledge graph.",
               "Every step scored; a failure goes back to its cause."]),
             (YELLOW, "3", "People approve", "people",
              ["Delivery: product owner, architect, tech lead, business, then CAB.",
               "Campaigns: marketer, clinical, compliance."]),
             (CYAN, "4", "Release safely", "salesforce",
              ["Delivery: the owning team deploys through Copado; no agent credential in production.",
               "Campaigns: patients only through the activation gateway; consent re-checked at send."])]
    for k, (c, n, t, ic, body) in enumerate(steps):
        x = 0.45 + k * (w + 0.32)
        cv.rect(x, y, w, h, TINTS[c], c, 1.0, R, 0.04)
        cv.rect(x, y, w, 0.5, c, None)
        cv.text(x + 0.2, y, w - 0.4, 0.5, [[(n + "  ", {"bold": True}), (t, {"bold": True})]], size=12,
                color=ON if c == MAGENTA else "000000", anchor=MSO_ANCHOR.MIDDLE)
        logo(cv, ic, x + 0.2, y + 0.7, 0.32)
        cv.text(x + 0.2, y + 1.2, w - 0.4, h - 1.3, [[(b, {})] for b in body], size=11, color=T2, after=8)
        if k < 3:
            cv.arrow(x + w + 0.04, y + h / 2, x + w + 0.28, y + h / 2, T2, 2.0)
    # underneath
    by = y + h + 0.3
    half = (12.43 - 0.2) / 2
    under = [("Change Guardian", "Checks every change against the pods that own and use it, and flags deployment "
              "blockers before release. No pod breaks another."),
             ("Change passport", "Every step, approval, test and release is recorded, so every change can be "
              "audited end to end.")]
    for k, (t, d) in enumerate(under):
        x = 0.45 + k * (half + 0.2)
        cv.rect(x, by, half, 1.05, CARD, None)
        gbar(cv, x, by, 0.06, 1.05)
        cv.text(x + 0.25, by + 0.12, half - 0.4, 0.3, t, size=11, bold=True, color=CYAN if k == 0 else T1,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.25, by + 0.45, half - 0.4, 0.55, d, size=10, color=T2)
    set_notes(s, HOW_NOTES)
    return s


# ================================================================== 5 safe by design
SAFE_NOTES = """SLIDE: Safe by design   (about 2 minutes)

KEY MESSAGE
Safety is built into the platform, not added at the end. Six guarantees hold for every pod and every job.

WALK IT
- Agents never see raw PHI: aggregates and masked views only; patient data leaves only through the activation gateway.
- No production credentials: three identities (read, develop, release); the owning team deploys.
- People decide: every gate is a named person; agents recommend, people approve.
- The right patient only: eligibility, consent and suppression are checked by deterministic services, and consent is re-checked at send.
- No pod breaks another: the Change Guardian checks every change against who owns and uses it.
- Every step recorded: the change passport and evals make every change auditable.
And on HIPAA: masking alone is not compliance. Classification and steward approval stay, and de-identification, permitted use, tenant isolation and the full HIPAA controls are designed in, not assumed."""


def safe(prs, layout):
    s, cv = new_slide(prs, layout)
    head(cv, "Safe by design",
         "Six guarantees hold for every pod and every job. Safety is part of the platform, not a step at the end.")
    y = 1.45
    w = (12.43 - 2 * 0.2) / 3
    h = 1.6
    items = [(MAGENTA, "Agents never see raw PHI", "Aggregates and masked views only. Patient data leaves only "
              "through the activation gateway."),
             (CYAN, "No production credentials", "Three identities: read, develop, release. The owning team "
              "deploys after CAB."),
             (YELLOW, "People decide", "Every gate is a named person. Agents recommend; people approve."),
             (MAGENTA, "The right patient only", "Eligibility, consent and suppression checked by deterministic "
              "services; consent re-checked at send."),
             (CYAN, "No pod breaks another", "The Change Guardian checks every change against who owns and uses "
              "it."),
             (YELLOW, "Every step recorded", "The change passport and evals make every change auditable.")]
    for k, (c, t, d) in enumerate(items):
        x = 0.45 + (k % 3) * (w + 0.2)
        yy = y + (k // 3) * (h + 0.2)
        cv.rect(x, yy, w, h, CARD, None)
        cv.rect(x, yy, 0.07, h, c, None)
        cv.text(x + 0.3, yy + 0.2, w - 0.5, 0.4, t, size=13, bold=True, color=c, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.3, yy + 0.7, w - 0.5, h - 0.8, d, size=10, color=T2)
    by = y + 2 * h + 0.4 + 0.1
    fx, fy, fw, fh = cv.frame(0.45, by, 12.43, 1.0, "0A0A0A", 0.02)
    cv.text(fx + 0.25, fy, fw - 0.5, fh, [[("HIPAA: masking alone is not compliance. ", {"bold": True,
                                                                                         "color": YELLOW}),
                                          ("Classification and steward approval stay; safe de-identification, "
                                           "permitted use, tenant isolation and the full HIPAA controls are designed "
                                           "in, not assumed.", {"color": T1})]], size=11, anchor=MSO_ANCHOR.MIDDLE)
    set_notes(s, SAFE_NOTES)
    return s


# ================================================================== 6 next step
NEXT_NOTES = """SLIDE: Recommended next step: a measured pilot   (about 2 minutes)

KEY MESSAGE
Start with one Health Cloud pod and one campaign, measure against today's baseline, then add pods one at a time.

THE PILOT
- Health Cloud: Network Services pod stories, Jira to production, measured against today's ~2 weeks.
- Campaigns: the CHF post-discharge nudge, measured against today's ~90 days, built and tested on synthetic data before any real patient.
- The Sierra voice agent follows on the same governance and knowledge base.

WHAT WE MEASURE
Cycle time per story and per campaign, first-pass approval at each gate, defects after release, eval scores per agent, and zero unauthorized access to patient data.

WHAT WE NEED FROM ASCENSION
Named pilot owners, access to the governed data with a Collibra steward, Codex usage limits, and agreed success criteria. Then the full deck walks through each part in detail."""


def next_step(prs, layout):
    s, cv = new_slide(prs, layout)
    head(cv, "Recommended next step: a measured pilot",
         "Start with one Health Cloud pod and one campaign, measure against today's baseline, then add pods one at a "
         "time.")
    y, h = 1.45, 2.65
    w = (12.43 - 2 * 0.2) / 3
    cols = [(CYAN, "THE PILOT", [("Health Cloud", "Network Services pod stories, Jira to production"),
                                 ("Campaigns", "CHF post-discharge nudge, proven on synthetic data first"),
                                 ("Sierra voice agent", "follows on the same governance and knowledge")]),
            (YELLOW, "WHAT WE MEASURE", [("Speed", "cycle time per story and per campaign"),
                                         ("Quality", "first-pass approval, defects after release, eval scores"),
                                         ("Safety", "zero unauthorized access to patient data")]),
            (MAGENTA, "WHAT WE NEED FROM ASCENSION", [("Owners", "a named owner for each pilot"),
                                                      ("Data access", "governed Google Cloud data and a Collibra "
                                                       "steward"),
                                                      ("Limits + criteria", "Codex usage limits and agreed success "
                                                       "criteria")])]
    for k, (c, t, rows) in enumerate(cols):
        x = 0.45 + k * (w + 0.2)
        cv.rect(x, y, w, h, TINTS[c], c, 1.0, R, 0.03)
        cv.rect(x, y, w, 0.4, c, None)
        cv.text(x + 0.2, y, w - 0.4, 0.4, t, size=9, bold=True, track=120, color=ON if c == MAGENTA else "000000",
                anchor=MSO_ANCHOR.MIDDLE)
        for j, (a, b) in enumerate(rows):
            yy = y + 0.55 + j * 0.68
            cv.text(x + 0.2, yy, w - 0.4, 0.62, [[(a, {"bold": True, "color": T1, "size": 10.5})],
                                                 [(b, {"color": T2})]], size=9, after=2)
    # path
    py = y + h + 0.35
    cv.text(0.45, py, 12.43, 0.25, "THE PATH", size=8, bold=True, color=T2, track=150, anchor=MSO_ANCHOR.MIDDLE)
    stages = [("Phase 0", "baseline today: cycle times, costs, quality"), ("Pilot", "one Health Cloud pod, one "
              "campaign"), ("Measure", "against the baseline; adjust"), ("Scale", "add pods one at a time; each "
              "adds its skills")]
    sw = (12.43 - 3 * 0.25) / 4
    for k, (t, d) in enumerate(stages):
        x = 0.45 + k * (sw + 0.25)
        cv.rect(x, py + 0.35, sw, 0.95, CARD, None)
        gbar(cv, x, py + 0.35, sw, 0.05, 0)
        cv.text(x + 0.2, py + 0.45, sw - 0.4, 0.8, [[(t, {"bold": True, "color": T1, "size": 11})],
                                                    [(d, {"color": T2})]], size=9, anchor=MSO_ANCHOR.MIDDLE,
                after=2)
        if k < 3:
            cv.arrow(x + sw + 0.03, py + 0.82, x + sw + 0.22, py + 0.82, T2, 1.5)
    set_notes(s, NEXT_NOTES)
    return s


# ================================================================== build
def main(theme="dark"):
    import combined_flow
    import muted_theme as mt
    white = theme == "white"
    if white:
        combined_flow.light_logos()
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst)[1 if not white else 0:]:
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    title_slide(prs, white, layout)
    for build in (one_page, outcomes, how, safe, next_step):
        build(prs, layout)
    if white:
        for sl in list(prs.slides)[1:]:
            set_bg(sl, "FFFFFF")
            mt.to_light(sl)
    for sl in prs.slides:
        mt.min_font(sl, 7.0)
    for line in fix_footers(prs):
        print(line)
    out = OUT_WHITE if white else OUT
    prs.save(out)
    print("wrote", out, "-", len(prs.slides), "slides")


if __name__ == "__main__":
    main(*sys.argv[1:])
