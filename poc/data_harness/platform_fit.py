"""One-slide preview: Platform fit. Core benefits, how one agent serves many pods at once, and how the platform is
built, measured and controlled. Black muted IG style. Not part of the deck yet.

    python3 platform_fit.py [out.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from build_data_harness_deck import R, SRC, logo, new_slide, path
from build_mad_ig_deck import CYAN, YELLOW, set_notes
from footer import fix_footers
from sierra_voice import CARD, CARD2, EDGE, GATE, T1, T2, T3, gbar
from story_slides import dark_header

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "build", "platform_fit.pptx")
GREY = "8C8C8C"

NOTES = """SLIDE: Platform fit   (about 3 minutes)

KEY MESSAGE
One harness serves every pod and three jobs: delivery, campaigns and the Sierra voice agent. An agent is one versioned definition; every request becomes its own isolated run, so ten pods can use the Requirements Agent at once and each gets the same quality. Load changes the wait, not the answer.

ONE AGENT, MANY PODS
- Each pod's request goes into a fair-share queue, so no pod can crowd out another.
- The Requirements Agent is one definition, versioned and evaluated once. Each request starts its own run: its own context slice from the graph, its own pod skill pack, its own branch, its own checkpoint. Runs share no memory.
- The Change Guardian checks whether two runs touch the same object; if they do, the second is flagged and sequenced. Evals score every run; people approve at H1.
- Concurrency (proposed, sized in the Phase 0 load test): one run per container, autoscaled; up to 3 runs in flight per pod; a global cap set by the model and API quotas, for example 30; extra requests wait in the queue, nothing is dropped.

BUILD, MEASURE, CONTROL
- Build: one agent definition plus pod skill packs; versioned prompts, skills and models; a new agent version must pass its eval set before release; pods onboard one at a time.
- Measure: eval score per agent and pod, first-pass approval at gates, retries, cited-fact rate, latency and queue wait, defects found after release, and zero PHI exposure in tests.
- Control: three identities with no production access, a tool gateway allowlist, Model Armor guardrails, human gates, the change passport, per-pod quotas and a kill switch per agent or pod.

CORE BENEFITS
Consistency, reuse, faster onboarding of a new pod, quality on every step, auditability, and safety by design."""


def pill_box(cv, x, y, w, h, t, line=EDGE, fill=CARD, tc=T1, size=6.2, bold=True):
    cv.rect(x, y, w, h, fill, line, 0.75, R, 0.3)
    cv.text(x, y, w, h, t, size=size, bold=bold, color=tc, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def panel(cv, x, y, w, h, title, sub=""):
    cv.rect(x, y, w, h, CARD2, None, shape=R, radius=0.02)
    cv.rect(x, y, w, 0.28, "2B2B2B", None)
    cv.text(x + 0.12, y, w - 0.24, 0.28, [[(title, {"bold": True, "track": 120, "color": T1}),
                                           ("   " + sub, {"color": T2})] if sub else
                                          [(title, {"bold": True, "track": 120, "color": T1})]],
            size=6.4, anchor=MSO_ANCHOR.MIDDLE)


def bullets(cv, x, y, w, items, step=0.27, size=6.1, dot=CYAN):
    for k, (b, d) in enumerate(items):
        yy = y + k * step
        cv.rect(x, yy + 0.08, 0.06, 0.06, dot, None, shape=MSO_SHAPE.OVAL)
        cv.text(x + 0.12, yy, w - 0.12, step, [[(b, {"bold": True, "color": T1}), ("  " + d, {"color": T2})]],
                size=size, anchor=MSO_ANCHOR.TOP)


def build(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "ASCENSION AGENTIC HARNESS · PLATFORM FIT",
                "Platform fit: one harness for every pod, every job",
                "Delivery, campaigns and the Sierra voice agent share one graph, one governance layer and one approval "
                "flow. Each agent is built once and serves every pod at once.")
    # ---------------------------------------------------------------- top: three jobs, one platform
    ty, th = 1.3, 0.5
    cv.text(0.45, ty, 1.5, th, [[("ONE HARNESS", {"bold": True, "track": 120, "color": T1})],
                                [("three jobs", {"color": T3, "size": 5.4})]], size=6.4, anchor=MSO_ANCHOR.MIDDLE,
            after=1)
    jobs = [("openai-white", "Delivery", "Codex to production · every pod"),
            ("salesforce", "Campaigns", "Marketing Cloud on governed GCP data"),
            (None, "Sierra voice agent", "bookings, reschedules, questions")]
    jw = 2.45
    for k, (ic, t, d) in enumerate(jobs):
        x = 1.95 + k * (jw + 0.12)
        cv.rect(x, ty, jw, th, CARD, CYAN, 0.75, R, 0.12)
        tx = x + 0.12
        if ic:
            logo(cv, ic, tx, ty + 0.14, 0.22)
            tx += 0.3
        cv.text(tx, ty, x + jw - tx - 0.08, th, [[(t, {"bold": True, "color": T1})], [(d, {"color": T2,
                                                                                         "size": 5.6})]],
                size=7.2, anchor=MSO_ANCHOR.MIDDLE)
    sx = 1.95 + 3 * (jw + 0.12) + 0.05
    cv.text(sx, ty, 12.88 - sx, th, [[("SHARED  ", {"bold": True, "color": CYAN, "track": 100}),
                                      ("knowledge graph · governance · evals · approvals · change passport",
                                       {"color": T2})]], size=6.2, anchor=MSO_ANCHOR.MIDDLE)
    # ---------------------------------------------------------------- left: one agent, many pods
    lx, ly, lw, lh = 0.45, 1.95, 6.95, 3.62
    panel(cv, lx, ly, lw, lh, "ONE AGENT, MANY PODS AT ONCE", "example: 10 pods call the Requirements Agent together")
    pods = ["Network Services", "Clinical Transformation", "Contact Center", "Nudge / Demand", "VBC outreach",
            "QA", "Pod 7", "Pod 8", "Pod 9", "Pod 10"]
    px, pw, p0, pstep, ph = lx + 0.12, 1.32, ly + 0.4, 0.232, 0.19
    for k, t in enumerate(pods):
        pill_box(cv, px, p0 + k * pstep, pw, ph, t, size=5.4, bold=False,
                 tc=T1 if k < 6 else T3, line=EDGE if k < 6 else "262626")
    pods_bot = p0 + 9 * pstep + ph
    # queue
    qx, qw = px + pw + 0.22, 1.12
    qy, qh = p0, pods_bot - p0
    cv.rect(qx, qy, qw, qh, CARD, EDGE, 0.75, R, 0.06)
    cv.text(qx + 0.08, qy + 0.06, qw - 0.14, 0.4, [[("Fair-share", {})], [("queue", {})]], size=6.8, bold=True,
            color=T1, line=0.9)
    for k in range(6):
        cv.rect(qx + 0.14, qy + 0.62 + k * 0.2, qw - 0.28, 0.13, "2B2B2B", None, shape=R, radius=0.3)
    cv.text(qx + 0.08, qy + 1.9, qw - 0.14, 0.38, "Per-pod cap, so no pod crowds out another.", size=5.5,
            color=T2)
    for k in range(10):
        yy = p0 + k * pstep + ph / 2
        cv.arrow(px + pw, yy, qx, qy + qh / 2, "4A4A4A", 0.5, head=False)
    # agent + isolated runs
    ax, aw = qx + qw + 0.22, 2.0
    cv.rect(ax, p0, aw, 0.52, CARD, CYAN, 0.9, R, 0.08)
    logo(cv, "harness", ax + 0.1, p0 + 0.08, 0.2)
    cv.text(ax + 0.36, p0 + 0.03, aw - 0.42, 0.26, "Requirements Agent", size=7.2, bold=True, color=T1,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(ax + 0.1, p0 + 0.3, aw - 0.16, 0.2, "one definition · versioned · evaluated once", size=5.5, color=T2,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.arrow(qx + qw, qy + qh / 2, ax, p0 + 0.26, CYAN, 1.0)
    runs = [("Run · NS-1234", "Network Services"), ("Run · VBC-88", "VBC outreach"), ("Run · CT-412", "Clinical "
                                                                                          "Transformation")]
    ry = p0 + 0.66
    for k, (t, pod) in enumerate(runs):
        y = ry + k * 0.5
        cv.rect(ax, y, aw, 0.42, CARD, EDGE, 0.75, R, 0.1)
        cv.text(ax + 0.1, y + 0.02, aw - 0.16, 0.2, [[(t, {"bold": True, "color": T1}), ("  " + pod,
                                                                                       {"color": T3})]],
                size=6.0, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(ax + 0.1, y + 0.21, aw - 0.16, 0.19, "own context slice · skill pack · branch", size=5.3,
                color=T2, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(ax, ry + 1.5, aw, 0.18, "+ one isolated run per request, autoscaled", size=5.4, color=CYAN,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cv.arrow(ax + aw / 2, p0 + 0.52, ax + aw / 2, ry, GREY, 0.75)
    # guard, evals, gate
    gx = ax + aw + 0.22
    gw = lx + lw - 0.12 - gx
    steps = [("Change Guardian", "two runs on the same object: flagged and sequenced", EDGE, CARD, T1),
             ("Evals on every run", "same score bar for every pod", EDGE, CARD, T1),
             ("H1 · people approve", "product owner of that pod", YELLOW, GATE, YELLOW)]
    for k, (t, d, ln, fl, tc) in enumerate(steps):
        y = p0 + k * 0.6
        cv.rect(gx, y, gw, 0.5, fl, ln, 0.75, R, 0.1)
        if k == 2:
            logo(cv, "people", gx + 0.08, y + 0.07, 0.18)
        cv.text(gx + (0.32 if k == 2 else 0.1), y + 0.04, gw - 0.4, 0.22, t, size=6.6, bold=True, color=tc,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(gx + 0.1, y + 0.25, gw - 0.16, 0.24, d, size=5.4, color=T2)
        if k < 2:
            cv.arrow(gx + gw / 2, y + 0.5, gx + gw / 2, y + 0.6, GREY, 0.75)
    path(cv, [(ax + aw, ry + 0.75), (gx - 0.11, ry + 0.75), (gx - 0.11, p0 + 0.3), (gx, p0 + 0.3)], CYAN, 1.0)
    # accuracy callout
    cy = p0 + 3 * 0.6 + 0.02
    cv.rect(gx, cy, gw, pods_bot - cy, "06262B", CYAN, 0.75, R, 0.08)
    cv.text(gx + 0.1, cy, gw - 0.16, pods_bot - cy, [[("Accuracy holds at 10 pods", {"bold": True,
                                                                                       "color": CYAN})],
                                                     [("Runs share no memory, so load changes the wait, not the "
                                                       "answer.", {"color": T1})]], size=5.8,
            anchor=MSO_ANCHOR.MIDDLE, after=1)
    # concurrency facts
    fy = pods_bot + 0.14
    fh = ly + lh - 0.1 - fy
    facts = [("1 run per container", "autoscaled on Cloud Run"),
             ("Up to 3 in flight per pod", "fair share across pods"),
             ("Global cap ~30 runs", "set by model and API quotas"),
             ("Queued, never dropped", "resumes from checkpoints")]
    fw = (lw - 0.24 - 3 * 0.1) / 4
    for k, (t, d) in enumerate(facts):
        x = lx + 0.12 + k * (fw + 0.1)
        cv.rect(x, fy, fw, fh, CARD, None)
        cv.rect(x, fy, 0.04, fh, CYAN, None)
        cv.text(x + 0.12, fy, fw - 0.16, fh, [[(t, {"bold": True, "color": T1})], [(d, {"color": T2,
                                                                                      "size": 5.4})]],
                size=6.4, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(lx + lw - 3.2, ly + 0.02, 3.08, 0.24, "concurrency: proposed, sized in the Phase 0 load test",
            size=5.2, color=T3, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # ---------------------------------------------------------------- right: build, measure, control
    rx, rw = 7.55, 12.88 - 7.55
    blocks = [("HOW WE BUILD", [("One definition", "per agent, plus pod skill packs"),
                                ("Versioned", "prompts, skills, model in config"),
                                ("Eval gate", "a new version ships only if it passes"),
                                ("Pod by pod", "onboard one, measure, then the next")]),
              ("HOW WE MEASURE", [("Eval score", "per agent, per pod, every run"),
                                  ("First-pass approval", "share accepted at H1-H4, G1-G3"),
                                  ("Grounding + retries", "cited facts, retries per run"),
                                  ("Wait + defects", "queue time, issues after release")]),
              ("HOW WE CONTROL", [("No production access", "three identities, tool allowlist"),
                                  ("Guardrails", "Model Armor in and out, no raw PHI"),
                                  ("People decide", "human gates, change passport"),
                                  ("Quotas + kill switch", "per pod, per agent")])]
    bh = (lh - 2 * 0.1) / 3
    for k, (t, items) in enumerate(blocks):
        y = ly + k * (bh + 0.1)
        panel(cv, rx, y, rw, bh, t)
        cw = (rw - 0.36) / 2
        for j, it in enumerate(items):
            x = rx + 0.14 + (j % 2) * (cw + 0.08)
            yy = y + 0.38 + (j // 2) * 0.36
            cv.rect(x, yy + 0.08, 0.06, 0.06, YELLOW if (k == 2 and j == 2) else CYAN, None,
                    shape=MSO_SHAPE.OVAL)
            cv.text(x + 0.12, yy, cw - 0.12, 0.34, [[(it[0], {"bold": True, "color": T1})],
                                                    [(it[1], {"color": T2, "size": 5.6})]], size=6.4,
                    anchor=MSO_ANCHOR.TOP)
    # ---------------------------------------------------------------- bottom: core benefits
    by, bh2 = 5.72, 1.1
    cv.text(0.45, by, 1.3, bh2, [[("CORE BENEFITS", {"bold": True, "color": T1, "track": 120})],
                                 [("beyond cost", {"color": T3, "size": 5.4})]], size=6.4,
            anchor=MSO_ANCHOR.MIDDLE, after=1)
    benefits = [("Consistency", "Same standards and checks in every pod, every time."),
                ("Reuse", "Build once; a new pod adds skills, not agents."),
                ("Knowledge kept", "The graph keeps what people know when they move on."),
                ("Quality", "Every step scored before it moves."),
                ("Auditability", "Every step, approval and test recorded."),
                ("Safe by design", "No production credentials, no raw PHI for agents.")]
    w = (12.88 - 1.8 - 5 * 0.1) / 6
    for k, (t, d) in enumerate(benefits):
        x = 1.8 + k * (w + 0.1)
        cv.rect(x, by, w, bh2, CARD, None)
        gbar(cv, x, by, 0.05, bh2)
        cv.text(x + 0.16, by + 0.04, w - 0.24, bh2 - 0.08, [[(t, {"bold": True, "size": 8, "color": T1})],
                                                            [(d, {"color": T2})]], size=6.2,
                anchor=MSO_ANCHOR.MIDDLE, after=2)
    set_notes(s, NOTES)
    return s


def main(out=OUT):
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    build(prs, layout)
    fix_footers(prs)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print("wrote", os.path.abspath(out))


if __name__ == "__main__":
    main(*sys.argv[1:])
