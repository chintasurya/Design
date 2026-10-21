"""One slide: change management. People and adoption (how pods move onto the harness) and the Change Guardian
(no pod can break another pod without knowing first). Colourful black IG style. Preview until approved.

    python3 change_mgmt.py [out.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from build_data_harness_deck import R, SRC, badge, logo, new_slide
from build_mad_ig_deck import CYAN, MAGENTA, YELLOW, set_notes
from footer import fix_footers
from sierra_voice import CARD, EDGE, T1, T2, T3, gbar
from story_slides import dark_header

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "build", "change_mgmt.pptx")
TINTS = {CYAN: "06262B", YELLOW: "2B2508", MAGENTA: "2B0716"}

NOTES = """SLIDE: Change management   (about 3 minutes)

KEY MESSAGE
Two kinds of change, both managed. People: pods move onto the harness one wave at a time, every role knows what changes for them, and adoption is measured. Code: the Change Guardian checks every change against who owns and uses each component, so no pod can break another pod without knowing first.

LEFT · PEOPLE AND ADOPTION
- Waves: a pilot on two pods (Network Services and a Nudge campaign), measured against the manual baseline; then wave 1 (five pods) and wave 2 (all pods), each after its gate is signed.
- What changes for each role: developers direct the agents and approve; QA verifies and signs off; product owners approve the brief at H1; architects approve the design at H2 with impact already found; a small platform team keeps agents, graph and rules current.
- Enablement: role-based training, a champion in every pod, playbooks and office hours, and a show-and-tell after each pilot so the next pods see real results.

RIGHT · CHANGE GUARDIAN
- How it decides: it reads the change, the ownership map from the knowledge graph and the guardrail rulebook, then returns PASS (moves on with its change passport), NEEDS APPROVAL (affected pod leads sign first) or BLOCK + FIX (reason given; the Developer Agent drafts the fix).
- Checkpoints no change can skip: story intake (blast radius before build), every commit (blocking, not advisory), Copado promotion (no passport, no promotion) and every night (drift outside the pipeline flagged within a day).
- Rulebook examples: a flow on a shared object without record-type criteria; an edit to a shared Apex class without tests for every pod that uses it; Modify All in a permission set; any change outside the pipeline.

BOTTOM · PILOT TARGETS
Every change checked, cross-pod impact known before build, drift caught within a day, CAB evidence in minutes, new engineers productive in about a week, and adoption measured in every pod. Figures are pilot targets."""


def band(cv, x, y, w, h, n, title, sub, color):
    cv.rect(x, y, w, h, TINTS[color], color, 0.75, R, 0.02)
    cv.rect(x, y, w, 0.3, color, None)
    tc = "FFFFFE" if color == MAGENTA else "000000"
    cv.text(x + 0.12, y, w - 0.24, 0.3, [[("%d  %s" % (n, title), {"bold": True, "track": 120}),
                                          ("   " + sub, {})]], size=7, color=tc, anchor=MSO_ANCHOR.MIDDLE)


def label(cv, x, y, w, t, color):
    cv.text(x, y, w, 0.2, t, size=7, bold=True, color=color, track=100, anchor=MSO_ANCHOR.MIDDLE)


def build(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "ASCENSION AGENTIC HARNESS · CHANGE MANAGEMENT",
                "Change management: people adopt it in waves, and no pod breaks another",
                "Pods move onto the harness in waves, every role prepared. Every change passes the Change Guardian, so "
                "cross-pod impact is known before build.")
    top, h = 1.3, 4.42
    # ---------------------------------------------------------------- 1 people and adoption
    lx, lw = 0.45, 6.12
    band(cv, lx, top, lw, h, 1, "PEOPLE AND ADOPTION", "pods move one wave at a time", CYAN)
    ix, iw = lx + 0.14, lw - 0.28
    label(cv, ix, top + 0.38, 3, "ROLLOUT IN WAVES", CYAN)
    waves = [("Pilot", "2 pods", "Network Services + a Nudge campaign, measured vs baseline"),
             ("Wave 1", "5 pods", "after the pilot gate is signed"),
             ("Wave 2", "all pods", "then the run team takes over")]
    ww = (iw - 2 * 0.08) / 3
    for k, (t, n, d) in enumerate(waves):
        x = ix + k * (ww + 0.08)
        cv.rect(x, top + 0.62, ww, 0.3, CYAN if k == 0 else "0A4A52", None,
                shape=MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON)
        cv.text(x + (0.24 if k else 0.1), top + 0.62, ww - 0.34, 0.3, [[(t + "  ", {"bold": True}), (n, {})]],
                size=7.4, color="000000" if k == 0 else T1, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.04, top + 0.96, ww - 0.1, 0.36, d, size=7, color=T2)
    label(cv, ix, top + 1.42, 3, "WHAT CHANGES FOR EACH ROLE", CYAN)
    roles = [("Developer", "writes the code", "directs the agents, reviews and approves"),
             ("QA", "writes tests by hand", "verifies against criteria and signs off"),
             ("Product owner", "chases unclear stories", "approves the brief at H1"),
             ("Architect", "finds impact by hand", "approves the design at H2, impact already found"),
             ("Platform team", "new", "keeps agents, graph and rules current")]
    c1, c2 = ix + 1.2, ix + 2.75
    cv.text(c1, top + 1.64, 1.5, 0.18, "TODAY", size=7, bold=True, color=T3, track=80, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(c2, top + 1.64, 2.6, 0.18, "WITH THE HARNESS", size=7, bold=True, color=CYAN, track=80,
            anchor=MSO_ANCHOR.MIDDLE)
    for k, (r, a, b) in enumerate(roles):
        y = top + 1.86 + k * 0.32
        cv.rect(ix, y, iw, 0.28, CARD, None)
        cv.text(ix + 0.1, y, 1.1, 0.28, r, size=7.2, bold=True, color=T1, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(c1, y, 1.5, 0.28, a, size=7, color=T2, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(c2, y, iw - (c2 - ix) - 0.06, 0.28, b, size=7, bold=True, color=CYAN, anchor=MSO_ANCHOR.MIDDLE)
    label(cv, ix, top + 3.5, 3, "ENABLEMENT", CYAN)
    chips = ["Role-based training", "A champion in every pod", "Playbooks + office hours", "Show-and-tell per pilot"]
    cw = (iw - 3 * 0.08) / 4
    for k, t in enumerate(chips):
        x = ix + k * (cw + 0.08)
        cv.rect(x, top + 3.74, cw, 0.5, CARD, CYAN, 0.6, R, 0.2)
        cv.text(x + 0.04, top + 3.74, cw - 0.08, 0.5, t, size=7.2, bold=True, color=T1, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
    # ---------------------------------------------------------------- 2 change guardian
    rx = lx + lw + 0.14
    rw = 12.88 - rx
    band(cv, rx, top, rw, h, 2, "CHANGE GUARDIAN", "no pod breaks another pod without knowing first", MAGENTA)
    jx, jw = rx + 0.14, rw - 0.28
    label(cv, jx, top + 0.38, 3, "HOW IT DECIDES", MAGENTA)
    ins = [("The change", "code diff and metadata"), ("Ownership map", "who owns and uses each part"),
           ("Guardrail rulebook", "Ascension rules as code")]
    inw = 1.55
    for k, (t, d) in enumerate(ins):
        y = top + 0.62 + k * 0.42
        cv.rect(jx, y, inw, 0.36, CARD, EDGE, 0.6, R, 0.12)
        cv.text(jx + 0.08, y, inw - 0.12, 0.36, [[(t, {"bold": True, "color": T1})], [(d, {"color": T2})]],
                size=7, anchor=MSO_ANCHOR.MIDDLE)
        cv.arrow(jx + inw, y + 0.18, jx + inw + 0.18, top + 0.62 + 0.6, "8C8C8C", 0.75)
    gx, gw = jx + inw + 0.18, 1.45
    fx, fy, fw, fh = cv.frame(gx, top + 0.62, gw, 1.2, CARD, 0.02)
    logo(cv, "harness", fx + 0.08, fy + 0.08, 0.2)
    cv.text(fx + 0.08, fy + 0.32, fw - 0.14, 0.8, [[("Change Guardian", {"bold": True, "size": 8, "color": T1})],
                                                   [("reads the graph, checks every rule, decides in minutes",
                                                     {"color": T2})]], size=7, after=2)
    ox = gx + gw + 0.18
    ow = jx + jw - ox
    outs = [("PASS", "moves on with its change passport", CYAN),
            ("NEEDS APPROVAL", "affected pod leads sign first", YELLOW),
            ("BLOCK + FIX", "reason given; the agent drafts the fix", MAGENTA)]
    for k, (t, d, c) in enumerate(outs):
        y = top + 0.62 + k * 0.42
        cv.arrow(gx + gw, top + 0.62 + 0.6, ox, y + 0.18, "8C8C8C", 0.75)
        cv.rect(ox, y, ow, 0.36, TINTS[c], c, 0.75, R, 0.12)
        cv.text(ox + 0.1, y, ow - 0.16, 0.36, [[(t, {"bold": True, "color": c})], [(d, {"color": T1})]], size=7,
                anchor=MSO_ANCHOR.MIDDLE)
    label(cv, jx, top + 1.98, 4, "CHECKPOINTS NO CHANGE CAN SKIP", MAGENTA)
    cps = [("1", "Story intake", "blast radius before build; affected leads added as approvers"),
           ("2", "Every commit", "diff checked against the rulebook; blocking, not advisory"),
           ("3", "Copado promotion", "impact tests per affected pod; no passport, no promotion"),
           ("4", "Every night", "orgs compared with Git and the graph; drift flagged in a day")]
    cw2 = (jw - 3 * 0.08) / 4
    for k, (n, t, d) in enumerate(cps):
        x = jx + k * (cw2 + 0.08)
        y = top + 2.22
        cv.rect(x, y, cw2, 1.18, CARD, None)
        cv.rect(x, y, cw2, 0.04, MAGENTA, None)
        badge(cv, x + 0.08, y + 0.12, n, MAGENTA, 0.24)
        cv.text(x + 0.38, y + 0.12, cw2 - 0.44, 0.24, t, size=7.4, bold=True, color=T1, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.08, y + 0.44, cw2 - 0.14, 0.7, d, size=7, color=T2)
    label(cv, jx, top + 3.5, 4, "RULEBOOK · EXAMPLES THAT BLOCK A CHANGE", MAGENTA)
    rules = ["Flow on a shared object without record-type criteria", "Shared Apex edit without tests for every pod",
             "Modify All or View All in a permission set", "Any change outside the pipeline"]
    for k, t in enumerate(rules):
        x = jx + (k % 2) * (jw / 2)
        y = top + 3.74 + (k // 2) * 0.26
        cv.rect(x, y + 0.08, 0.07, 0.07, MAGENTA, None, shape=MSO_SHAPE.OVAL)
        cv.text(x + 0.14, y, jw / 2 - 0.2, 0.24, t, size=7, color=T1, anchor=MSO_ANCHOR.MIDDLE)
    # ---------------------------------------------------------------- pilot targets
    by, bh = top + h + 0.14, 0.78
    cv.text(0.45, by, 1.3, bh, [[("PILOT TARGETS", {"bold": True, "color": T1, "track": 120})],
                                [("measured per pod", {"color": T3})]], size=7, anchor=MSO_ANCHOR.MIDDLE, after=1)
    targets = [("100%", "changes checked by the Guardian"), ("Pre-build", "cross-pod impact known"),
               ("< 24 h", "out-of-pipeline edits caught"), ("Minutes", "to assemble CAB evidence"),
               ("~1 week", "new engineer productive"), ("Every pod", "adoption measured and shared")]
    w = (12.88 - 1.8 - 5 * 0.1) / 6
    for k, (b, d) in enumerate(targets):
        x = 1.8 + k * (w + 0.1)
        cv.rect(x, by, w, bh, CARD, None)
        gbar(cv, x, by, 0.05, bh)
        cv.text(x + 0.16, by, w - 0.22, bh, [[(b, {"size": 13, "color": T1})], [(d, {"color": T2})]], size=7,
                anchor=MSO_ANCHOR.MIDDLE)
    set_notes(s, NOTES)
    return s


def main(out=OUT):
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    build(prs, layout)
    fix_footers(prs)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print("wrote", os.path.abspath(out))


if __name__ == "__main__":
    main(*sys.argv[1:])
