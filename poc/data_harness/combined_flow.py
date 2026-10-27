"""One review image: the target-state 'two access paths' architecture merged with the campaign flow.

Nothing from either view is dropped. Where they differ, the final proposed solution wins: Silver keeps every
patient (consent and contact filtering move to the per-campaign eligibility service), Golden splits into Gold A
and Gold B, the release store becomes the activation gateway, the Developer Agent is the Campaign Design Agent
and the testing loop runs under the Activation Readiness Agent.

    python3 combined_flow.py [out.pptx] [dark|light]
"""
import os
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from build_data_harness_deck import CARD, DIM, EDGE, R, SOFT, SRC, TINT, VIO, badge, logo, new_slide, path, pill
from build_mad_ig_deck import BLACK, CYAN, GREEN, MAGENTA, WHITE, YELLOW
from final_solution import gatecard, item
from footer import fix_footers
from story_slides import dark_header

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "build", "combined_flow.pptx")
CH1 = 0.96              # height of the band-1 data cards
ON_COLOR = "FFFFFE"     # white text on a brand-colour bar: stays white in the light theme


def band(cv, y, h, color, fill, title, icon, tcolor):
    cv.rect(0.45, y, 12.25, h, fill, color, 1.0, R, 0.02)
    cv.rect(0.45, y, 12.25, 0.26, color, None)
    logo(cv, icon, 0.53, y + 0.03, 0.2)
    cv.text(0.8, y, 11.75, 0.26, title, size=6.4, bold=True, color=tcolor, track=100, align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE)


def slots(x0, widths, gap=0.16):
    xs = []
    for w in widths:
        xs.append(x0)
        x0 += w + gap
    return xs


def build(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "AGENTIC HEALTHCARE CAMPAIGN PLATFORM · END TO END (FOR REVIEW)",
                "From governed data to a tested, approved campaign",
                "Read top to bottom: the data is governed in Google Cloud first, then the user prompts in Codex, then "
                "approved patients are released to Marketing Cloud.")
    X0 = 0.6
    # ---------------------------------------------------------------- 1 govern the data (top band)
    b1, h1 = 1.28, 1.32
    band(cv, b1, h1, MAGENTA, BLACK, "1  GOVERN THE DATA IN GOOGLE CLOUD", "googlecloud", ON_COLOR)
    cy, ch = b1 + 0.34, CH1
    w1 = (11.95 - 4 * 0.16) / 5
    xs1 = slots(X0, [w1] * 5)
    data = [("Sources", "Epic, athenahealth, ADSI, Strata, SFDC, SFMC engagement. Approved datasets only.", EDGE,
             CARD, "googlecloud"),
            ("DLP + Collibra", "Before ingest: the Collibra steward approves, DLP classifies, policy tags mask "
             "PHI.", MAGENTA, TINT[MAGENTA], None),
            ("Bronze", "Approved, tagged extracts with lineage and timestamps. Restricted; DLP re-scans.", "A0703C",
             CARD, None),
            ("Silver", "Validated identities, encounters, contacts, consent evidence. Keeps everyone.", SOFT, CARD,
             None)]
    for k, (t, d, ln, fl, ic) in enumerate(data):
        item(cv, xs1[k], cy, w1, ch, t, d, line=ln, fill=fl, icon=ic, n=k + 1, ncolor=MAGENTA, bsize=5.6,
             tsize=6.6)
        if k < 3:
            cv.arrow(xs1[k] + w1, cy + ch / 2, xs1[k + 1], cy + ch / 2, WHITE, 1.25)
    gh = (ch - 0.06) / 2
    gx = xs1[4]
    item(cv, gx, cy, w1, gh, "Gold B · eligibility", "Identity-linked; restricted services only", line=YELLOW,
         n=5, ncolor=YELLOW, bsize=5.2, tsize=6.2)
    item(cv, gx, cy + gh + 0.06, w1, gh, "Gold A · intelligence", "Aggregates: Heart, Maternity, General, more", line=YELLOW, n=6, ncolor=YELLOW, icon="googlebigquery", bsize=5.2, tsize=6.2)
    sxr = xs1[3] + w1
    cv.arrow(sxr, cy + gh / 2, gx, cy + gh / 2, WHITE, 1.25)
    cv.arrow(sxr, cy + gh + 0.06 + gh / 2, gx, cy + gh + 0.06 + gh / 2, WHITE, 1.25)
    # ---------------------------------------------------------------- boundary 1
    l1 = 2.75
    cv.arrow(0.45, l1, 12.7, l1, MAGENTA, 1.25, dash="dash", head=False)
    pill(cv, 0.6, l1 - 0.09, 3.0, "TRUST BOUNDARY · AGENTS: NO RAW-PHI ACCESS", MAGENTA, True, h=0.18, size=5,
         tcolor=ON_COLOR)
    # ---------------------------------------------------------------- 2 prompt, plan, build + test (centre band)
    b2, h2 = 2.9, 1.75
    band(cv, b2, h2, VIO, TINT[VIO], "2  AGENTIC HARNESS · PROMPT IN CODEX, PLAN, BUILD + TEST · MASKED DATA ONLY · RUNS IN GOOGLE CLOUD", "harness",
         ON_COLOR)
    my, mh = b2 + 0.32, 0.78
    w2 = (11.95 - 5 * 0.16) / 6
    xs2 = slots(X0, [w2] * 6)
    item(cv, xs2[0], my, w2, mh, "Prompt in Codex", "“Build a post-discharge follow-up campaign for heart "
         "patients”. Or the business console.", line=CYAN, n=7, icon="openai-white", bsize=5.3, tsize=6.4)
    item(cv, xs2[1], my, w2, mh, "Planning agents", "Orchestrator, Strategist and Analytics read the graph and "
         "Gold A; counts only.", line=VIO, n=8, ncolor=VIO, bsize=5.4, tsize=5.9)
    gatecard(cv, xs2[2], my, w2, mh, "G1", "Confirm brief", "People only, no agent: marketer + clinical approve "
             "the brief: size, channel, content plan.", n=9, bsize=5.2)
    item(cv, xs2[3], my, w2, mh, "Campaign Design Agent", "Invokes Journey Builder (native); builds the journey and "
         "data extension, paused; seeds risks.", line=VIO, n=10, ncolor=VIO, bsize=5.2, tsize=6.2)
    tx = xs2[4]
    cv.rect(tx, my, w2, mh, "0B2418", GREEN, 1.0, R, 0.07)
    badge(cv, tx + 0.08, my + 0.06, 11, GREEN, 0.22)
    cv.text(tx + 0.36, my + 0.04, w2 - 0.4, 0.22, "Readiness · testing", size=6.2, bold=True, color=GREEN,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(tx + 0.08, my + 0.28, w2 - 0.14, mh - 0.3, [[("Test contacts only: consent, STOP, links, hours. ",
                                                         {}),
                                                        ("Fail → its cause: journey 10, rule 8 or the data.",
                                                         {"bold": True, "color": MAGENTA})]], size=5.0, color=SOFT)
    gatecard(cv, xs2[5], my, w2, mh, "G3", "Approve journey", "Clinical + compliance, when every test passes. "
             "The journey then waits, active.", n=12, bsize=5.2)
    for k in range(5):
        cv.arrow(xs2[k] + w2, my + mh / 2, xs2[k + 1], my + mh / 2, WHITE, 1.25)
    # failed tests go back to their cause, under the row: the journey (10), the rule (8) or the data (band 1,
    # up the left margin)
    lp = my + mh + 0.05
    cv.arrow(tx + w2 / 2, my + mh, tx + w2 / 2, lp, MAGENTA, 1.5, head=False)
    cv.arrow(tx + w2 / 2, lp, 0.5, lp, MAGENTA, 1.5, head=False)
    for k in (3, 1):                                       # right edge: clear of the graph and Tableau arrows
        cv.arrow(xs2[k] + w2 - 0.3, lp, xs2[k] + w2 - 0.3, my + mh, MAGENTA, 1.5)
    path(cv, [(0.5, lp), (0.5, cy + ch - 0.18), (X0, cy + ch - 0.18)], MAGENTA, 1.5)
    # context + oversight row
    sy, sh = my + mh + 0.12, 0.47
    cv.text(xs2[0], sy, w2, sh, [[("CONTEXT + OVERSIGHT", {"bold": True, "color": DIM, "track": 100})],
                                 [("masked data only", {"color": DIM, "size": 5})]], size=5.4,
            anchor=MSO_ANCHOR.MIDDLE, after=1)
    item(cv, xs2[1], sy, w2, sh, "Knowledge graph", "Metadata, rules, lineage; from 6.", line=VIO,
         icon="googlecloudspanner", bsize=5.0, tsize=5.8)
    item(cv, xs2[2], sy, w2, sh, "Tableau", "Masked dashboards from 6.", icon="tableau", bsize=5.0, tsize=5.8)
    item(cv, xs2[3], sy, w2, sh, "Compliance Guardian", "Explains exclusions for 14.", line=VIO, bsize=5.0,
         tsize=5.8)
    item(cv, xs2[4], sy, w2, sh, "Measurement Agent", "Reads results from 20.", line=VIO, bsize=5.0, tsize=5.8)
    gatecard(cv, xs2[5], sy, w2, sh, "G1-3", "People approve", "Approvals go to the gateway (15).", bsize=5.0)
    cv.arrow(xs2[1] + w2 / 2, sy, xs2[1] + w2 / 2, my + mh, VIO, 1.25)
    cv.arrow(xs2[2] + w2 / 2, sy, xs2[2] + w2 / 2, my + mh, "E97627", 1.25)
    # Gold A -> agents (8), across boundary 1
    ga_y = cy + gh + 0.06 + gh
    ch1 = 2.68
    path(cv, [(gx + w1 / 2, ga_y), (gx + w1 / 2, ch1), (xs2[1] + w2 / 2 + 0.3, ch1), (xs2[1] + w2 / 2 + 0.3, my)],
         YELLOW, 1.5)
    # ---------------------------------------------------------------- boundary 2
    l2 = 4.74
    cv.arrow(0.45, l2, 12.7, l2, MAGENTA, 1.25, dash="dash", head=False)
    pill(cv, 5.0, l2 - 0.09, 3.0, "RECOMMENDATION ≠ RELEASE", MAGENTA, True, h=0.18, size=5, tcolor=ON_COLOR)
    # ---------------------------------------------------------------- 3 release + send (bottom band)
    b3, h3 = 4.95, 1.27
    band(cv, b3, h3, YELLOW, BLACK, "3  RELEASE + SEND · ONLY AFTER G2 + G3", "salesforce", BLACK)
    ry, rh = b3 + 0.32, 0.88
    widths = [1.7, 1.6, 2.0, 1.45, 1.45, 1.45, 1.34]
    xs3 = slots(X0, widths)
    item(cv, xs3[0], ry, widths[0], rh, "Eligibility service", "Gold B + approved rules: no consent, no phone or "
         "email, STOP, deceased, duplicates out, with reasons.", line=MAGENTA, n=13, ncolor=MAGENTA, bsize=5.0,
         tsize=6.2)
    gatecard(cv, xs3[1], ry, widths[1], rh, "G2", "Pre-send", "SME + tech lead review a masked report: counts, "
             "consent, exclusions. No PHI on screen.", n=14, bsize=5.0)
    fx, fy, fw, fh = cv.frame(xs3[2], ry, widths[2], rh, BLACK, 0.04)
    badge(cv, fx + 0.05, fy + 0.04, 15, YELLOW, 0.2)
    cv.text(fx + 0.3, fy + 0.01, fw - 0.34, 0.36, "Policy engine + activation gateway", size=6.2, bold=True,
            color=WHITE, line=0.9)
    cv.text(fx + 0.08, fy + 0.38, fw - 0.14, fh - 0.39, "Purpose, consent, suppression, minimum fields, human "
            "authorization. Compliance Passport; time-limited.", size=5.0, color=SOFT)
    ph = (rh - 0.06) / 2
    item(cv, xs3[3], ry, widths[3], ph, "A · via Data Cloud", None, line=CYAN, n=16, icon="salesforce", tsize=5.6)
    item(cv, xs3[3], ry + ph + 0.06, widths[3], ph, "B · direct pull", None, line=CYAN, dash="dash", tsize=5.6)
    item(cv, xs3[4], ry, widths[4], rh, "Data extension", "Approved records, this run only.", n=17,
         icon="salesforce", bsize=5.2, tsize=5.8)
    item(cv, xs3[5], ry, widths[5], rh, "Journey starts", "Built in 10, tested in 11, approved in 12.", n=18,
         bsize=5.2, tsize=5.8)
    item(cv, xs3[6], ry, widths[6], ph, "Send-time check", None, line=MAGENTA, fill=TINT[MAGENTA], n=19,
         ncolor=MAGENTA, tsize=5.4)
    item(cv, xs3[6], ry + ph + 0.06, widths[6], ph, "Measurement", None, line=MAGENTA, n=20, ncolor=MAGENTA,
         tsize=5.4)
    cv.arrow(xs3[6] + widths[6] / 2, ry + ph, xs3[6] + widths[6] / 2, ry + ph + 0.06, WHITE, 1.0)
    mid3 = ry + rh / 2
    for k in (0, 1):
        cv.arrow(xs3[k] + widths[k], mid3, xs3[k + 1], mid3, WHITE, 1.25)
    for yy in (ry + ph / 2, ry + ph + 0.06 + ph / 2):
        cv.arrow(xs3[2] + widths[2], yy, xs3[3], yy, WHITE, 1.0)
        cv.arrow(xs3[3] + widths[3], yy, xs3[4], yy, WHITE, 1.0)
    cv.arrow(xs3[4] + widths[4], mid3, xs3[5], mid3, WHITE, 1.25)
    cv.arrow(xs3[5] + widths[5], ry + ph / 2, xs3[6], ry + ph / 2, WHITE, 1.25)
    # Gold B -> eligibility: down the right margin, across boundary 2
    gb_y = cy + gh / 2
    ch2 = 4.86
    path(cv, [(gx + w1, gb_y), (12.82, gb_y), (12.82, ch2), (xs3[0] + 0.5, ch2), (xs3[0] + 0.5, ry)], MAGENTA,
         1.5)
    # engagement back: measurement -> sources, round the left margin
    fb = 6.27
    path(cv, [(xs3[6] + widths[6] / 2, ry + rh), (xs3[6] + widths[6] / 2, fb), (0.33, fb), (0.33, cy + ch / 2),
              (X0, cy + ch / 2)], GREEN, 1.25)
    pill(cv, 8.6, fb - 0.07, 2.4, "ENGAGEMENT BACK → NEXT RUN", GREEN, True, h=0.14, size=4.6)
    # ---------------------------------------------------------------- safeguards
    sgy, sgh = 6.36, 0.5
    cv.text(0.45, sgy, 1.25, sgh, [[("SAFEGUARDS", {"bold": True, "color": MAGENTA, "track": 120})],
                                   [("no wrong or no-consent patient", {"color": DIM, "size": 5})]], size=6.2,
            anchor=MSO_ANCHOR.MIDDLE, after=1)
    guards = ["Build and test before any data lands", "G2: masked pre-send report",
              "Minimum fields, only after G2 + G3", "Gateway release; agents hold no raw PHI",
              "Collibra steward approves masking", "Retry cap: 5, then an owner", "Send-time re-checks + feedback"]
    gw_ = (12.7 - 1.75 - 6 * 0.08) / 7
    for k, t in enumerate(guards):
        x = 1.75 + k * (gw_ + 0.08)
        cv.rect(x, sgy, gw_, sgh, CARD, None)
        cv.rect(x, sgy, 0.05, sgh, MAGENTA, None)
        cv.text(x + 0.12, sgy, gw_ - 0.16, sgh, t, size=5.8, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    return s


# light theme: same slide, recoloured for a white background
LIGHT_FILL = {BLACK: "FFFFFF", CARD: "F7F7F7", "1E1E1E": "EFEFEF", "2A2A2A": "E6E6E6", "06262B": "E0FAFE",
              "2B2508": "FFF7CC", "2B0716": "FFE6F0", "1A1630": "ECE9F5", "0E2A08": "EAFBE6", "0B2418": "EAFBE6",
              "0A1A24": "E6EEF2"}
LIGHT_LINE = {EDGE: "C8C8C8", SOFT: "8C8C8C", WHITE: "404040", BLACK: "404040"}
LIGHT_TEXT = {WHITE: "1A1A1A", SOFT: "4D4D4D", DIM: "6B6B6B", YELLOW: "9A7B00", GREEN: "2E8F1C", CYAN: "0089A0"}
LIGHT_LOGOS = {"openai-white": "openai", "people": "people-black", "harness": "harness-black"}


def light_logos():
    """A logo folder for white slides: white icons swapped for their dark versions."""
    import shutil
    import build_data_harness_deck as bd
    src = bd.LOGOS
    dst = os.path.join(HERE, "build", "logos_light")
    os.makedirs(dst, exist_ok=True)
    for f in os.listdir(src):
        if f.endswith(".png"):
            shutil.copy(os.path.join(src, f), os.path.join(dst, f))
    for name, dark in LIGHT_LOGOS.items():
        shutil.copy(os.path.join(src, dark + ".png"), os.path.join(dst, name + ".png"))
    bd.LOGOS = dst


def recolor_light(slide):
    A_ = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for clr in slide._element.iter("{%s}srgbClr" % A_):
        tags = set()
        node = clr.getparent()
        while node is not None:
            tags.add(node.tag.split("}")[-1])
            node = node.getparent()
        if "gradFill" in tags or "bg" in tags:
            continue
        v = clr.get("val").upper()
        if tags & {"rPr", "defRPr", "endParaRPr"}:
            table = LIGHT_TEXT
        elif "ln" in tags:
            table = LIGHT_LINE
        else:
            table = LIGHT_FILL
        if v in table:
            clr.set("val", table[v])


def main(out=OUT, theme="dark"):
    from story_slides import set_bg
    light = theme == "light"
    if light:
        light_logos()
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    s = build(prs, layout)
    if light:
        set_bg(s, "FFFFFF")
        recolor_light(s)
    for line in fix_footers(prs):
        print(line)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print("wrote", os.path.abspath(out), "-", len(prs.slides), "slide", "(%s)" % theme)


if __name__ == "__main__":
    main(*sys.argv[1:])
