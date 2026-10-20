"""One-slide review reference: the campaign data flow from the whiteboard diagram, with the gaps closed.

Kept from the diagram: Google Cloud -> bronze -> silver (consent filter) -> golden segments -> DLP; Tableau,
graph and the Analytics Agent on masked data; Codex; data extension; Journey Builder; the Developer Agent and
Testing Agent loop (test cases, failure modes catalog, classify and report).
Added to close the gaps: build and test before data lands, the G2 masked pre-send report, a minimum-field
release store, Marketing Cloud pulls, Collibra steward approval, a retry cap, send-time checks.

Not part of the deck: the review image is review/Campaign flow - for review.png; the .pptx is a build file.

    python3 campaign_flow_review.py [out.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from build_data_harness_deck import (CARD, CARD2, DIM, EDGE, R, SOFT, SRC, TINT, VIO, badge, card, logo,
                                     new_slide, path, pill)
from build_mad_ig_deck import BLACK, CYAN, GREEN, MAGENTA, WHITE, YELLOW, set_notes
from footer import fix_footers
from story_slides import dark_header

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "build", "campaign_flow_review.pptx")

NOTES = """SLIDE: From governed data to a tested, approved campaign   (about 3 minutes)

KEY MESSAGE
The data is governed and masked before any agent sees it, the journey is built and tested before any patient data exists, and people approve twice before the minimum fields are released. The data extension starts the journey by itself, so everything that must be checked is checked before records land there.

WALK IT
- Govern (left): Google Cloud sources land in the bronze layer. The silver layer removes anyone we must not contact: no consent, no phone or email, STOP, deceased, duplicates. The golden layer holds campaign-ready segments; heart, maternity and general are examples, not the full list. DLP classifies, a Collibra steward approves, and PHI is masked.
- Plan (middle): Tableau, the graph and the Analytics Agent read masked data only. 1 The marketer enters "build the campaign" in Codex. 2 Codex fetches the segment, consent and history from the graph, and the Analytics Agent sizes the audience as counts. 3 The marketer confirms the brief (G1).
- Build and test first (right): 4 The campaign build invokes Journey Builder; the Developer Agent builds the journey and data extension, paused, and seeds the known risks. 5 The Testing Agent runs test cases with test phone numbers and emails only, checks the failure modes catalog, and classifies each failed test. Failures go back to the Developer Agent, at most five times, then to a named owner. 6 Clinical and compliance approve the journey and content (G3).
- Release and send (bottom): 7 The Privacy Guard builds the masked audience. 8 The SME and tech lead approve a masked pre-send report (G2). 9 Only after G2 and G3, the minimum fields go to a time-limited release store. 10 Marketing Cloud pulls them into 11 the data extension. 12 The tested journey starts by itself. 13 Consent, STOP and frequency caps are re-checked at send. Engagement flows back to the bronze layer.
"""



def box(cv, x, y, w, h, title, body=None, line=EDGE, fill=CARD, icon=None, n=None, ncolor=CYAN, tsize=7,
        bsize=5.8, tag=None, dash=None):
    card(cv, x, y, w, h, title, body, line=line, fill=fill, icon=icon, n=n, ncolor=ncolor, tsize=tsize,
         bsize=bsize, tag=tag, dash=dash)


def panel(cv, x, y, w, h, n, title, color, icon=None):
    cv.rect(x, y, w, h, BLACK, color, 1.25, R, 0.03)
    cv.rect(x, y, w, 0.3, color, None)
    tx = x + 0.1
    if icon:
        logo(cv, icon, tx, y + 0.04, 0.22)
        tx += 0.3
    cv.text(tx, y, w - (tx - x) - 0.6, 0.3, [[("%d  " % n, {"bold": True}), (title, {"bold": True, "track": 120})]],
            size=6.8, color=BLACK if color in (YELLOW, CYAN, WHITE) else WHITE, anchor=MSO_ANCHOR.MIDDLE)


def gate(cv, x, y, w, h, g, title, body, n):
    cv.rect(x, y, w, h, TINT[YELLOW], YELLOW, 1.25, R, 0.08)
    badge(cv, x + 0.07, y + 0.07, n, YELLOW)
    logo(cv, "people", x + 0.34, y + 0.07, 0.2)
    cv.text(x + 0.6, y + 0.05, w - 1.1, 0.24, title, size=7, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    pill(cv, x + w - 0.4, y + 0.08, 0.32, g, YELLOW, True, h=0.16, size=5.2)
    cv.text(x + 0.1, y + 0.32, w - 0.2, h - 0.34, body, size=5.8, color=SOFT)


def build(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "AGENTIC CAMPAIGNS · FOR REVIEW",
                "From governed data to a tested, approved campaign",
                "Codex builds the campaign from the graph and the Analytics Agent on masked data. The journey is "
                "built, tested and approved before any patient record reaches the data extension.")
    top, th = 1.3, 3.0

    # ---------------------------------------------------------------- 1 govern (Google Cloud)
    ax, aw = 0.45, 3.55
    panel(cv, ax, top, aw, th, 1, "GOVERN · GOOGLE CLOUD", MAGENTA, "googlecloud")
    ix, iw = ax + 0.12, aw - 0.24
    box(cv, ix, top + 0.4, iw, 0.36, "Sources", None, icon="googlecloud", tsize=6.6)
    cv.text(ix + 0.95, top + 0.4, iw - 1.0, 0.36, "Epic · athenahealth · Strata · SFDC · SFMC engagement",
            size=5.6, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    box(cv, ix, top + 0.86, iw, 0.36, "Bronze", None, line="A0703C", tsize=6.6)
    cv.text(ix + 0.75, top + 0.86, iw - 0.8, 0.36, "every record useful for a campaign", size=5.8, color=SOFT,
            anchor=MSO_ANCHOR.MIDDLE)
    box(cv, ix, top + 1.32, iw, 0.5, "Silver · consent filter",
        "Drops no consent, no phone or email, STOP, deceased, duplicates.", line=SOFT, tsize=6.6, bsize=5.6)
    gy = top + 1.92
    cv.rect(ix, gy, iw, 0.5, TINT[YELLOW], YELLOW, 1.0, R, 0.08)
    cv.text(ix + 0.08, gy, 0.6, 0.5, [[("Golden", {})], [("segments, e.g.", {"size": 4.8, "bold": False,
                                                                             "color": SOFT})]],
            size=6.6, bold=True, color=YELLOW, anchor=MSO_ANCHOR.MIDDLE)
    segs = ["Heart", "Maternity", "General", "+ more"]
    sw = (iw - 0.72 - 3 * 0.06) / 4
    for k, t in enumerate(segs):
        x = ix + 0.66 + k * (sw + 0.06)
        cv.rect(x, gy + 0.09, sw, 0.32, CARD, EDGE, 0.75, R, 0.15, dash="dash" if t == "+ more" else None)
        cv.text(x, gy + 0.09, sw, 0.32, t, size=5.8, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
    dy = top + 2.52
    box(cv, ix, dy, iw, 0.4, "DLP masks PHI · Collibra steward approves", None, line=MAGENTA,
        fill=TINT[MAGENTA], tsize=6.4)
    for y0, y1 in ((top + 0.76, top + 0.86), (top + 1.22, top + 1.32), (top + 1.82, top + 1.92),
                   (top + 2.42, top + 2.52)):
        cv.arrow(ix + iw / 2, y0, ix + iw / 2, y1, SOFT, 1.0)

    # ---------------------------------------------------------------- 2 plan on masked data
    bx, bw = 4.15, 3.3
    panel(cv, bx, top, bw, th, 2, "PLAN ON MASKED DATA", VIO, "googlebigquery")
    jx = bx + 0.12
    cw3 = (bw - 0.24 - 2 * 0.08) / 3
    cons = [("tableau", "Tableau", "masked dashboards"), (None, "Graph", "segment, consent, history"),
            ("harness", "Analytics Agent", "sizes the audience: counts")]
    for k, (ic, t, d) in enumerate(cons):
        x = jx + k * (cw3 + 0.08)
        cv.rect(x, top + 0.42, cw3, 0.62, CARD, VIO if k == 2 else EDGE, 1.0, R, 0.1)
        if ic:
            logo(cv, ic, x + 0.07, top + 0.48, 0.18)
        cv.text(x + (0.3 if ic else 0.08), top + 0.46, cw3 - 0.34, 0.24, t, size=6.4, bold=True, color=WHITE,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.08, top + 0.72, cw3 - 0.12, 0.28, d, size=5.5, color=SOFT)
    badge(cv, jx + 2 * (cw3 + 0.08) + cw3 - 0.26, top + 0.78, 2, VIO, 0.2)
    # Codex
    cy = top + 1.32
    cv.rect(jx, cy, bw - 0.24, 0.62, CARD, CYAN, 1.0, R, 0.1)
    badge(cv, jx + 0.07, cy + 0.07, 1)
    logo(cv, "openai-white", jx + 0.34, cy + 0.07, 0.2)
    cv.text(jx + 0.6, cy + 0.05, 2.4, 0.24, "Build the campaign in Codex", size=7, bold=True, color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(jx + 0.1, cy + 0.32, bw - 0.44, 0.28, "“Build a post-discharge follow-up campaign for heart patients”", size=5.8,
            color=SOFT, font="Courier New")
    ax_ = jx + 2 * (cw3 + 0.08) + cw3 / 2
    gx_ = jx + (cw3 + 0.08) + cw3 / 2
    for xx_ in (gx_, ax_):                 # Codex fetches from the graph and calls the Analytics Agent
        cv.arrow(xx_ - 0.08, cy, xx_ - 0.08, top + 1.04, CYAN, 1.25)
        cv.arrow(xx_ + 0.08, top + 1.04, xx_ + 0.08, cy, VIO, 1.25)
    gate(cv, jx, top + 2.08, bw - 0.24, 0.62, "G1", "Marketer confirms the brief",
         "Audience size, channel and content plan, as counts only; agents never see a patient.", 3)
    # masked feed: DLP -> consumers
    path(cv, [(ix + iw, dy + 0.2), (bx - 0.075, dy + 0.2), (bx - 0.075, top + 0.73), (jx, top + 0.73)], MAGENTA,
         1.5)

    # ---------------------------------------------------------------- 3 build and test first
    cx_, cw_ = 7.6, 12.88 - 7.6
    panel(cv, cx_, top, cw_, th, 3, "BUILD AND TEST BEFORE ANY PATIENT DATA", CYAN, None)
    kx = cx_ + 0.12
    lw_ = 1.75
    box(cv, kx, top + 0.42, lw_, 0.5, "Campaign build", "Invokes Journey Builder", icon="salesforce", tsize=6.6)
    box(cv, kx, top + 1.12, lw_, 0.7, "Developer Agent", "Builds the journey and data extension, paused.",
        line=VIO, fill=TINT[VIO], n=4, ncolor=VIO, tsize=6.6)
    cv.arrow(kx + lw_ / 2, top + 0.92, kx + lw_ / 2, top + 1.12, SOFT, 1.0)
    # testing agent
    tx, tw = kx + lw_ + 0.5, cw_ - lw_ - 0.74
    cv.rect(tx, top + 0.4, tw, 1.88, "0B2418", GREEN, 1.0, R, 0.06)
    badge(cv, tx + 0.08, top + 0.45, 5, GREEN, 0.2)
    cv.text(tx + 0.34, top + 0.42, tw - 0.4, 0.24, "Testing Agent", size=7, bold=True, color=GREEN,
            anchor=MSO_ANCHOR.MIDDLE)
    tests = [("Test cases", "Test phone numbers and emails only; never a patient.", EDGE, CARD),
             ("Failure modes catalog", "Seeded by the developer, owned here: consent, STOP, links, quiet hours.",
              MAGENTA, TINT[MAGENTA]),
             ("Classify and report", "A failure mode for every failed test.", EDGE, CARD)]
    yy = top + 0.72
    for k, (t, d, ln, fl) in enumerate(tests):
        cv.rect(tx + 0.1, yy, tw - 0.2, 0.42, fl, ln, 0.9, R, 0.1)
        cv.text(tx + 0.18, yy + 0.02, tw - 0.36, 0.18, t, size=6.4, bold=True,
                color=MAGENTA if k == 1 else WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(tx + 0.18, yy + 0.2, tw - 0.36, 0.2, d, size=5.5, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
        if k < 2:
            cv.arrow(tx + tw / 2, yy + 0.42, tx + tw / 2, yy + 0.5, SOFT, 1.0)
        yy += 0.5
    # developer -> testing, seeds risks, failed tests back
    cv.arrow(kx + lw_, top + 1.3, tx, top + 1.3, SOFT, 1.25)
    cv.arrow(kx + lw_, top + 1.56, tx + 0.1, top + 1.56, SOFT, 1.0, dash="dash")
    cv.text(kx + lw_ + 0.02, top + 1.6, 0.48, 0.2, "seeds known risks", size=4.6, color=DIM)
    path(cv, [(tx + 0.1, top + 1.93), (kx + 0.3, top + 1.93), (kx + 0.3, top + 1.82)], MAGENTA, 1.5)
    cv.text(kx + 0.38, top + 1.97, 1.5, 0.2, "failed tests + failure mode", size=5.3, color=MAGENTA,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(kx + 0.38, top + 2.16, 1.3, 0.2, "max 5, then an owner", size=5.3, bold=True, color=MAGENTA,
            anchor=MSO_ANCHOR.MIDDLE)
    gate(cv, kx, top + 2.38, cw_ - 0.24, 0.54, "G3", "Clinical + compliance approve the journey and content",
         "Only when every test passes. The journey then waits, active, for entries.", 6)
    # G1 -> campaign build
    path(cv, [(jx + bw - 0.24, top + 2.39), (cx_ - 0.075, top + 2.39), (cx_ - 0.075, top + 0.67), (kx, top + 0.67)],
         YELLOW, 1.5)

    # ---------------------------------------------------------------- 4 release and send
    ry, rh = 4.45, 1.45
    cv.rect(0.45, ry, 12.43, rh, BLACK, YELLOW, 1.25, R, 0.03)
    cv.text(8.2, ry + 0.04, 4.6, 0.22, [[("4  RELEASE AND SEND", {"bold": True, "track": 120}),
                                         ("   only after G2 + G3", {"color": SOFT})]],
            size=6.8, color=YELLOW, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    steps = [("Privacy Guard", "Masked audience: tokens and counts per run.", 7, "googlebigquery", EDGE, CARD,
              None),
             ("Pre-send report", "SME + tech lead approve counts, consent, exclusions. No PHI on screen.", 8,
              "tableau", YELLOW, TINT[YELLOW], 2),
             ("Release store", "Minimum fields only, time-limited, read-only API; no agent access.", 9,
              "googlecloudspanner", YELLOW, BLACK, 3),
             ("Marketing Cloud pulls", "Automation Studio pulls the run; nothing is pushed by Codex.", 10,
              "salesforce", CYAN, TINT[CYAN], 4),
             ("Data extension", "Only approved records, for this run.", 11, "salesforce", EDGE, CARD, None),
             ("Journey starts itself", "The journey built and tested in step 3.", 12, "salesforce", EDGE, CARD,
              None),
             ("Send-time check", "Consent, STOP and frequency caps re-checked; then SMS or email.", 13, None,
              MAGENTA, TINT[MAGENTA], 7)]
    n_ = len(steps)
    gap = 0.14
    sw_ = (12.43 - 0.24 - (n_ - 1) * gap) / n_
    sy = ry + 0.36
    sh = 0.86
    xs = []
    for k, (t, d, n, ic, ln, fl, fx) in enumerate(steps):
        x = 0.57 + k * (sw_ + gap)
        xs.append(x)
        cv.rect(x, sy, sw_, sh, fl, ln, 1.25 if ln != EDGE else 0.9, R, 0.08)
        badge(cv, x + 0.06, sy + 0.06, n, YELLOW if ln == YELLOW else CYAN, 0.2)
        tx0 = x + 0.3
        if ic:
            logo(cv, ic, tx0, sy + 0.06, 0.18)
            tx0 += 0.23
        cv.text(tx0, sy + 0.04, x + sw_ - tx0 - 0.04, 0.24, t, size=6.2, bold=True, color=WHITE,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.08, sy + 0.32, sw_ - 0.14, 0.38, d, size=5.4, color=SOFT)
        if k < n_ - 1:
            cv.arrow(x + sw_, sy + sh / 2, x + sw_ + gap, sy + sh / 2, WHITE, 1.25)
    pill(cv, xs[1] + 0.08, sy + sh - 0.2, 0.32, "G2", YELLOW, True, h=0.16, size=5.2)
    # G3 -> release starts the Privacy Guard step
    path(cv, [(kx + 0.5, top + th), (kx + 0.5, ry + 0.2), (xs[0] + sw_ / 2, ry + 0.2), (xs[0] + sw_ / 2, sy)],
         YELLOW, 1.5)
    # engagement back to bronze
    ey = ry + rh - 0.12
    path(cv, [(xs[-1] + sw_ / 2, sy + sh), (xs[-1] + sw_ / 2, ey), (0.3, ey), (0.3, top + 1.04), (ix, top + 1.04)],
         GREEN, 1.25)
    pill(cv, 5.4, ey - 0.08, 2.5, "ENGAGEMENT BACK TO BRONZE → NEXT RUN", GREEN, True, h=0.16, size=5)

    # ---------------------------------------------------------------- what changed
    fy = 6.05
    cv.text(0.45, fy, 1.3, 0.62, [[("SAFEGUARDS", {"bold": True, "color": MAGENTA, "track": 120})],
                                  [("so no wrong or no-consent patient is sent to", {"color": DIM, "size": 5.4})]],
            size=6.4, anchor=MSO_ANCHOR.MIDDLE, after=1)
    guards = ["Build and test before any data lands", "G2: masked pre-send report",
             "Minimum fields released only after G2 + G3", "Marketing Cloud pulls; agents hold no PHI",
             "Collibra steward approves masking", "Retry cap: 5, then an owner",
             "Send-time re-checks + feedback loop"]
    fw = (12.88 - 1.8 - 6 * 0.08) / 7
    for k, t in enumerate(guards):
        x = 1.8 + k * (fw + 0.08)
        cv.rect(x, fy, fw, 0.62, CARD, None)
        cv.rect(x, fy, 0.05, 0.62, MAGENTA, None)
        cv.text(x + 0.14, fy, fw - 0.2, 0.62, t, size=6.4, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    set_notes(s, NOTES)
    return s


def main(out=OUT):
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    build(prs, layout)
    for line in fix_footers(prs):
        print(line)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print("wrote", os.path.abspath(out), "-", len(prs.slides), "slide")


if __name__ == "__main__":
    main(*sys.argv[1:])
