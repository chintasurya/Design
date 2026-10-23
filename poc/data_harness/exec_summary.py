"""Two-slide executive summary for the Ascension Agentic Factory, in the IG proposal style: white page, spaced
section headings, prose and bullets; slide 2 ends in a black timeline band.

1 Our understanding · Results & value · Why Insight Global
2 Solution overview · Key considerations · Solution timeline (the build phases of the full deck)

    python3 exec_summary.py
"""
import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from build_data_harness_deck import SRC
from build_mad_ig_deck import CYAN, Canvas, set_notes
from footer import fix_footers
from sierra_voice import gbar
from story_slides import set_bg

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Ascension Agentic Factory - Executive Summary (2 slides).pptx")
INK, SOFT_INK, RULE = "1A1A1A", "4D4D4D", "BFBFBF"
BODY = 12


def heading(cv, x, y, w, t, color=INK):
    cv.text(x, y, w, 0.3, t, size=13, bold=True, color=color, track=300, anchor=MSO_ANCHOR.MIDDLE)


def bullets(cv, x, y, w, h, items, size=BODY, color=INK, gap=4):
    tb = cv.text(x, y, w, h, list(items), size=size, color=color, after=gap)
    for p in tb.text_frame.paragraphs:
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", "228600")
        ppr.set("indent", "-228600")
        bu = ppr.makeelement(qn("a:buChar"), {"char": "•"})
        ppr.append(bu)
    return tb


def page_head(cv):
    cv.text(0.45, 0.2, 6, 0.2, "EXECUTIVE SUMMARY", size=8, bold=True, color=INK, track=300,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(0.45, 0.45, 12.43, 0.7, "Executive Summary", size=36, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)


# ================================================================== 1
NOTES_1 = """SLIDE: Executive summary (1 of 2)   (about 2 minutes)

OUR UNDERSTANDING
Ascension wants faster Salesforce delivery and faster Marketing Cloud campaigns without loosening control of patient data. Today a Health Cloud change takes about two weeks and a campaign about 90 days; much of that is manual research, impact analysis, test writing, and data preparation repeated for every campaign. The Agentic Factory brings AI agents to Health Cloud delivery, Marketing Cloud campaigns and the Sierra voice agent on one governed platform. We start with one Health Cloud pod and one campaign, then add pods one at a time.

RESULTS & VALUE
About 2 days per Health Cloud change instead of about 2 weeks; about 10-12 days per campaign instead of about 90; squads of about 2.4 people instead of 5-7; about $19.9M saved over three years. All are pilot targets and planning estimates, confirmed against the Phase 0 baseline.

WHY INSIGHT GLOBAL
It runs on what Ascension already owns; agent-assisted delivery is already measured at Ascension (1-2 weeks to 2-4 days); safety is designed in; the Change Guardian protects every pod; and the rollout is measured and incremental."""


def slide_1(prs, layout):
    s = prs.slides.add_slide(layout)
    set_bg(s, "FFFFFF")
    cv = Canvas(s)
    page_head(cv)
    heading(cv, 0.45, 1.22, 8, "OUR UNDERSTANDING")
    cv.text(0.45, 1.58, 12.43, 1.5, [
        "Ascension is seeking to speed up Salesforce delivery and Marketing Cloud campaigns while keeping patient "
        "data governed. Today a Health Cloud change takes about two weeks and a campaign about 90 days, much of it "
        "manual research, impact analysis, test writing and data preparation repeated for every campaign. The "
        "Agentic Factory brings AI agents to Health Cloud delivery, Marketing Cloud campaigns and the Sierra voice "
        "agent, on one governed platform.",
        "The initial focus will be one Health Cloud pod (Network Services) and one campaign (the CHF post-discharge "
        "nudge), followed by a pod-by-pod rollout, using a measured pilot and incremental releases."],
        size=BODY, color=INK, after=10, line=1.08)
    # results & value (left) | why Insight Global (right)
    top = 3.25
    lx, lw = 0.45, 6.75
    heading(cv, lx, top, lw, "RESULTS & VALUE")
    cv.text(lx, top + 0.38, lw, 1.0, "The solution addresses slow, manual delivery, data preparation rebuilt for "
            "every campaign, and the risk of one pod's change breaking another. Success will be measured by cycle "
            "time, first-pass approval, defects after release and zero unauthorized access to patient data.",
            size=BODY, color=INK, line=1.08)
    bullets(cv, lx, top + 1.5, lw, 1.9, [
        "Health Cloud change: ~2 weeks to ~2 days, story to production",
        "Campaign launch: ~90 days to ~10-12 days",
        "Leaner squads: 5-7 people to about 2.4",
        "About $19.9M saved over three years",
        "Agents never see raw PHI; people approve every gate",
        "A reusable foundation: new pods add skills, not rebuilds"])
    # dotted divider
    cv.arrow(7.55, top + 0.05, 7.55, 6.85, RULE, 2.0, dash="sysDot", head=False)
    rx, rw = 7.9, 4.98
    heading(cv, rx, top, rw, "WHY INSIGHT GLOBAL")
    bullets(cv, rx, top + 0.42, rw, 3.0, [
        "Built on what Ascension already runs: Codex, Salesforce, Google Cloud, Collibra, Copado",
        "Agent-assisted delivery already measured at Ascension: 1-2 weeks to 2-4 days",
        "HIPAA-aligned by design: no raw PHI for agents, no production credentials",
        "Change Guardian: no pod breaks another",
        "Outcome-led pilot with measurable success criteria",
        "Incremental, pod-by-pod rollout, not a big-bang launch"], gap=5)
    set_notes(s, NOTES_1)
    return s


# ================================================================== 2
NOTES_2 = """SLIDE: Executive summary (2 of 2)   (about 2 minutes)

SOLUTION OVERVIEW
Insight Global designs and implements the Ascension Agentic Factory: seven delivery agents and seven campaign agents, an engineering knowledge graph of Ascension's Salesforce estate, and a data harness in Google Cloud that governs patient data through Bronze, Silver and Gold layers with DLP and Collibra. Agents plan, build and test on masked data; people approve at every gate; patients reach Marketing Cloud only through a policy-checked activation gateway.

KEY CONSIDERATIONS
No raw PHI for agents, and masking alone is not HIPAA compliance; people approve every gate; existing platforms and data are reused; everything is proven on synthetic data before any real patient; releases are incremental; outcomes are measured.

SOLUTION TIMELINE
The seven build phases of the full deck, in order: discovery and baseline; the data foundation; PHI governance (its exit is a hard gate: zero unauthorized access in tests); agentic intelligence; build and test on synthetic data; governed activation with the first real CHF audience; then measure against the baseline and scale pod by pod. Phases 1 and 2 run partly in parallel. Durations are set in Phase 0. The Network Services Health Cloud pilot runs alongside, on the same agents and knowledge graph."""

PHASES = [("Phase 0 · Discovery & baseline", "owners, architecture, today's measures"),
          ("Phase 1 · Data foundation", "Bronze, Silver, Gold in Google Cloud"),
          ("Phase 2 · PHI governance", "DLP, Collibra; zero unauthorized access"),
          ("Phase 3 · Agentic intelligence", "knowledge graph, agents, first brief"),
          ("Phase 4 · Build & test", "tested; no real patient yet"),
          ("Phase 5 · Governed activation", "the first real CHF audience"),
          ("Phase 6 · Measure & scale", "vs the baseline; pod by pod")]


def slide_2(prs, layout):
    s = prs.slides.add_slide(layout)
    set_bg(s, "000000")                 # black page: the footer turns white over the timeline band
    cv = Canvas(s)
    band = 4.2
    cv.rect(0, 0, 13.333, band, "FFFFFF", None)
    page_head(cv)
    lx, lw = 0.45, 6.6
    heading(cv, lx, 1.45, lw, "SOLUTION OVERVIEW")
    cv.text(lx, 1.83, lw, 2.4, "Insight Global will design and implement the Ascension Agentic Factory: seven "
            "delivery agents and seven campaign agents, an engineering knowledge graph of Ascension's Salesforce "
            "estate, and a data harness in Google Cloud that governs patient data through Bronze, Silver and Gold "
            "layers with DLP and Collibra. Agents plan, build and test on masked data; people approve at every "
            "gate; patients reach Marketing Cloud only through a policy-checked activation gateway. Delivery "
            "follows a measured pilot, then pod by pod.", size=11.5, color=INK, line=1.08)
    rx, rw = 7.55, 5.33
    heading(cv, rx, 1.45, rw, "KEY CONSIDERATIONS")
    bullets(cv, rx, 1.85, rw, 2.4, [
        "Agents never see raw PHI; masking alone is not HIPAA compliance",
        "People approve every gate: H1-H4 and CAB, G1-G3",
        "Existing platforms and data reused: Codex, Salesforce, Google Cloud, Collibra, Copado",
        "Proven on synthetic data before any real patient",
        "Incremental releases, not a big-bang launch",
        "Measurable speed, quality and safety outcomes"], size=11.5, gap=3)
    # ---------------------------------------------------------------- timeline band
    cv.text(0.45, band + 0.22, 12.4, 0.34, [[("SOLUTION TIMELINE – ", {"color": "FFFFFF"}),
                                            ("Phased pilot", {"color": CYAN})]], size=14, track=300,
            anchor=MSO_ANCHOR.MIDDLE)
    ly = band + 1.7
    gbar(cv, 0, ly, 13.333, 0.04, 0)
    n = len(PHASES)
    x0, x1 = 0.7, 12.6
    for k, (t, d) in enumerate(PHASES):
        xx = x0 + (x1 - x0) * k / (n - 1)
        up = k % 2 == 0
        dy0, dy1 = (ly - 0.36, ly - 0.08) if up else (ly + 0.12, ly + 0.4)
        cv.arrow(xx, dy0, xx, dy1, "FFFFFF", 1.5, dash="sysDot", head=False)
        cv.rect(xx - 0.08, ly - 0.06, 0.16, 0.16, "FFFFFF", None, shape=MSO_SHAPE.OVAL)
        tw = 2.05 if k >= n - 3 else 2.55         # the last labels share the right end
        tx = xx - 0.05 if k < n - 1 else xx - tw + 0.05
        ty = ly - 0.95 if up else ly + 0.45
        cv.text(tx, ty, tw, 0.55, [[(t, {"color": "FFFFFF"})], [(d, {"color": "B3B3B3", "size": 9.5})]],
                size=10.5, anchor=MSO_ANCHOR.BOTTOM if up else MSO_ANCHOR.TOP,
                align=PP_ALIGN.RIGHT if k == n - 1 else None, after=1)
    cv.text(0.45, band + 0.58, 12.4, 0.22, "Phases 1 and 2 run partly in parallel. Durations are set in Phase 0. The "
            "Network Services Health Cloud pilot runs alongside, on the same agents and knowledge graph.", size=8.5,
            color="8C8C8C", anchor=MSO_ANCHOR.MIDDLE)
    set_notes(s, NOTES_2)
    return s


def main():
    import muted_theme as mt
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    slide_1(prs, layout)
    slide_2(prs, layout)
    for sl in prs.slides:
        mt.min_font(sl, 7.0)
    for line in fix_footers(prs):
        print(line)
    prs.save(OUT)
    print("wrote", OUT, "-", len(prs.slides), "slides")


if __name__ == "__main__":
    main()
