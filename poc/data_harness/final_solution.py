"""Final proposed solution, as review views (images for approval, not the deck).

Governed Healthcare Audience Intelligence & Agentic Activation Platform (proposed target state).

1  Solution architecture: two access paths (analytical for agents, restricted for patient activation)
2  How it works: phases A-E across people, agents, non-agent services and outputs
3  PHI governance: today's DLP -> Collibra -> policy flow, masked vs de-identified, the ten-layer control model
4  Knowledge graph (CHF context) and reuse of the MAD harness building blocks
5  Roadmap, success metrics and architecture decisions

    python3 final_solution.py [out.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from build_data_harness_deck import (CARD, CARD2, DIM, EDGE, R, SOFT, SRC, TINT, VIO, badge, card, logo,
                                     new_slide, path, pill)
from build_mad_ig_deck import BLACK, CYAN, GREEN, MAGENTA, WHITE, YELLOW
from footer import fix_footers
from story_slides import dark_header

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "build", "final_solution.pptx")
ORANGE = "FA8B35"
EYEBROW = "AGENTIC HEALTHCARE CAMPAIGN PLATFORM · "
AGENTS = ["Orchestrator", "Campaign Strategist", "Audience Analytics", "Compliance Guardian", "Campaign Design",
          "Activation Readiness", "Measurement"]


# ------------------------------------------------------------------ helpers
def side(cv, x, y, h, t, color, icon=None, tcolor=WHITE):
    cv.rect(x, y, 0.3, h, color, None)
    off = 0.0
    if icon:
        logo(cv, icon, x + 0.04, y + 0.07, 0.22)
        off = 0.3
    tb = cv.text(x + 0.15 - (h - off) / 2, y + off + (h - off) / 2 - 0.14, h - off, 0.28, t, size=6, bold=True,
                 color=tcolor, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, track=120)
    tb.rotation = 270


def gatecard(cv, x, y, w, h, g, title, body, n=None, bsize=5.8):
    cv.rect(x, y, w, h, TINT[YELLOW], YELLOW, 1.25, R, 0.08)
    tx = x + 0.08
    if n is not None:
        badge(cv, tx, y + 0.07, n, YELLOW, 0.2)
        tx += 0.25
    logo(cv, "people", tx, y + 0.07, 0.18)
    tx += 0.23
    cv.text(tx, y + 0.04, x + w - tx - 0.46, 0.24, title, size=6.8, bold=True, color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE)
    pill(cv, x + w - 0.42, y + 0.08, 0.34, g, YELLOW, True, h=0.15, size=5)
    cv.text(x + 0.08, y + 0.32, w - 0.14, h - 0.34, body, size=bsize, color=SOFT)


def item(cv, x, y, w, h, title, body, line=EDGE, fill=CARD, n=None, ncolor=CYAN, icon=None, tsize=6.8,
         bsize=6.0, tag=None, dash=None):
    card(cv, x, y, w, h, title, body, line=line, fill=fill, n=n, ncolor=ncolor, icon=icon, tsize=tsize,
         bsize=bsize, tag=tag, dash=dash)


def cells(cv, y, h, items, colors, x0=0.45, x1=12.88, tsize=7.4, size=6):
    w = (x1 - x0 - (len(items) - 1) * 0.12) / len(items)
    for k, (t, d) in enumerate(items):
        x = x0 + k * (w + 0.12)
        cv.rect(x, y, w, h, CARD, None)
        cv.rect(x, y, 0.05, h, colors[k % len(colors)], None)
        cv.text(x + 0.16, y, w - 0.24, h, [[(t, {"bold": True, "size": tsize, "color": WHITE})],
                                           [(d, {"color": SOFT})]], size=size, anchor=MSO_ANCHOR.MIDDLE, after=1)


def table(cv, x, y, w, title, color, hdr, rows, colw, rh=0.29, size=5.9, first_bold=True):
    cv.text(x, y, w, 0.22, title, size=6.4, bold=True, color=color, track=120, anchor=MSO_ANCHOR.MIDDLE)
    xs = [x + 0.08]
    for cw in colw[:-1]:
        xs.append(xs[-1] + cw)
    for cx, cw, t in zip(xs, colw, hdr):
        cv.text(cx, y + 0.24, cw, 0.18, t, size=5.4, bold=True, color=DIM)
    for k, row in enumerate(rows):
        yy = y + 0.44 + k * rh
        cv.rect(x, yy, w, rh - 0.03, CARD if k % 2 == 0 else CARD2, None)
        for j, (cx, cw, t) in enumerate(zip(xs, colw, row)):
            cv.text(cx, yy, cw - 0.06, rh - 0.03, t, size=size, bold=first_bold and j == 0,
                    color=WHITE if j == 0 else SOFT, anchor=MSO_ANCHOR.MIDDLE)
    return y + 0.44 + len(rows) * rh


# ================================================================== 1 solution architecture
def architecture(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, EYEBROW + "TARGET STATE",
                "Two access paths: analytical for agents, restricted for patient activation",
                "Proposed target state, not deployed. Agents work with masked, de-identified or aggregated "
                "information; only authorized services touch identifiable data, for an approved activity.")
    zy, zh = 1.3, 1.5
    cv.rect(0.45, zy, 12.43, zh, TINT[VIO], VIO, 1.0, R, 0.03)
    side(cv, 0.45, zy, zh, "ANALYTICAL", VIO, "harness")
    ty, th_ = zy + 0.1, zh - 0.2
    x = 0.88
    item(cv, x, ty, 1.55, th_, "Business interface", "Codex for delivery teams; a vendor-neutral console for "
         "business users.", line=CYAN, icon="openai-white", bsize=6.2)
    x += 1.55 + 0.12
    aw = 4.4
    cv.rect(x, ty, aw, th_, CARD, VIO, 1.0, R, 0.07)
    logo(cv, "harness", x + 0.1, ty + 0.07, 0.2)
    cv.text(x + 0.36, ty + 0.05, aw - 0.4, 0.24, "Agentic intelligence + compliance layer · 7 agents", size=6.8,
            bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    chw = (aw - 0.2 - 3 * 0.08) / 4
    for k, a in enumerate(AGENTS):
        cx_ = x + 0.1 + (k % 4) * (chw + 0.08)
        cy_ = ty + 0.38 + (k // 4) * 0.34
        cv.rect(cx_, cy_, chw, 0.27, TINT[VIO], VIO, 0.75, R, 0.3)
        cv.text(cx_, cy_, chw, 0.27, a, size=5.8, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
    cv.text(x + 0.1 + 3 * (chw + 0.08), ty + 0.72, chw, 0.27, "approved context only", size=5.2, color=DIM,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(x + 0.1, ty + 1.06, aw - 0.2, 0.2, "Propose, analyse and explain. They cannot unmask PHI or invoke "
            "patient-data exports.", size=5.8, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    x += aw + 0.12
    item(cv, x, ty, 2.05, th_, "Knowledge graph", "Metadata, rules, lineage and protected aggregate context. Not "
         "a patient store, not the system of record.", line=VIO, icon="googlecloudspanner", bsize=6.2)
    x += 2.05 + 0.12
    gold_a_x = x
    item(cv, x, ty, 1.95, th_, "Gold A · intelligence", "Aggregate cohorts, opportunity insights, engagement "
         "trends, protected features.", line=YELLOW, icon="googlebigquery", bsize=6.2)
    x += 1.95 + 0.12
    ax_ = x
    gatecard(cv, ax_, ty, 12.8 - ax_, th_, "G1-3", "People approve",
             "Product, clinical, privacy and security: purpose, audience rules and release.", bsize=6.2)
    by = 2.93
    cv.arrow(0.45, by, 12.88, by, MAGENTA, 1.25, dash="dash", head=False)
    pill(cv, 0.8, by - 0.09, 5.0, "NO RAW-PHI PERMISSION FOR AGENTS · A RECOMMENDATION IS NEVER PERMISSION TO "
         "RELEASE", MAGENTA, True, h=0.18, size=5, tcolor=WHITE)
    ry, rh = 3.06, 2.5
    cv.rect(0.45, ry, 12.43, rh, BLACK, MAGENTA, 1.0, R, 0.03)
    side(cv, 0.45, ry, rh, "RESTRICTED PATIENT PATH", MAGENTA, "googlecloud")
    cy, ch = ry + 0.38, rh - 0.5
    cols = [("Sources", ["Epic / EHR, ADSI and other approved clinical and engagement sources.",
                         [("Approved datasets only, not the whole estate.", {"color": DIM})]], "googlecloud",
             EDGE, 1.38),
            ("Bronze · governed raw", ["Approved extracts in BigQuery or Cloud Storage; lineage, classification.",
                                       [("Restricted PHI access.", {"bold": True, "color": MAGENTA})]], None,
             "A0703C", 1.4),
            ("Silver · patient foundation", ["Identities, encounters, contacts, attribution, consent evidence.",
                                             [("Patients without a phone or consent are kept.",
                                               {"color": DIM})]], None, SOFT, 1.58),
            ("Gold B · eligibility", ["Identity-linked, current eligibility inputs.",
                                      [("Eligibility and activation services only.",
                                        {"bold": True, "color": MAGENTA})]], None, YELLOW, 1.42)]
    x = 0.88
    xs = []
    for t, d, ic, ln, w in cols:
        item(cv, x, cy, w, ch, t, d, line=ln, icon=ic, bsize=6.2)
        xs.append((x, w))
        x += w + 0.15
    gx, gw = x, 2.15
    fx, fy, fw, fh = cv.frame(gx, cy, gw, ch, BLACK, 0.04)
    cv.text(fx + 0.08, fy + 0.04, fw - 0.16, 0.36, "Privacy policy engine + authorized activation gateway",
            size=6.6, bold=True, color=WHITE, line=0.9)
    for k, t in enumerate(["Approved purpose", "Patient eligibility", "Consent + restrictions", "Suppression",
                           "Minimum fields", "Human authorization"]):
        yy = fy + 0.46 + k * 0.24
        cv.rect(fx + 0.08, yy + 0.06, 0.09, 0.09, YELLOW, None, shape=MSO_SHAPE.OVAL)
        cv.text(fx + 0.24, yy, fw - 0.3, 0.22, t, size=6.4, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    x = gx + gw + 0.15
    pw_ = 1.5
    ah = (ch - 0.1) / 2
    item(cv, x, cy, pw_, ah, "Path A · Data Cloud", "Data Cloud → Marketing Cloud.", line=CYAN, icon="salesforce",
         bsize=6.0)
    item(cv, x, cy + ah + 0.1, pw_, ah, "Path B · direct", "Governed direct activation → Marketing Cloud.",
         line=CYAN, dash="dash", bsize=6.0)
    mx = x + pw_ + 0.15
    item(cv, mx, cy, 12.8 - mx, ch, "Marketing Cloud", ["Engagement or Next: native orchestration, patient "
         "communications, send-time checks.", [("Protected measurement refreshes Gold A and the graph.",
                                                {"color": DIM})]], line=EDGE, icon="salesforce", bsize=6.0)
    mid = cy + ch / 2
    for (x0, w0), (x1, _w) in zip(xs, xs[1:] + [(gx, gw)]):
        cv.arrow(x0 + w0, mid, x1, mid, WHITE, 1.25)
    for yy in (cy + ah / 2, cy + ah + 0.1 + ah / 2):
        cv.arrow(gx + gw, yy, x, yy, WHITE, 1.25)
        cv.arrow(x + pw_, yy, mx, yy, WHITE, 1.25)
    sx_ = xs[2][0] + 0.3
    hy = ry + 0.2
    path(cv, [(sx_, cy), (sx_, hy), (gold_a_x + 0.3, hy), (gold_a_x + 0.3, ty + th_)], YELLOW, 1.5)
    cv.text(sx_ + 0.08, hy + 0.01, 2.2, 0.16, "protected aggregates only", size=5.4, color=YELLOW,
            anchor=MSO_ANCHOR.MIDDLE)
    apx = ax_ + 0.4
    path(cv, [(apx, ty + th_), (apx, hy), (gx + gw - 0.3, hy), (gx + gw - 0.3, cy)], YELLOW, 1.5)
    cv.text(apx + 0.08, hy + 0.01, 2.2, 0.16, "approved purpose + rules", size=5.4, color=YELLOW,
            anchor=MSO_ANCHOR.MIDDLE)
    cells(cv, 5.72, 0.95, [("Data · prepared once", "Governed, continuously refreshed data products: Bronze, "
                            "Silver, Gold A, Gold B."),
                           ("AI · proposes", "Analyses protected insights and proposes campaigns; explains risks "
                            "and readiness."),
                           ("Governance · controls", "Purpose, eligibility, access and release, enforced by "
                            "deterministic services and people.")], [GREEN, VIO, MAGENTA])
    return s


# ================================================================== 2 how it works (phases A-E)
PHASES = [("A", "Build the data foundation", GREEN, BLACK), ("B", "Generate opportunities", VIO, WHITE),
          ("C", "Determine the eligible audience", YELLOW, BLACK), ("D", "Authorize activation", MAGENTA, WHITE),
          ("E", "Execute, measure, learn", CYAN, BLACK)]
PASSPORT = ["Approved purpose", "Audience version", "Source snapshot", "Policy version", "Field allowlist",
            "Approvals", "Destination", "Expiry", "Reconciliation"]


def how_it_works(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, EYEBROW + "HOW IT WORKS",
                "Five phases, from a reusable data foundation to measured outcomes",
                "Agents propose, analyse and explain. Non-agent services resolve identity, evaluate eligibility, "
                "enforce policy, release data and attribute outcomes.")
    x0 = 0.85
    w = (12.88 - x0 - 4 * 0.12) / 5
    X = [x0 + k * (w + 0.12) for k in range(5)]
    for k, (L, t, c, tc) in enumerate(PHASES):
        cv.rect(X[k], 1.3, w, 0.38, c, None, shape=MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON)
        cv.text(X[k] + (0.32 if k else 0.18), 1.3, w - 0.45, 0.38, [[(L + "  ", {"bold": True}), (t, {"bold": True})]], size=7,
                color=tc, anchor=MSO_ANCHOR.MIDDLE)
    lanes = [("PEOPLE", 1.78, 0.9, "2A2A2A", "people"), ("AGENTS", 2.76, 1.05, "0A4A52", "harness"),
             ("SERVICES", 3.89, 1.2, "3A0A1E", "googlecloud"), ("OUTPUT", 5.17, 0.58, "1E1E1E", None)]
    for t, y, h, c, ic in lanes:
        cv.rect(0.45, y, 12.43, h, CARD, None)
        side(cv, 0.45, y, h, t, c, ic)
    py, ph = 1.84, 0.78
    item(cv, X[0], py, w, ph, "Data stewards", "Collibra stewards approve classifications and ownership.",
         icon="people")
    gatecard(cv, X[1], py, w, ph, "G1", "Approve direction", "Product owner and clinical stakeholders approve "
             "the opportunity.", bsize=6.0)
    item(cv, X[2], py, w, ph, "Policy owners", "Eligibility rules are human-owned and approved per campaign.",
         icon="people")
    gatecard(cv, X[3], py, w, ph, "G2-3", "Authorize release", "Clinical, privacy, security and business approve "
             "on masked evidence.", bsize=6.0)
    item(cv, X[4], py, w, ph, "Marketing team", "Runs approved journeys in Marketing Cloud; reviews results.",
         icon="people")
    ay, ah = 2.82, 0.93
    item(cv, X[0], ay, w, ah, "No agents here", "Agents never read Bronze, Silver or Gold B.", line=EDGE,
         dash="dash")
    item(cv, X[1], ay, w, ah, "Orchestrator · Strategist · Analytics", "Query Gold A through aggregate APIs and "
         "the graph; estimate reach, engagement, cost.", line=VIO, tsize=6.4)
    item(cv, X[2], ay, w, ah, "Campaign Design Agent", "Requirements and platform-ready specs, reusing Marketing "
         "Cloud's native automation.", line=VIO)
    item(cv, X[3], ay, w, ah, "Compliance Guardian · Readiness", "Investigate risks, check readiness, assemble "
         "evidence. Cannot override policy or see raw PHI.", line=VIO, tsize=6.4)
    item(cv, X[4], ay, w, ah, "Measurement Agent", "Aggregated outcomes, forecasts and experiment results.",
         line=VIO)
    sy, sh = 3.95, 1.08
    item(cv, X[0], sy, w, sh, "Discover, classify, publish", "Inventory APIs, Sensitive Data Protection and "
         "Collibra; Bronze → Silver → Gold A and Gold B.", line=MAGENTA, icon="googlecloud")
    item(cv, X[1], sy, w, sh, "Aggregate APIs + graph", "Gold A only, with disclosure-controlled outputs.",
         line=MAGENTA, icon="googlebigquery")
    item(cv, X[2], sy, w, sh, "Eligibility service", "Deterministic rules on Gold B; resolves eligible patients "
         "and records why others are excluded.", line=MAGENTA)
    fx, fy, fw, fh = cv.frame(X[3], sy, w, sh, BLACK, 0.04)
    cv.text(fx + 0.08, fy + 0.03, fw - 0.16, 0.24, "Policy engine + activation gateway", size=6.8, bold=True,
            color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(fx + 0.08, fy + 0.3, fw - 0.16, fh - 0.32, "The identity service resolves identifiers only after "
            "approval; release through Path A (Data Cloud) or Path B (direct).", size=6.0, color=SOFT)
    item(cv, X[4], sy, w, sh, "Marketing Cloud + measurement", "Send-time checks; patient-level attribution "
         "stays restricted; aggregates refresh Gold A.", line=MAGENTA, icon="salesforce", tsize=6.4)
    oy, oh = 5.23, 0.46
    outs = ["Governed data products, refreshed", "Campaign Opportunity Brief", "Release candidate + exclusion "
            "reasons", "Campaign Compliance Passport", "Performance + ROI report"]
    for k, t in enumerate(outs):
        c = PHASES[k][2]
        cv.rect(X[k], oy, w, oh, TINT.get(c, CARD), c, 1.0, R, 0.2)
        cv.text(X[k], oy, w, oh, t, size=6.6, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        if k < 4:
            cv.arrow(X[k] + w, oy + oh / 2, X[k + 1], oy + oh / 2, WHITE, 1.25)
    # bottom: passport + learning
    by = 5.9
    pw = 8.6
    cv.rect(0.45, by, pw, 0.92, CARD, MAGENTA, 1.0, R, 0.05)
    cv.text(0.6, by + 0.04, pw - 0.3, 0.24, [[("CAMPAIGN COMPLIANCE PASSPORT", {"bold": True, "track": 120}),
                                              ("   versioned; checked again just before activation",
                                               {"color": SOFT})]], size=6.4, color=MAGENTA,
            anchor=MSO_ANCHOR.MIDDLE)
    chw = (pw - 0.3 - 4 * 0.08) / 5
    for k, t in enumerate(PASSPORT):
        cx_ = 0.6 + (k % 5) * (chw + 0.08)
        cy_ = by + 0.34 + (k // 5) * 0.28
        cv.rect(cx_, cy_, chw, 0.23, CARD2, EDGE, 0.6, R, 0.3)
        cv.text(cx_, cy_, chw, 0.23, t, size=5.9, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cells(cv, by, 0.92, [("Learns from evidence", "Refreshed aggregate results, policy changes and controlled "
                          "agent evaluation. Never by training on patient records.")], [CYAN], x0=0.45 + pw + 0.12)
    return s


# ================================================================== 3 PHI governance
ADDG = {"A": (0, 1, "DLP classification", "UPSTREAM"), "B": (1, 0, "Steward review", "UPSTREAM"),
        "C": (2, 1, "Reviewed classifications", "BUILT"), "G": (2, 2, "Rollout gate", "BUILT"),
        "T": (3, 2, "Data Policy Tagger", "LIVE IN DEV"), "P": (4, 3, "Policies applied", "LIVE IN DEV"),
        "V": (5, 1, "Bindings + validation", "BUILT"), "R": (4, 2, "Policy change review", "PROPOSED")}
ST = {"UPSTREAM": SOFT, "BUILT": GREEN, "LIVE IN DEV": CYAN, "IN PROGRESS": ORANGE, "PROPOSED": VIO}
CONTROLS = [("Discovery", "Continuous scanning of approved datasets, including new tables."),
            ("Classification", "Sensitive Data Protection findings plus Collibra stewardship decisions."),
            ("Protection", "Column and row access, encryption, restricted datasets, masking."),
            ("Agent isolation", "No raw-PHI permission, scoped tools, controlled aggregate outputs, prompt and "
             "output inspection."),
            ("Re-identification", "A separate token and identity service for authorized workloads only."),
            ("Consent + purpose", "Approved policies for HIPAA, organizational restrictions and communication "
             "laws."),
            ("Activation", "Dedicated service identity, destination allowlist, short-lived authorization, final "
             "checks."),
            ("Audit", "Access logs, policy versions, approval history, protected evidence, exception "
             "monitoring."),
            ("Remediation", "Findings go to authorized people; weakening a policy needs independent approval."),
            ("Multi-tenancy", "Per-customer identity, access, data isolation, policies, retention and "
             "contracts.")]


def addg_flow(cv, fx0, fx1, fy0):
    cv.text(fx0, fy0, 5, 0.22, "TODAY'S FLOW STAYS, WITH STATUS", size=6.4, bold=True, color=WHITE, track=120,
            anchor=MSO_ANCHOR.MIDDLE)
    lanes = [("REVIEW", "Collibra", "people"), ("GOVERNANCE", "datarepository", "googlecloud"),
             ("COMPUTE", "datalake", "googlecloud"), ("TARGET", "dataplace", "googlecloud")]
    lh, lg, ly0 = 0.56, 0.05, fy0 + 0.3
    nx0, pitch, nw, nh = fx0 + 1.26, 0.98, 0.9, 0.5
    for i, (role, proj, ic) in enumerate(lanes):
        y = ly0 + i * (lh + lg)
        cv.rect(fx0, y, fx1 - fx0, lh, CARD, None)
        cv.rect(fx0, y, 1.2, lh, CARD2, None)
        logo(cv, ic, fx0 + 0.06, y + (lh - 0.18) / 2, 0.18)
        cv.text(fx0 + 0.3, y, 0.88, lh, [[(role, {"bold": True, "color": WHITE, "size": 5.4})],
                                         [(proj, {"color": DIM, "size": 4.8, "font": "Courier New"})]],
                size=5.4, anchor=MSO_ANCHOR.MIDDLE)

    def box(nid):
        c, l = ADDG[nid][:2]
        return nx0 + c * pitch, ly0 + l * (lh + lg) + (lh - nh) / 2

    def rt(nid, dy=0.0):
        x, y = box(nid)
        return x + nw, y + nh / 2 + dy

    def lf(nid):
        x, y = box(nid)
        return x, y + nh / 2
    for a, b in (("A", "B"), ("B", "C"), ("C", "T"), ("G", "T"), ("T", "P"), ("P", "V")):
        (x1, y1), (x2, y2) = rt(a), lf(b)
        xm = x1 + (x2 - x1) / 2
        path(cv, [(x1, y1), (xm, y1), (xm, y2), (x2, y2)], CYAN, 1.25)
    rx_, ry_ = box("R")
    x1, y1 = rt("C", 0.12)
    path(cv, [(x1, y1), (rx_ + nw / 2, y1), (rx_ + nw / 2, ry_)], VIO, 1.25, dash="dash")
    px_, py_ = box("P")
    path(cv, [(rx_ + nw / 2, ry_ + nh), (px_ + nw / 2, py_)], VIO, 1.25, dash="dash")
    for nid, (c, l, t, st) in ADDG.items():
        x, y = box(nid)
        gate_ = nid in ("B", "G")
        cv.rect(x, y, nw, nh, TINT[YELLOW] if gate_ else BLACK,
                YELLOW if gate_ else (VIO if nid == "R" else EDGE), 1.0, R, 0.1,
                dash="dash" if nid in ("G", "R") else None)
        cv.text(x + 0.05, y + 0.03, nw - 0.08, 0.24, t, size=5.6, bold=True, color=WHITE, line=0.85)
        sw_ = min(nw - 0.1, 0.16 + 0.056 * len(st))
        cv.rect(x + 0.05, y + nh - 0.2, sw_, 0.16, None, ST[st], 0.75, R, 0.5)
        cv.text(x + 0.05, y + nh - 0.2, sw_, 0.16, st, size=6, bold=True, color=ST[st], align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
    ly = ly0 + 4 * (lh + lg) + 0.02
    cv.text(fx0, ly, fx1 - fx0, 0.2, [[("Fail-closed: ", {"bold": True, "color": VIO}),
                                       ("rejecting a classification never removes existing protection; relaxing a "
                                        "policy needs independent approval.", {})]], size=5.8, color=SOFT,
            anchor=MSO_ANCHOR.MIDDLE)


def governance(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, EYEBROW + "PHI GOVERNANCE",
                "Minimum-data agent access, not a zero-PHI promise",
                "The DLP → Collibra → policy enforcement pattern stays. Masked data can still be PHI, so agent "
                "access is minimized and tested, and ten control layers surround it.")
    addg_flow(cv, 0.45, 7.6, 1.3)
    rx, rw = 7.78, 12.88 - 7.78
    cv.rect(rx, 1.3, rw, 2.72, CARD, EDGE, 0.75, R, 0.04)
    cv.text(rx + 0.12, 1.34, rw - 0.24, 0.24, "MASKED IS NOT DE-IDENTIFIED", size=6.4, bold=True, color=YELLOW,
            track=120, anchor=MSO_ANCHOR.MIDDLE)
    colx = [rx + 0.12, rx + 1.35, rx + 3.05]
    colw = [1.2, 1.65, rw - 3.12]
    for x, w, t in zip(colx, colw, ["", "What it does", "HIPAA implication"]):
        cv.text(x, 1.62, w, 0.18, t, size=5.6, bold=True, color=DIM)
    rows = [("Masking", "Hides values from readers without access", "Access control; the data underneath is "
             "still PHI", MAGENTA),
            ("Tokenization", "Replaces identifiers; reversible through a separate map", "Can still be PHI, "
             "depending on linkability and context", ORANGE),
            ("De-identification", "Safe Harbor or Expert Determination met", "Can leave PHI status when the "
             "standard is met", GREEN)]
    for k, (a, b, c, col) in enumerate(rows):
        y = 1.84 + k * 0.56
        cv.rect(rx + 0.08, y, rw - 0.16, 0.5, CARD2, None)
        cv.rect(rx + 0.08, y, 0.05, 0.5, col, None)
        cv.text(colx[0] + 0.06, y, colw[0], 0.5, a, size=6.6, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(colx[1], y, colw[1], 0.5, b, size=5.8, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(colx[2], y, colw[2], 0.5, c, size=5.8, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(rx + 0.12, 3.52, rw - 0.24, 0.46, [[("A token, an aggregate count, a diagnosis combination or a small "
                                                 "cohort can remain sensitive. ", {"color": WHITE}),
                                                ("Each is assessed, not assumed safe.", {"bold": True,
                                                                                       "color": YELLOW})]],
            size=6, anchor=MSO_ANCHOR.MIDDLE)
    gy = 4.3
    cv.text(0.45, gy, 8, 0.22, [[("TARGET CONTROL MODEL", {"bold": True, "track": 120, "color": VIO}),
                                 ("   ten layers · proposed design, not production evidence", {"color": SOFT})]],
            size=6.4, anchor=MSO_ANCHOR.MIDDLE)
    cw = (12.43 - 4 * 0.12) / 5
    for k, (t, d) in enumerate(CONTROLS):
        x = 0.45 + (k % 5) * (cw + 0.12)
        y = gy + 0.3 + (k // 5) * 1.04
        cv.rect(x, y, cw, 0.96, CARD, VIO, 0.9, R, 0.06)
        badge(cv, x + 0.08, y + 0.08, k + 1, VIO, 0.22)
        cv.text(x + 0.38, y + 0.06, cw - 0.46, 0.26, t, size=7.2, bold=True, color=WHITE,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.1, y + 0.38, cw - 0.18, 0.56, d, size=6.1, color=SOFT)
    return s


# ================================================================== 4 knowledge graph + MAD reuse
KG = [("Business capability", "CHF readmission prevention"),
      ("Campaign definition", "Approved eligibility rules"),
      ("Source datasets", "Field classifications"),
      ("Policy requirements", "Purpose, consent and privacy"),
      ("Approved cohorts", "Aggregate audience insights"),
      ("History", "Past performance, approvals, audit evidence")]
REUSE = [("Knowledge graph", "Healthcare compliance and campaign context graph"),
         ("Cloud Run + durable workflows", "Agent execution, checkpoints, retries, human gates"),
         ("Tool and policy gateway", "Allowed agent tools and sensitive-data access"),
         ("Change Guardian", "Compliance Guardian with privacy-specific rules"),
         ("Three service identities", "Expanded: discovery, analysis, policy admin, activation, audit"),
         ("Approval service", "Clinical, privacy, security and business approvals"),
         ("Change passport", "Campaign Compliance Passport"),
         ("Agent evaluations", "PHI-exposure tests, false positives, policy accuracy, groundedness, reliability"),
         ("Audit store", "Tamper-evident compliance and campaign decision records")]


def graph_reuse(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, EYEBROW + "KNOWLEDGE GRAPH AND REUSE",
                "A campaign context graph, built on the harness we already designed",
                "The graph holds metadata, rules, lineage and aggregate context, never a patient-level table. Most "
                "platform blocks are reused from the Ascension Agentic Harness.")
    lx, lw = 0.45, 5.9
    cv.rect(lx, 1.3, lw, 5.52, CARD, VIO, 1.0, R, 0.03)
    logo(cv, "googlecloudspanner", lx + 0.12, 1.36, 0.2)
    cv.text(lx + 0.4, 1.34, lw - 0.5, 0.26, [[("EXAMPLE · CHF CAMPAIGN CONTEXT", {"bold": True, "track": 120})]],
            size=6.4, color=VIO, anchor=MSO_ANCHOR.MIDDLE)
    y = 1.7
    for k, (t, d) in enumerate(KG):
        cv.rect(lx + 0.15, y, lw - 0.3, 0.42, TINT[VIO] if k else TINT[YELLOW], YELLOW if k == 0 else VIO, 0.9, R,
                0.2)
        cv.text(lx + 0.3, y, 1.9, 0.42, t, size=7, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(lx + 2.25, y, lw - 2.5, 0.42, d, size=6.6, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
        if k < len(KG) - 1:
            cv.arrow(lx + 1.0, y + 0.42, lx + 1.0, y + 0.54, VIO, 1.0)
        y += 0.54
    y += 0.06
    cv.text(lx + 0.15, y, lw - 0.3, 0.2, "EVERY FACT CARRIES", size=5.8, bold=True, color=YELLOW, track=120)
    tags = ["provenance", "version", "owner", "classification", "freshness"]
    tw = (lw - 0.3 - 4 * 0.08) / 5
    for k, t in enumerate(tags):
        x = lx + 0.15 + k * (tw + 0.08)
        cv.rect(x, y + 0.24, tw, 0.24, CARD2, YELLOW, 0.75, R, 0.4)
        cv.text(x, y + 0.24, tw, 0.24, t, size=6, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.6
    cv.text(lx + 0.15, y, lw - 0.3, 0.2, "AGENTS ASK THE GRAPH", size=5.8, bold=True, color=CYAN, track=120)
    qs = ["Which datasets support a CHF campaign?", "Which approved rule identifies eligible patients?",
          "What consent evidence is required?", "Which past campaigns used a comparable population?"]
    for k, q in enumerate(qs):
        cv.text(lx + 0.15 + (k % 2) * (lw - 0.3) / 2, y + 0.22 + (k // 2) * 0.26, (lw - 0.3) / 2 - 0.05, 0.24,
                "› " + q, size=6, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(lx + 0.15, 6.5, lw - 0.3, 0.26, "The graph is not the legal authority or the system of record.",
            size=6.2, bold=True, color=MAGENTA, anchor=MSO_ANCHOR.MIDDLE)
    rx, rw = 6.5, 12.88 - 6.5
    end = table(cv, rx, 1.3, rw, "REUSE FROM THE AGENTIC HARNESS (MAD)", CYAN, ["Harness block", "Reused as"],
                REUSE, [2.15, rw - 2.15], rh=0.44, size=6.4)
    cv.rect(rx, end + 0.12, rw, 0.62, TINT[MAGENTA], MAGENTA, 0.9, R, 0.1)
    cv.text(rx + 0.15, end + 0.12, rw - 0.3, 0.62, [[("Reuse cuts engineering, not privacy work. ",
                                                      {"bold": True, "color": WHITE}),
                                                     ("Healthcare privacy controls need their own design and tests; "
                                                      "the MAD engineering graph is not a compliance engine.", {})]],
            size=6.4, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    return s


# ================================================================== 5 roadmap, metrics, decisions
STAGES = [("0", "Discovery", "GCP estate, ADSI, Salesforce licences, consent sources, BAAs, owners",
           "Approved architecture + baseline"),
          ("1", "Data foundation", "Bronze, Silver, Gold A / B and selected sources",
           "Reusable governed data products"),
          ("2", "PHI governance", "Discovery, classification, stewards, masking, access, security tests",
           "Verified protected access"),
          ("3", "Agentic intelligence", "Knowledge graph, Campaign Strategist, Analytics",
           "Opportunities + insights"),
          ("4", "Governed activation", "Eligibility service, policy gateway, approvals, release",
           "End-to-end CHF pilot"),
          ("5", "Scale + measure", "Monitoring, outcome feedback, templates, adapters",
           "Multi-campaign accelerator"),
          ("6", "Productization", "Tenant isolation, deployment templates, configuration, SecOps",
           "Cross-organization product")]
KPIS = [("Data preparation time per new campaign", "Reduce"),
        ("Approved request → audience ready", "Reduce"),
        ("Campaign datasets reused", "Increase"),
        ("In-scope sensitive fields classified", "Increase"),
        ("Unauthorized PHI access in security tests", "Zero"),
        ("Eligibility and consent validation errors", "Reduce"),
        ("Release reconciliation accuracy", "Improve"),
        ("Time preparing compliance evidence", "Reduce"),
        ("Agent recommendations approved", "Increase, quality held"),
        ("Campaign ROI", "Improve, where validated")]
DECISIONS = [("Processing location", "Google Cloud first, portable"),
             ("Data architecture", "Bronze → Silver → Gold A + Gold B"),
             ("Raw PHI for agents", "Not permitted"),
             ("Universal ready audience", "No: refreshed foundation, per-purpose audiences"),
             ("Knowledge graph", "Metadata, rules, lineage, aggregates"),
             ("Consent + eligibility", "Deterministic, purpose-specific, human-owned"),
             ("Marketing Cloud", "Data Cloud and approved direct activation"),
             ("Journey automation", "Salesforce native where practical"),
             ("Compliance model", "Policy-as-code, controlled release, approvals, audit"),
             ("Market", "Reusable, tenant-specific policies + connectors")]


def roadmap(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, EYEBROW + "ROADMAP",
                "Foundation first, then intelligence, then governed activation",
                "Pilot the CHF post-discharge nudge on synthetic data first. No speed or financial number is "
                "committed until today's workflow and the pilot are measured on equal scope.")
    w = (12.43 - 6 * 0.08) / 7
    cols = [GREEN, GREEN, MAGENTA, VIO, YELLOW, CYAN, CYAN]
    for k, (n, t, work, dlv) in enumerate(STAGES):
        x = 0.45 + k * (w + 0.08)
        c = cols[k]
        cv.rect(x, 1.3, w, 0.34, c, None, shape=MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON)
        cv.text(x + (0.28 if k else 0.14), 1.3, w - 0.36, 0.34, [[(n + "  ", {"bold": True}), (t, {"bold": True})]], size=6.6,
                color=BLACK if c in (GREEN, YELLOW, CYAN) else WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.rect(x, 1.7, w, 1.12, CARD, c if k == 4 else EDGE, 1.25 if k == 4 else 0.75, R, 0.06)
        cv.text(x + 0.08, 1.74, w - 0.14, 0.62, work, size=6, color=SOFT)
        cv.text(x + 0.08, 2.36, w - 0.14, 0.42, dlv, size=6.3, bold=True, color=WHITE, anchor=MSO_ANCHOR.BOTTOM)
    x4 = 0.45 + 4 * (w + 0.08)
    pill(cv, x4 + 0.1, 2.9, w - 0.2, "FIRST MILESTONE · CHF PILOT", YELLOW, True, h=0.18, size=5)
    cv.text(0.45, 2.9, x4 - 0.6, 0.2, "Prove a protected, pre-prepared foundation produces an approved CHF audience "
            "faster, with equal or stronger identity and consent controls.", size=6, color=SOFT,
            anchor=MSO_ANCHOR.MIDDLE)
    end1 = table(cv, 0.45, 3.25, 5.6, "SUCCESS METRICS", GREEN, ["Measure", "Direction"], KPIS, [3.85, 1.75],
                 rh=0.27, size=6.0, first_bold=False)
    end2 = table(cv, 6.2, 3.25, 6.68, "ARCHITECTURE DECISIONS", CYAN, ["Decision", "Recommendation"], DECISIONS,
                 [2.0, 4.68], rh=0.27, size=6.0)
    y = max(end1, end2) + 0.1
    cv.rect(0.45, y, 12.43, 6.85 - y, TINT[VIO], VIO, 1.0, R, 0.1)
    cv.text(0.6, y, 12.1, 6.85 - y, [[("Governed Healthcare Audience Intelligence & Agentic Activation Platform  ",
                                       {"bold": True, "color": WHITE}),
                                      ("HIPAA-aligned and vendor-neutral: governed data products, agents on protected "
                                       "insights, approved audiences released through auditable services.",
                                       {"color": SOFT})]], size=6.8, anchor=MSO_ANCHOR.MIDDLE)
    return s


def main(out=OUT):
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    for fn in (architecture, how_it_works, governance, graph_reuse, roadmap):
        fn(prs, layout)
    for line in fix_footers(prs):
        print(line)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print("wrote", os.path.abspath(out), "-", len(prs.slides), "slides")


if __name__ == "__main__":
    main(*sys.argv[1:])
