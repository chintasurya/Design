"""Muted slide styles for the Agentic Harness deck: neutral greys, one accent colour, thinner borders, and the
IG gradient (yellow, magenta, cyan) only on differentiators and hooks.

accent_dark(slide, (line, text))   dark slide, greys plus one accent on gates and campaign items (slide 2)
muted(slide, theme, hooks, ...)    dark or white slide, greys plus one accent, thinner lines, gradient hooks
metals(slide, cards)               bronze, silver and gold colours for the data-layer boxes
"""
import copy

from lxml import etree
from pptx.util import Emu

import build_data_harness_deck as bd
from build_mad_ig_deck import BLACK, CYAN, GREEN, MAGENTA, NAVY, WHITE, YELLOW

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
q = lambda ns, t: "{%s}%s" % (ns, t)
EMU = 914400
BRAND = {CYAN, GREEN, MAGENTA, NAVY, YELLOW, bd.VIO, "FA8B35", "A0703C", "0A4A52", "3A0A1E", "0A2A44", "2A2A2A"}
TINTS = set(bd.TINT.values()) | {"0B2418"}
ACCENT_SRC = {YELLOW, CYAN}
TH = {"dark": dict(bg="000000", card="141414", card2="1E1E1E", tint="161616", brand="2B2B2B", blue="00D6F2",
                   oval="2B2B2B", edge="333333", bline="4A4A4A", aline="00D6F2", arrow="8C8C8C", soft="6B6B6B",
                   t1="F2F2F2", t2="B3B3B3", t3="8C8C8C", tacc="00D6F2"),
      "white": dict(bg="FFFFFF", card="F7F7F7", card2="EFEFEF", tint="F7F7F7", brand="E9E9E9", blue="E0FAFE",
                    oval="E3E3E3", edge="D6D6D6", bline="BDBDBD", aline="00B8D4", arrow="595959", soft="9E9E9E",
                    t1="1A1A1A", t2="555555", t3="707070", tacc="0089A0")}
ACC = {"cyan": {"dark": dict(fill="00D6F2", oval="00D6F2", line="00D6F2", text="00D6F2"),
                "white": dict(fill="E0FAFE", oval="00D6F2", line="00B8D4", text="0089A0")},
       "yellow": {"dark": dict(fill="FFD700", oval="FFD700", line="FFD700", text="FFD700"),
                  "white": dict(fill="FFF3B0", oval="FFD700", line="C9A100", text="8A6D00")}}
METAL = {"bronze": ("CD7F32", "F8ECE0", "2A1A0C"), "silver": ("A8A9AD", "F0F1F2", "1F2124"),
         "gold": ("C9A227", "FBF4DC", "2A2310")}     # line, white-theme tint, dark-theme tint
GRAD = ('<a:gradFill xmlns:a="%s" rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="%s"/></a:gs>'
        '<a:gs pos="50000"><a:srgbClr val="%s"/></a:gs><a:gs pos="100000"><a:srgbClr val="%s"/></a:gs></a:gsLst>'
        '<a:lin ang="%d" scaled="0"/></a:gradFill>')


def grad(ang=0):
    return etree.fromstring(GRAD % (A, YELLOW, MAGENTA, CYAN, ang))


def lum(h):
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def solid(el):
    sp = el.find(q(P, "spPr"))
    if sp is None:
        return None
    sf = sp.find(q(A, "solidFill"))
    if sf is None or sf.find(q(A, "srgbClr")) is None:
        return None
    return sf.find(q(A, "srgbClr"))


def box(shape):
    return shape.left, shape.top, shape.left + shape.width, shape.top + shape.height


def muted(slide, theme, hooks=(), accent="cyan", yellow_if=None, src=None, src_accent=None, neutral_if=None):
    """src: the original colours that become the accent (default yellow and cyan); src_accent maps one of
    them to another palette; neutral_if forces a shape's brand fill to grey."""
    t = dict(TH[theme])
    a = ACC[accent][theme]
    t.update(blue=a["fill"], aline=a["line"], tacc=a["text"])
    ya = ACC["yellow"][theme]
    src = set(src or ACCENT_SRC)
    src_accent = src_accent or {}
    pal = {v: ACC[src_accent.get(v, accent)][theme] for v in src}
    accent_fills = {a["fill"], a["oval"], ya["fill"], ya["oval"]} | {p["fill"] for p in pal.values()}
    shapes = list(slide.shapes)
    # fills
    for sh in shapes:
        c = solid(sh._element)
        if c is None:
            continue
        v = c.get("val").upper()
        geom = sh._element.find(".//" + q(A, "prstGeom"))
        oval = geom is not None and geom.get("prst") == "ellipse"
        if v == BLACK:
            c.set("val", t["bg"])
        elif v == bd.CARD:
            c.set("val", t["card"])
        elif v == "1E1E1E":
            c.set("val", t["card2"])
        elif v in TINTS:
            c.set("val", t["tint"])
        elif neutral_if is not None and (v in BRAND or v in src) and neutral_if(sh):
            c.set("val", t["oval"] if oval else t["brand"])
        elif yellow_if is not None and v in src and yellow_if(sh):
            c.set("val", ya["fill"])
        elif v in src and not oval and sh.width > 6 * EMU:
            c.set("val", t["brand"])                 # full-width header bars stay neutral
        elif v in src and not oval:
            c.set("val", pal[v]["fill"])
        elif v == YELLOW and oval and YELLOW in src:
            c.set("val", pal[v]["oval"])
        elif v in BRAND or v in ACCENT_SRC:
            c.set("val", t["oval"] if oval else t["brand"])
    # lines: recolour and thin
    for sh in shapes:
        el = sh._element
        conn = el.tag == q(P, "cxnSp")
        for ln in el.iter(q(A, "ln")):
            w = ln.get("w")
            if w:
                ln.set("w", str(max(9525 if conn else 6350, int(int(w) * 0.6))))
            c = ln.find(q(A, "solidFill") + "/" + q(A, "srgbClr"))
            if c is None:
                continue
            v = c.get("val").upper()
            if v in src:
                c.set("val", pal[v]["line"])
            elif v in (WHITE, BLACK):
                c.set("val", t["arrow"])
            elif v == "E97627":                      # Tableau orange line: IG cyan
                c.set("val", t["aline"])
            elif v == bd.EDGE:
                c.set("val", t["edge"])
            elif v == bd.SOFT:
                c.set("val", t["soft"])
            elif v in BRAND:
                c.set("val", t["bline"])
    # hooks: a thin gradient frame behind the card (like the IG differentiator border), or a gradient bar
    d = 18288                                     # 0.02 in frame
    for sh in shapes:
        if solid(sh._element) is None:
            continue
        for kind, test in hooks:
            if not test(sh):
                continue
            el = sh._element
            sp = el.find(q(P, "spPr"))
            if kind == "frame":
                fr = copy.deepcopy(el)
                fsp = fr.find(q(P, "spPr"))
                xfrm = fsp.find(q(A, "xfrm"))
                off, ext = xfrm.find(q(A, "off")), xfrm.find(q(A, "ext"))
                off.set("x", str(int(off.get("x")) - d)); off.set("y", str(int(off.get("y")) - d))
                ext.set("cx", str(int(ext.get("cx")) + 2 * d)); ext.set("cy", str(int(ext.get("cy")) + 2 * d))
                for ch in list(fsp):
                    if ch.tag in (q(A, "solidFill"), q(A, "noFill"), q(A, "gradFill"), q(A, "ln")):
                        fsp.remove(ch)
                fsp.insert(list(fsp).index(fsp.find(q(A, "prstGeom"))) + 1, grad())
                ln = etree.SubElement(fsp, q(A, "ln")); etree.SubElement(ln, q(A, "noFill"))
                tb = fr.find(q(P, "txBody"))
                if tb is not None:
                    fr.remove(tb)
                el.addprevious(fr)
                for ch in list(sp):
                    if ch.tag == q(A, "ln"):
                        sp.remove(ch)
                ln = etree.SubElement(sp, q(A, "ln")); etree.SubElement(ln, q(A, "noFill"))
            else:
                for ch in list(sp):
                    if ch.tag in (q(A, "solidFill"), q(A, "noFill"), q(A, "gradFill")):
                        sp.remove(ch)
                sp.insert(list(sp).index(sp.find(q(A, "prstGeom"))) + 1, grad(5400000))
    shapes = list(slide.shapes)
    # text: contrast against whatever is underneath
    for i, sh in enumerate(shapes):
        if not sh.has_text_frame:
            continue
        cx, cy = sh.left + sh.width // 2, sh.top + sh.height // 2
        under = t["bg"]
        for prev in shapes[:i]:
            c = solid(prev._element)
            if c is None:
                continue
            x0, y0, x1, y1 = box(prev)
            if x0 <= cx <= x1 and y0 <= cy <= y1:
                under = c.get("val").upper()
        dark = lum(under) < 0.5
        on_blue = under in accent_fills
        for rc in sh._element.iter(q(A, "srgbClr")):
            par = rc.getparent().getparent()
            if par is None or par.tag.split("}")[-1] not in ("rPr", "defRPr", "endParaRPr"):
                continue
            v = rc.get("val").upper()
            if v in src and not on_blue:
                name = src_accent.get(v, accent)
                new = ACC[name]["dark"]["text"] if dark else ACC[name]["white"]["text"]
            elif v in (bd.SOFT, "4D4D4D"):
                new = TH["dark"]["t2"] if dark else TH["white"]["t2"]
            elif v == bd.DIM:
                new = TH["dark"]["t3"] if dark else TH["white"]["t3"]
            elif v == "E97627":
                continue
            else:
                new = TH["dark"]["t1"] if dark else TH["white"]["t1"]
            rc.set("val", new)


def metals(slide, cards, theme="white"):
    """cards: list of (metal, x, y, w, h) in inches. Recolour the card, and any badge inside it."""
    for sh in slide.shapes:
        c = solid(sh._element)
        sp = sh._element.find(q(P, "spPr"))
        for metal, x, y, w, h in cards:
            line, light, dark = METAL[metal]
            tint = dark if theme == "dark" else light
            X, Y, W, H = (int(v * EMU) for v in (x, y, w, h))
            if abs(sh.left - X) < 30000 and abs(sh.top - Y) < 30000 and abs(sh.width - W) < 30000:
                if c is not None:
                    c.set("val", tint)
                ln = sp.find(q(A, "ln"))
                lc = None if ln is None else ln.find(q(A, "solidFill") + "/" + q(A, "srgbClr"))
                if lc is not None:
                    lc.set("val", line)
                    ln.set("w", "15875")
            geom = sh._element.find(".//" + q(A, "prstGeom"))
            cx, cy = sh.left + sh.width // 2, sh.top + sh.height // 2
            if c is not None and geom is not None and geom.get("prst") == "ellipse" and X <= cx <= X + W                     and Y <= cy <= Y + H:
                c.set("val", line)


def near(v, target, tol=0.03):
    return abs(v / EMU - target) < tol



def accent_dark(slide, accent):
    """accent: None, or (line colour, text colour) for human gates and campaign items."""
    fill = {t: "161616" for t in TINTS}
    fill.update({c: "2B2B2B" for c in BRAND})
    line = {c: "5C5C5C" for c in BRAND}
    text = {BLACK: "FFFFFF"}
    text.update({c: "E6E6E6" for c in BRAND})
    if accent:
        for c in (YELLOW, MAGENTA):
            line[c], text[c] = accent
    for clr in slide._element.iter("{%s}srgbClr" % A):
        tags = set()
        n = clr.getparent()
        while n is not None:
            tags.add(n.tag.split("}")[-1]); n = n.getparent()
        if "bg" in tags or "gradFill" in tags:
            continue
        v = clr.get("val").upper()
        table = text if tags & {"rPr", "defRPr", "endParaRPr"} else line if "ln" in tags else fill
        if v in table:
            clr.set("val", table[v])


def min_font(slide, pt=7.0):
    """Raise every text size below `pt` to `pt`, so nothing on the slide is too small to read."""
    floor = int(pt * 100)
    for tag in ("rPr", "endParaRPr", "defRPr"):
        for r in slide._element.iter(q(A, tag)):
            sz = r.get("sz")
            if sz and int(sz) < floor:
                r.set("sz", str(floor))


# ------------------------------------------------------------------ white version of the dark muted slides
LIGHT_FILLS = {"000000": "FFFFFF", "0E0E0E": "F7F7F7", "141414": "F5F5F5", "161616": "F5F5F5", "1E1E1E": "EEEEEE",
               "262626": "ECECEC", "2B2B2B": "E6E6E6", "333333": "E0E0E0", "06262B": "E0FAFE", "0A4A52": "E0FAFE",
               "2B2508": "FFF7CC", "0B2418": "EAFBE6", "2B0716": "FFE6F0", "3A0A1E": "FFE6F0", "1A1630": "ECE9F5",
               "0A1A24": "E6EEF2", "2A1A0C": "F8ECE0", "1F2124": "F0F1F2", "2A2310": "FBF4DC"}
LIGHT_ACCENT_TEXT = {CYAN: "0089A0", YELLOW: "8A6D00", GREEN: "2E8F1C", "FF3D8B": "D6005A"}
LIGHT_ACCENT_LINE = {CYAN: "00B8D4", YELLOW: "C9A100", GREEN: "2E8F1C"}


def _sl(v):
    import colorsys
    r, g, b = (int(v[k:k + 2], 16) / 255 for k in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    return s, l


def to_light(slide):
    """Recolour a dark muted slide for a white background: light cards, dark text, accents darkened enough to
    read on white. Gradients and the slide background are left to the caller."""
    for clr in slide._element.iter(q(A, "srgbClr")):
        tags = set()
        n = clr.getparent()
        while n is not None:
            tags.add(n.tag.split("}")[-1])
            n = n.getparent()
        if "gradFill" in tags or "bg" in tags:
            continue
        v = clr.get("val").upper()
        sat, l = _sl(v)
        grey = sat < 0.15
        if tags & {"rPr", "defRPr", "endParaRPr"}:
            if v in LIGHT_ACCENT_TEXT:
                new = LIGHT_ACCENT_TEXT[v]
            elif grey and l > 0.85:
                new = "1A1A1A"
            elif grey and l > 0.6:
                new = "4D4D4D"
            elif grey and l > 0.45:
                new = "6B6B6B"
            else:
                continue
        elif "ln" in tags:
            if v in LIGHT_ACCENT_LINE:
                new = LIGHT_ACCENT_LINE[v]
            elif grey and l > 0.85:
                new = "595959"
            elif grey and l > 0.6:
                new = "8C8C8C"
            elif grey and l < 0.3:
                new = "D0D0D0"
            elif grey and l < 0.5:
                new = "BDBDBD"
            else:
                continue
        else:
            if v in LIGHT_FILLS:
                new = LIGHT_FILLS[v]
            elif l < 0.25:
                new = "F2F2F2"
            elif grey and l < 0.6:
                new = "E6E6E6"
            else:
                continue
        clr.set("val", new)


# ------------------------------------------------------------------ colourful black: IG colours only, thin lines
VIVID = {"8B7CD8": MAGENTA, GREEN: CYAN, "FA8B35": YELLOW, "E97627": CYAN, NAVY: "06262B",
         "1A1630": "2B0716", "0E2A08": "06262B", "0B2418": "06262B", "0A2A44": "06262B"}


def vivid(slide):
    """Keep the colourful design, but only in IG cyan, yellow and magenta (plus greys, metals and the gradient),
    with borders and arrows about 40% thinner."""
    for clr in slide._element.iter(q(A, "srgbClr")):
        anc = [n.tag.split("}")[-1] for n in clr.iterancestors()]
        if "gradFill" in anc or "bg" in anc:
            continue
        v = clr.get("val").upper()
        if v in VIVID:
            clr.set("val", VIVID[v])
    for sh in slide.shapes:
        el = sh._element
        conn = el.tag == q(P, "cxnSp")
        for ln in el.iter(q(A, "ln")):
            w = ln.get("w")
            if w:
                ln.set("w", str(max(9525 if conn else 6350, int(int(w) * 0.6))))

