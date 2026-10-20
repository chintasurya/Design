"""Review views after the architecture review: four slides, rendered to images for approval (not the deck).

A  Target architecture: governed foundation (Bronze, Silver, Gold A / Gold B), restricted activation gateway,
   two activation adapters (Data Cloud kept as a supported route)
B  Campaign harness (replaces deck slide 5): agents on Gold A only, policy engine, identity service,
   activation gateway with a Campaign Release Manifest
C  PHI governance (replaces deck slide 6): today's ADDG flow with statuses, masking vs tokenization vs
   de-identification, and the proposed additions (gaps A-F)
D  Gap register, Data Cloud decision and the CHF validation plan

    python3 manager_review.py [out.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from build_data_harness_deck import (CARD, CARD2, DIM, EDGE, R, SOFT, SRC, TINT, VIO, badge, card, logo,
                                     new_slide, path, pill)
from build_mad_ig_deck import BLACK, CYAN, GREEN, MAGENTA, WHITE, YELLOW
from footer import fix_footers
from story_slides import dark_header

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "build", "manager_review.pptx")
ORANGE = "FA8B35"


def side(cv, x, y, h, t, color, icon=None, tcolor=WHITE):
    cv.rect(x, y, 0.3, h, color, None)
    off = 0.0
    if icon:
        logo(cv, icon, x + 0.04, y + 0.07, 0.22)
        off = 0.3
    tb = cv.text(x + 0.15 - (h - off) / 2, y + off + (h - off) / 2 - 0.14, h - off, 0.28, t, size=6, bold=True,
                 color=tcolor, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, track=120)
    tb.rotation = 270


def gatecard(cv, x, y, w, h, g, title, body, n=None):
    cv.rect(x, y, w, h, TINT[YELLOW], YELLOW, 1.25, R, 0.08)
    tx = x + 0.08
    if n is not None:
        badge(cv, tx, y + 0.07, n, YELLOW, 0.2)
        tx += 0.25
    logo(cv, "people", tx, y + 0.07, 0.18)
    tx += 0.23
    cv.text(tx, y + 0.04, x + w - tx - 0.42, 0.24, title, size=6.6, bold=True, color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE)
    pill(cv, x + w - 0.38, y + 0.08, 0.3, g, YELLOW, True, h=0.15, size=5)
    cv.text(x + 0.08, y + 0.32, w - 0.14, h - 0.34, body, size=5.6, color=SOFT)


def item(cv, x, y, w, h, title, body, line=EDGE, fill=CARD, n=None, ncolor=CYAN, icon=None, tsize=6.6,
         bsize=5.6, tag=None, dash=None):
    card(cv, x, y, w, h, title, body, line=line, fill=fill, n=n, ncolor=ncolor, icon=icon, tsize=tsize,
         bsize=bsize, tag=tag, dash=dash)


def cells(cv, y, h, items, colors):
    w = (12.43 - (len(items) - 1) * 0.12) / len(items)
    for k, (t, d) in enumerate(items):
        x = 0.45 + k * (w + 0.12)
        c = colors[k % len(colors)]
        cv.rect(x, y, w, h, CARD, None)
        cv.rect(x, y, 0.05, h, c, None)
        cv.text(x + 0.16, y, w - 0.24, h, [[(t, {"bold": True, "size": 7.4, "color": WHITE})],
                                           [(d, {"color": SOFT})]], size=6, anchor=MSO_ANCHOR.MIDDLE, after=1)


# ================================================================== A target architecture
def architecture(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "AGENTIC CAMPAIGNS · TARGET ARCHITECTURE",
                "A governed data foundation, an agentic intelligence layer, a restricted activation gateway",
                "Agents work only on aggregated intelligence (Gold A). Patient-level data stays in the restricted "
                "zone and leaves only through a campaign-scoped gateway.")
    # intelligence zone
    zy, zh = 1.3, 1.5
    cv.rect(0.45, zy, 12.43, zh, TINT[VIO], VIO, 1.0, R, 0.03)
    side(cv, 0.45, zy, zh, "NO PHI", VIO, "harness")
    ty, th_ = zy + 0.1, zh - 0.2
    x = 0.88
    top = [("Business interface", "Codex for delivery teams; a vendor-neutral campaign console for business "
            "users.", "openai-white", CYAN, 1.85),
           ("Campaign Strategist · Analytics Agent", "Recommend from approved, aggregated insight: volume, overlap, "
            "care gaps, forecast ranges. No clinical or treatment decisions.", "harness", VIO, 2.3),
           ("Knowledge graph: metadata, not patients", "Assets, fields, classifications, rules, policies, approvals, "
            "segment definitions, aggregate outcomes.", "googlecloudspanner", VIO, 2.45),
           ("Gold A · audience intelligence", "Aggregated, masked or de-identified cohorts with disclosure control. "
            "For agents and Tableau.", "googlebigquery", YELLOW, 2.2)]
    for t, d, ic, ln, w in top:
        item(cv, x, ty, w, th_, t, d, line=ln, icon=ic, bsize=6.6, tsize=6.8)
        x += w + 0.14
    ax_ = x
    gatecard(cv, ax_, ty, 12.8 - ax_, th_, "G1-3", "People approve",
             "Clinical, privacy and product approve the campaign, its purpose and the eligibility rule. Pre-send "
             "review is masked.")
    gold_a_x = 0.88 + 1.85 + 2.3 + 2.45 + 3 * 0.14
    # trust boundary
    by = 2.93
    cv.arrow(0.45, by, 12.88, by, MAGENTA, 1.25, dash="dash", head=False)
    pill(cv, 0.8, by - 0.09, 4.5, "TRUST BOUNDARY · NO PATIENT-LEVEL DATA CROSSES UP · AGENTS CANNOT CALL THE "
         "GATEWAY", MAGENTA, True, h=0.18, size=5, tcolor=WHITE)
    # restricted zone
    ry, rh = 3.06, 2.5
    cv.rect(0.45, ry, 12.43, rh, BLACK, MAGENTA, 1.0, R, 0.03)
    side(cv, 0.45, ry, rh, "RESTRICTED PHI ZONE", MAGENTA, "googlecloud")
    cy, ch = ry + 0.38, rh - 0.5
    cols = [("Sources", ["EHR (Epic, athenahealth), ADSI, Strata, SFDC, SFMC engagement.",
                         [("Owners keep authority; approved ingestion only.", {"color": DIM})]], "googlecloud",
             EDGE, 1.4),
            ("Bronze · source data", ["Authorized records with provenance, quality and lineage.",
                                      [("No agent access.", {"bold": True, "color": MAGENTA})]], None, "A0703C",
             1.38),
            ("Silver · standardized", ["Validated identities, encounters, preferences, consent evidence, "
                                       "attribution, quality checks.",
                                       [("Ineligible patients are kept; eligibility is per campaign.",
                                         {"color": DIM})]], None, SOFT, 1.62),
            ("Gold B · eligibility", ["Identity-linked facts, consent and eligibility inputs.",
                                      [("Eligibility service only.", {"bold": True, "color": MAGENTA})]], None,
             YELLOW, 1.45)]
    x = 0.88
    xs = []
    for t, d, ic, ln, w in cols:
        item(cv, x, cy, w, ch, t, d, line=ln, icon=ic, bsize=6.5, tsize=6.8)
        xs.append((x, w))
        x += w + 0.16
    # gateway
    gx, gw = x, 2.1
    fx, fy, fw, fh = cv.frame(gx, cy, gw, ch, BLACK, 0.04)
    cv.text(fx + 0.08, fy + 0.04, fw - 0.16, 0.36, "Campaign eligibility + activation gateway", size=6.8,
            bold=True, color=WHITE, line=0.9)
    steps = ["Approved rule → fresh audience", "Identity + encounter links valid", "Live consent + suppression",
             "Minimum fields, short-lived", "Release manifest + audit"]
    for k, t in enumerate(steps):
        yy = fy + 0.46 + k * 0.29
        cv.rect(fx + 0.08, yy + 0.06, 0.1, 0.1, YELLOW, None, shape=MSO_SHAPE.OVAL)
        cv.text(fx + 0.24, yy, fw - 0.3, 0.24, t, size=6.6, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    x = gx + gw + 0.16
    # adapters
    aw = 1.55
    ah = (ch - 0.1) / 2
    item(cv, x, cy, aw, ah, "Adapter A · Data Cloud", "Current VBC route: identity, segments, activation.",
         line=CYAN, icon="salesforce", bsize=6.2)
    item(cv, x, cy + ah + 0.1, aw, ah, "Adapter B · direct", "Only after parity and security validation.",
         line=CYAN, dash="dash", bsize=6.2)
    mx = x + aw + 0.16
    mw = 12.8 - mx
    item(cv, mx, cy, mw, ch, "Marketing Cloud", ["Engagement or Next: journeys, tracking, responses.",
                                                  [("Opt-outs and responses flow back to Silver.",
                                                    {"color": DIM})]], line=EDGE, icon="salesforce", bsize=6.4)
    mid = cy + ch / 2
    for (x0, w0), (x1, _w) in zip(xs, xs[1:] + [(gx, gw)]):
        cv.arrow(x0 + w0, mid, x1, mid, WHITE, 1.25)
    cv.arrow(gx + gw, cy + ah / 2, x, cy + ah / 2, WHITE, 1.25)
    cv.arrow(gx + gw, cy + ah + 0.1 + ah / 2, x, cy + ah + 0.1 + ah / 2, WHITE, 1.25)
    cv.arrow(x + aw, cy + ah / 2, mx, cy + ah / 2, WHITE, 1.25)
    cv.arrow(x + aw, cy + ah + 0.1 + ah / 2, mx, cy + ah + 0.1 + ah / 2, WHITE, 1.25)
    # silver -> gold A (aggregates only) and approvals -> gateway (rule only)
    sx_ = xs[2][0] + 0.3
    hy = ry + 0.2
    path(cv, [(sx_, cy), (sx_, hy), (gold_a_x + 0.3, hy), (gold_a_x + 0.3, ty + th_)], YELLOW, 1.5)
    cv.text(sx_ + 0.08, hy + 0.01, 2.2, 0.16, "aggregates only, disclosure-controlled", size=5.3, color=YELLOW,
            anchor=MSO_ANCHOR.MIDDLE)
    apx = ax_ + 0.4
    path(cv, [(apx, ty + th_), (apx, hy), (gx + gw - 0.3, hy), (gx + gw - 0.3, cy)], YELLOW, 1.5)
    cv.text(apx + 0.08, hy + 0.01, 2.2, 0.16, "approved campaign + rule, no patient data", size=5.3,
            color=YELLOW, anchor=MSO_ANCHOR.MIDDLE)
    # responsibilities
    cells(cv, 5.72, 0.95, [("Data engineering prepares", "Trusted, reusable patient data: Bronze, Silver, Gold A "
                            "and Gold B contracts, refresh SLAs, reconciliation."),
                           ("AI agents discover and accelerate", "Analyse aggregated populations, recommend "
                            "campaigns, build journeys and integration contracts. No PHI."),
                           ("Compliance + activation services enforce", "Identity, consent, authorized purpose, "
                            "minimum necessary and the approved release, independently of the agents.")],
          [CYAN, VIO, MAGENTA])
    return s


# ================================================================== B campaign harness
MANIFEST = ["Approved purpose", "Source snapshot", "Audience version", "Field allowlist", "Policy version",
            "Approval IDs (G1-G3)", "Expiry", "Destination", "Reconciliation results"]
CHECKS = ["Campaign and purpose are approved", "Audience came from the approved eligibility rule",
          "Identity and encounter links are valid", "Consent, restrictions and opt-outs are current",
          "No new suppression since the audience was prepared", "Only allowlisted fields leave",
          "Destination and credentials are approved", "Scoped to campaign, version, destination, time window"]


def harness(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "AGENTIC CAMPAIGNS · CAMPAIGN HARNESS (REPLACES SLIDE 5)",
                "Agents recommend on aggregates; governed services decide who can be contacted",
                "Follow the numbers. Agents work on Gold A only; deterministic services with no LLM decide who "
                "can be contacted, and consent is re-checked at release and at send.")
    X0, X1 = 0.45, 9.72
    lanes = [("PEOPLE", 1.3, 1.0, "2A2A2A", "people"), ("AGENTS", 2.4, 1.1, "0A4A52", "harness"),
             ("SERVICES", 3.6, 1.38, "3A0A1E", "googlecloud"),
             ("SEND", 5.08, 0.92, "0A2A44", "salesforce")]
    for t, y, h, c, ic in lanes:
        cv.rect(X0, y, X1 - X0, h, CARD, None)
        side(cv, X0, y, h, t, c, ic)
    cx = [0.85 + i * 1.78 for i in range(5)]
    cw = 1.66
    py, ph = 1.38, 0.84
    item(cv, cx[0], py, cw, ph, "Request", "Codex or the campaign console: “Build a CHF post-discharge "
         "campaign”.", line=CYAN, n=1, icon="openai-white", bsize=5.5)
    gatecard(cv, cx[1], py, cw, ph, "G1", "Opportunity", "Clinical, product and privacy approve the purpose and "
             "eligibility rule.", 4)
    gatecard(cv, cx[2], py, cw, ph, "G3", "Journey", "Clinical + compliance approve content after tests pass.", 7)
    gatecard(cv, cx[3], py, cw, ph, "G2", "Pre-send", "Masked report: identity, live consent, suppression, "
             "minimum fields, channel.", 8)
    item(cv, cx[4], py, cw, ph, "Results", "Measured outcomes against a holdout; no patient view.", n=13,
         icon="people", bsize=5.5)
    ay, ah = 2.48, 0.94
    item(cv, cx[0], ay, cw, ah, "Campaign Strategist", "Recommends only from approved, aggregated insight.",
         line=VIO, n=2, ncolor=VIO, bsize=5.5)
    item(cv, cx[1], ay, cw, ah, "Analytics Agent", "Volume, overlap, care gaps; forecast ranges from approved "
         "benchmarks.", line=VIO, n=3, ncolor=VIO, bsize=5.5)
    item(cv, cx[2], ay, cw, ah, "Developer + Testing", "Integration contract and journey from approved "
         "templates; tested on test contacts.", line=VIO, n=6, ncolor=VIO, bsize=5.5)
    item(cv, cx[3], ay, cw, ah, "Privacy Agent", "Explains exclusions, investigates exceptions. Does not decide "
         "consent.", line=VIO, bsize=5.5)
    item(cv, cx[4], ay, cw, ah, "Measurement Agent", "Reads aggregate results; keeps forecasts apart from "
         "measured outcomes.", line=VIO, n=12, ncolor=VIO, bsize=5.5)
    sy, sh = 3.68, 1.22
    item(cv, cx[0], sy, cw, sh, "Audience intelligence", ["Gold A: aggregates, cohorts, bounded queries, "
         "disclosure control."], line=YELLOW, icon="googlebigquery", bsize=5.5)
    item(cv, cx[1], sy, cw, sh, "Eligibility, consent + privacy policy engine", ["Gold B: approved rules; "
         "purpose- and channel-aware consent; exclusion reasons."], line=MAGENTA, n=5, ncolor=MAGENTA,
         bsize=5.5, tsize=6.2)
    item(cv, cx[2], sy, cw, sh, "Identity + token service", ["Governed identity resolution, token vault; "
         "re-identifies only for an approved release."], line=MAGENTA, bsize=5.5)
    fx, fy, fw, fh = cv.frame(cx[3], sy, cw, sh, BLACK, 0.04)
    badge(cv, fx + 0.05, fy + 0.05, 9, YELLOW, 0.2)
    cv.text(fx + 0.3, fy + 0.02, fw - 0.34, 0.36, "Authorized patient activation gateway", size=6.2, bold=True,
            color=WHITE, line=0.9)
    cv.text(fx + 0.08, fy + 0.4, fw - 0.14, fh - 0.42, "Validates the manifest, re-checks consent and "
            "suppression, releases minimum fields, expires, audits. Own isolated identity.", size=5.5,
            color=SOFT)
    item(cv, cx[4], sy, cw, sh, "Measurement service", ["Holdout attribution on protected identities; outputs "
         "aggregates only."], line=MAGENTA, n=11, ncolor=MAGENTA, bsize=5.5)
    vy, vh = 5.15, 0.78
    cv.text(cx[0], vy, cx[2] + cw - cx[0], vh, [[("Data Cloud stays a supported route. ", {"bold": True,
                                                                                         "color": WHITE}),
                                                 ("Both adapters are measured on latency, cost, freshness, "
                                                  "consent integrity, complexity and migration impact.", {})]],
            size=6.2, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    ah2 = (vh - 0.06) / 2
    item(cv, cx[3], vy, cw, ah2, "A · via Data Cloud", None, line=CYAN, icon="salesforce", tsize=6)
    item(cv, cx[3], vy + ah2 + 0.06, cw, ah2, "B · direct, after parity", None, line=CYAN, dash="dash", tsize=6)
    item(cv, cx[4], vy, cw, vh, "Marketing Cloud journey", "Send-time consent, STOP, caps; responses back.",
         line=CYAN, n=10, icon="salesforce", bsize=5.4, tsize=6.2)
    # a few arrows, all within a column or between neighbours
    m = cw / 2
    cv.arrow(cx[0] + m, py + ph, cx[0] + m, ay, WHITE, 1.25)
    cv.arrow(cx[0] + m - 0.1, ay + ah, cx[0] + m - 0.1, sy, VIO, 1.25)
    cv.arrow(cx[0] + m + 0.1, sy, cx[0] + m + 0.1, ay + ah, YELLOW, 1.25)
    cv.arrow(cx[1] + m, ay, cx[1] + m, py + ph, VIO, 1.25)
    cv.arrow(cx[3] + m, sy + sh, cx[3] + m, vy, YELLOW, 1.5)
    cv.arrow(cx[3] + cw, vy + ah2 / 2, cx[4], vy + ah2 / 2, WHITE, 1.0)
    cv.arrow(cx[3] + cw, vy + ah2 + 0.06 + ah2 / 2, cx[4], vy + ah2 + 0.06 + ah2 / 2, WHITE, 1.0)
    cv.arrow(cx[4] + m, vy, cx[4] + m, sy + sh, CYAN, 1.25)
    cv.arrow(cx[4] + m, sy, cx[4] + m, ay + ah, MAGENTA, 1.25)
    cv.arrow(cx[4] + m, ay, cx[4] + m, py + ph, VIO, 1.25)
    cv.arrow(cx[2] + cw, sy + sh / 2, cx[3], sy + sh / 2, MAGENTA, 1.25)
    cv.arrow(cx[1] + cw, sy + sh / 2, cx[2], sy + sh / 2, MAGENTA, 1.25)
    # right: manifest + checks
    rx, rw = 9.88, 12.88 - 9.88
    cv.rect(rx, 1.3, rw, 1.98, CARD, YELLOW, 1.0, R, 0.04)
    cv.text(rx + 0.12, 1.34, rw - 0.24, 0.24, "CAMPAIGN RELEASE MANIFEST", size=6.4, bold=True, color=YELLOW,
            track=120, anchor=MSO_ANCHOR.MIDDLE)
    for k, t in enumerate(MANIFEST):
        x = rx + 0.12 + (k % 2) * (rw - 0.24) / 2
        y = 1.64 + (k // 2) * 0.32
        cv.rect(x, y, (rw - 0.24) / 2 - 0.08, 0.26, CARD2, EDGE, 0.6, R, 0.2)
        cv.text(x, y, (rw - 0.24) / 2 - 0.08, 0.26, t, size=5.8, color=WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
    cv.rect(rx, 3.4, rw, 2.6, CARD, MAGENTA, 1.0, R, 0.04)
    cv.text(rx + 0.12, 3.44, rw - 0.24, 0.24, "GATEWAY CHECKS · AT RELEASE AND AT SEND", size=6,
            bold=True, color=MAGENTA, track=100, anchor=MSO_ANCHOR.MIDDLE)
    for k, t in enumerate(CHECKS):
        y = 3.74 + k * 0.265
        badge(cv, rx + 0.12, y + 0.03, k + 1, MAGENTA, 0.18)
        cv.text(rx + 0.38, y, rw - 0.46, 0.24, t, size=5.9, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    cells(cv, 6.12, 0.72, [("Agents: Gold A only", "Aggregates and bounded queries; no patient rows."),
                           ("Consent by approved rules", "Purpose- and channel-aware; checked at release and send."),
                           ("Release is scoped", "One campaign, one audience version, minimum fields, expiry."),
                           ("Marketing Cloud never browses Gold", "It receives only what the manifest allows.")],
          [VIO, MAGENTA, YELLOW, CYAN])
    return s


# ================================================================== C PHI governance
ADDG = {"A": (0, 1, "DLP classification", "UPSTREAM"), "B": (1, 0, "Steward review", "UPSTREAM"),
        "C": (2, 1, "Reviewed classifications", "BUILT"), "G": (2, 2, "Rollout gate", "BUILT"),
        "T": (3, 2, "Data Policy Tagger", "LIVE IN DEV"), "P": (4, 3, "Policies applied", "LIVE IN DEV"),
        "V": (5, 1, "Bindings + validation", "BUILT"), "R": (4, 2, "Policy change review", "PROPOSED")}
ST = {"UPSTREAM": SOFT, "BUILT": GREEN, "LIVE IN DEV": CYAN, "IN PROGRESS": ORANGE, "PROPOSED": VIO}
GAPS = [("A", "Re-identification risk", "Assess column, row, combination and aggregate disclosure risk: dates, "
         "rare diagnoses, facility and geography can identify a patient."),
        ("B", "Agent access", "Aggregates, privacy-protected segments and bounded queries by default. Tokenized "
         "rows only for an authorized task."),
        ("C", "Separate identities", "Discovery, analysis, policy admin, token resolution and release each get "
         "their own identity. Raw access is isolated from the LLM runtime."),
        ("D", "Beyond BigQuery", "Cloud Storage exports, temp tables, job results, logs, notebooks, API payloads "
         "and backups are governed too."),
        ("E", "Fail-closed policy lifecycle", "Rejecting a new classification never removes existing protection. "
         "Relaxing a policy needs independent approval."),
        ("F", "Compliance control register", "Maps technical evidence to HIPAA safeguards, BAAs, access reviews, "
         "incident and breach handling.")]


def governance(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "AGENTIC CAMPAIGNS · PHI GOVERNANCE (REPLACES SLIDE 6)",
                "Masking is an access control, not de-identification",
                "Today's classify, approve, protect and prove flow stays. The proposed additions close gaps A-F; "
                "statuses show what exists today, and PROPOSED is design, not production evidence.")
    # today's flow, compact
    fx0, fx1, fy0 = 0.45, 7.6, 1.3
    cv.text(fx0, fy0, 5, 0.22, "TODAY'S FLOW, WITH STATUS", size=6.4, bold=True, color=WHITE, track=120,
            anchor=MSO_ANCHOR.MIDDLE)
    lanes = [("REVIEW", "Collibra", "people"), ("GOVERNANCE", "datarepository", "googlecloud"),
             ("COMPUTE", "datalake", "googlecloud"), ("TARGET", "dataplace", "googlecloud")]
    lh, lg, ly0 = 0.56, 0.05, fy0 + 0.3
    nx0, pitch, nw, nh = fx0 + 1.0, 1.02, 0.9, 0.46
    for i, (role, proj, ic) in enumerate(lanes):
        y = ly0 + i * (lh + lg)
        cv.rect(fx0, y, fx1 - fx0, lh, CARD, None)
        cv.rect(fx0, y, 0.92, lh, CARD2, None)
        logo(cv, ic, fx0 + 0.06, y + (lh - 0.18) / 2, 0.18)
        cv.text(fx0 + 0.3, y, 0.62, lh, [[(role, {"bold": True, "color": WHITE, "size": 5.4})],
                                         [(proj, {"color": DIM, "size": 4.8, "font": "Courier New"})]],
                size=5.4, anchor=MSO_ANCHOR.MIDDLE)

    def box(nid):
        c, l = ADDG[nid][:2]
        return nx0 + c * pitch, ly0 + l * (lh + lg) + (lh - nh) / 2

    def rt(nid, dy=0.0):
        x, y = box(nid)
        return x + nw, y + nh / 2 + dy

    def lf(nid):
        x, y = box(nid)
        return x, y + nh / 2
    for a, b in (("A", "B"), ("B", "C"), ("C", "T"), ("G", "T"), ("T", "P"), ("P", "V")):
        (x1, y1), (x2, y2) = rt(a), lf(b)
        xm = x1 + (x2 - x1) / 2
        path(cv, [(x1, y1), (xm, y1), (xm, y2), (x2, y2)], CYAN, 1.25)
    rx_, ry_ = box("R")
    x1, y1 = rt("C", 0.12)
    path(cv, [(x1, y1), (rx_ + nw / 2, y1), (rx_ + nw / 2, ry_)], VIO, 1.25, dash="dash")
    px_, py_ = box("P")
    path(cv, [(rx_ + nw / 2, ry_ + nh), (px_ + nw / 2, py_)], VIO, 1.25, dash="dash")
    for nid, (c, l, t, st) in ADDG.items():
        x, y = box(nid)
        gate_ = nid in ("B", "G")
        cv.rect(x, y, nw, nh, TINT[YELLOW] if gate_ else BLACK,
                YELLOW if gate_ else (VIO if nid == "R" else EDGE), 1.0, R, 0.1,
                dash="dash" if nid in ("G", "R") else None)
        cv.text(x + 0.05, y + 0.03, nw - 0.08, 0.24, t, size=5.6, bold=True, color=WHITE, line=0.85)
        pill(cv, x + 0.05, y + nh - 0.17, 0.1 + 0.042 * len(st), st, ST[st], h=0.13, size=4.2)
    ly = ly0 + 4 * (lh + lg) + 0.02
    cv.text(fx0, ly, fx1 - fx0, 0.2, [[("Rejected path, changed: ", {"bold": True, "color": VIO}),
                                       ("a rejection never removes existing protection; relaxing a policy needs "
                                        "independent approval and a risk assessment.", {})]], size=5.8,
            color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    # right: three different things
    rx, rw = 7.78, 12.88 - 7.78
    cv.rect(rx, 1.3, rw, 2.72, CARD, EDGE, 0.75, R, 0.04)
    cv.text(rx + 0.12, 1.34, rw - 0.24, 0.24, "THREE DIFFERENT THINGS", size=6.4, bold=True, color=YELLOW,
            track=120, anchor=MSO_ANCHOR.MIDDLE)
    hdr = ["", "What it does", "HIPAA implication"]
    colx = [rx + 0.12, rx + 1.35, rx + 3.05]
    colw = [1.2, 1.65, rw - 3.12]
    for x, w, t in zip(colx, colw, hdr):
        cv.text(x, 1.62, w, 0.18, t, size=5.6, bold=True, color=DIM)
    rows = [("Masking", "Hides values from readers without access", "Access control; the data underneath is "
             "still PHI", MAGENTA),
            ("Tokenization", "Replaces identifiers; reversible through a separate map", "Can still be PHI, "
             "depending on linkability and context", ORANGE),
            ("De-identification", "Safe Harbor or Expert Determination met", "Can leave PHI status when the "
             "standard is met", GREEN)]
    for k, (a, b, c, col) in enumerate(rows):
        y = 1.84 + k * 0.56
        cv.rect(rx + 0.08, y, rw - 0.16, 0.5, CARD2, None)
        cv.rect(rx + 0.08, y, 0.05, 0.5, col, None)
        cv.text(colx[0] + 0.06, y, colw[0], 0.5, a, size=6.6, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(colx[1], y, colw[1], 0.5, b, size=5.8, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(colx[2], y, colw[2], 0.5, c, size=5.8, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(rx + 0.12, 3.54, rw - 0.24, 0.42, [[("Hashing or masking alone does not meet either HHS method. ",
                                                 {"bold": True, "color": WHITE}),
                                                ("Agents get aggregates by default.", {})]], size=6, color=SOFT,
            anchor=MSO_ANCHOR.MIDDLE)
    # gaps A-F
    gy = 4.3
    cv.text(0.45, gy, 6, 0.22, [[("PROPOSED ADDITIONS", {"bold": True, "track": 120, "color": VIO}),
                                 ("   closing gaps A-F", {"color": SOFT})]], size=6.4, anchor=MSO_ANCHOR.MIDDLE)
    gw = (12.43 - 2 * 0.12) / 3
    for k, (L, t, d) in enumerate(GAPS):
        x = 0.45 + (k % 3) * (gw + 0.12)
        y = gy + 0.28 + (k // 3) * 0.82
        cv.rect(x, y, gw, 0.74, CARD, VIO, 0.9, R, 0.06)
        badge(cv, x + 0.08, y + 0.08, L, VIO, 0.22)
        cv.text(x + 0.38, y + 0.06, gw - 1.2, 0.26, t, size=7, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        pill(cv, x + gw - 0.72, y + 0.1, 0.62, "PROPOSED", VIO, h=0.15, size=4.8)
        cv.text(x + 0.12, y + 0.36, gw - 0.24, 0.4, d, size=5.9, color=SOFT)
    ey = gy + 0.28 + 2 * 0.82 + 0.02
    cv.rect(0.45, ey, 12.43, 0.42, TINT[GREEN], GREEN, 0.9, R, 0.2)
    cv.text(0.6, ey, 12.1, 0.42, [[("ENFORCEMENT EVIDENCE  ", {"bold": True, "color": GREEN, "track": 100}),
                                   ("Discovery inventory across the approved scope · row-level security, authorized "
                                    "views and scoped datasets alongside dynamic masking · negative-access tests · "
                                    "drift detection · periodic access review", {"color": WHITE})]], size=6.2,
            anchor=MSO_ANCHOR.MIDDLE)
    return s


# ================================================================== D gaps, Data Cloud, validation
GAP_ROWS = [("Critical", "No defined patient identity service", "Governed identity resolution, token vault, "
             "authorized re-identification"),
            ("Critical", "Consent treated as a static filter", "Purpose- and channel-aware consent at release and "
             "send"),
            ("Critical", "Masking treated as de-identification", "De-identification assessment, minimization, "
             "access isolation"),
            ("Critical", "Direct BigQuery → Marketing Cloud bypass", "Parity assessment; two activation adapters"),
            ("Critical", "No rule on which campaigns are permitted", "Privacy-approved purpose and authorization "
             "rules"),
            ("High", "Always-ready Gold table", "Versioned intelligence plus fresh per-campaign eligibility"),
            ("High", "Source data quality", "Data quality and reconciliation engine"),
            ("High", "Graph fed from Gold patient data", "Metadata, lineage and approved aggregates only"),
            ("High", "ROI as prediction only", "Historical outcomes, approved costs, holdout measurement"),
            ("High", "Downstream privacy coverage", "SFMC, providers, tracking, devices: minimal content, "
             "retention, audit"),
            ("High", "Refresh not specified", "Refresh SLAs, event handling, source reconciliation"),
            ("High", "No tenant isolation", "Tenant-scoped identities, storage, policies, keys, audit")]
PARITY = [("Ingestion", "BigQuery → Data Cloud", "Stays in governed BigQuery"),
          ("Patient relationships", "Individual / Unified Individual", "Equivalent identity service"),
          ("Eligibility", "Data Cloud segmentation", "Governed BigQuery eligibility engine"),
          ("Consent filtering", "Eligibility + SFMC suppression", "Purpose- and channel-aware rules"),
          ("Activation", "Data Cloud → SFMC", "Authorized SFMC API or import"),
          ("Encounter dedup", "Campaign identity logic", "Defined selection rules"),
          ("Survey feedback", "SFMC → Data Cloud → BigQuery", "Direct governed feedback"),
          ("Frequency rules", "Not yet across campaigns", "Global suppression policy")]
CHECKLIST = ["Verify the BigQuery inventory, PHI classifications and Collibra ownership",
             "Map the CHF audience definition, identity relationships and consent sources",
             "Define Bronze, Silver, Gold A and Gold B data contracts",
             "Design the identity / token service and the activation gateway",
             "Test masked agent access, including attempts to re-identify",
             "Compare Data Cloud activation with direct SFMC activation",
             "Validate suppression changes between audience preparation and send",
             "Benchmark runtime, data quality and cost against today",
             "Secure privacy, security and architecture approval"]


def gaps(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "AGENTIC CAMPAIGNS · GAPS AND PROOF",
                "What we change, and how we prove it on one campaign",
                "Twelve gaps from the architecture review, the Data Cloud decision, and a CHF post-discharge pilot "
                "on synthetic data first.")
    # gap register
    lx, lw = 0.45, 7.25
    cv.text(lx, 1.3, lw, 0.22, "GAP REGISTER", size=6.4, bold=True, color=WHITE, track=120,
            anchor=MSO_ANCHOR.MIDDLE)
    cx = [lx + 0.08, lx + 0.85, lx + 3.5]
    for x, t in zip(cx, ["Priority", "Gap", "What we add"]):
        cv.text(x, 1.54, 2.5, 0.18, t, size=5.6, bold=True, color=DIM)
    for k, (p, g, f) in enumerate(GAP_ROWS):
        y = 1.74 + k * 0.335
        cv.rect(lx, y, lw, 0.3, CARD if k % 2 == 0 else CARD2, None)
        pill(cv, cx[0], y + 0.07, 0.62, p.upper(), MAGENTA if p == "Critical" else ORANGE, p == "Critical",
             h=0.16, size=4.8, tcolor=WHITE if p == "Critical" else None)
        cv.text(cx[1], y, 2.6, 0.3, g, size=6.2, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(cx[2], y, lw - 3.6, 0.3, f, size=6, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    # data cloud parity
    rx, rw = 7.88, 12.88 - 7.88
    cv.rect(rx, 1.3, rw, 2.78, CARD, CYAN, 1.0, R, 0.04)
    logo(cv, "salesforce", rx + 0.12, 1.35, 0.2)
    cv.text(rx + 0.4, 1.33, rw - 0.5, 0.24, "DATA CLOUD: KEEP AS ADAPTER A, TEST ADAPTER B", size=6.2, bold=True,
            color=CYAN, track=100, anchor=MSO_ANCHOR.MIDDLE)
    px = [rx + 0.12, rx + 1.3, rx + 3.05]
    pw = [1.15, 1.72, rw - 3.12]
    for x, w, t in zip(px, pw, ["Capability", "Today (VBC)", "Needed if bypassed"]):
        cv.text(x, 1.6, w, 0.18, t, size=5.4, bold=True, color=DIM)
    for k, row in enumerate(PARITY):
        y = 1.8 + k * 0.275
        cv.rect(rx + 0.08, y, rw - 0.16, 0.25, CARD2 if k % 2 == 0 else CARD, None)
        for j, (x, w, t) in enumerate(zip(px, pw, row)):
            cv.text(x, y, w, 0.25, t, size=5.6, bold=j == 0, color=WHITE if j < 2 else SOFT,
                    anchor=MSO_ANCHOR.MIDDLE)
    # validation plan
    vy = 4.2
    cv.rect(rx, vy, rw, 2.62, CARD, GREEN, 1.0, R, 0.04)
    cv.text(rx + 0.12, vy + 0.04, rw - 0.24, 0.24, [[("PROVE IT ON ONE CAMPAIGN", {"bold": True, "track": 100}),
                                                    ("   CHF post-discharge · synthetic data first",
                                                     {"color": SOFT})]], size=6.2, color=GREEN,
            anchor=MSO_ANCHOR.MIDDLE)
    for k, t in enumerate(CHECKLIST):
        y = vy + 0.34 + k * 0.252
        cv.rect(rx + 0.14, y + 0.05, 0.13, 0.13, None, GREEN, 0.9)
        cv.text(rx + 0.36, y, rw - 0.46, 0.23, t, size=5.9, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    # recommendation
    cv.rect(lx, 5.85, lw, 0.97, TINT[VIO], VIO, 1.0, R, 0.06)
    cv.text(lx + 0.15, 5.85, lw - 0.3, 0.97, [[("Recommendation  ", {"bold": True, "color": VIO}),
                                               ("Make the Bronze-Silver-Gold foundation and governed activation the "
                                                "primary design. The campaign harness and PHI governance views sit "
                                                "on top of it. Keep the goal: faster campaign launch, defensibly.",
                                                {"color": WHITE})]], size=7, anchor=MSO_ANCHOR.MIDDLE)
    return s


def main(out=OUT):
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    for fn in (architecture, harness, governance, gaps):
        fn(prs, layout)
    for line in fix_footers(prs):
        print(line)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print("wrote", os.path.abspath(out), "-", len(prs.slides), "slides")


if __name__ == "__main__":
    main(*sys.argv[1:])
