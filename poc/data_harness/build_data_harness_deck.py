"""Agentic Harness + data harness: a short deck on the IG template (13 slides).

1 Title
2 One agent team: seven delivery agents and seven campaign agents
3 The platform, now with the data harness
4 Health Cloud: a story from Codex to production (GCP patient data not needed, shown grey)
5 A Health Cloud change: ~2 weeks manual vs ~2 days on the Agentic Harness (hc_timeline.py)
6 Section: HIPAA-Aligned Agentic Healthcare Campaign Platform
7 Marketing campaigns end to end, from governed data to a tested, approved campaign (combined_flow.py)
8 Launching a campaign: ~90 days manual vs ~10-12 days on the data harness (campaign_timeline.py)
9 How we build it: safeguards first, release last (campaign_platform_deck.py)
10 PHI governance: classification and approval stay; masking alone is not HIPAA compliance
11 Sierra voice agent: governed knowledge for answers, scoped APIs for patient actions (sierra_voice.py)
12 Platform fit: one harness for every pod and job (platform_fit.py)
13 Lean squads and three-year savings (lean_team.py)

Black deck: the original colourful design in IG colours (muted_theme.vivid). White deck (theme "white"): the
muted styles in muted_theme.py, recoloured for white.

    python3 build_data_harness_deck.py [out.pptx]
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "..", "deck")
sys.path.insert(0, DECK)

from pptx import Presentation                                     # noqa: E402
from pptx.enum.shapes import MSO_SHAPE                             # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN                    # noqa: E402
from pptx.util import Inches                                       # noqa: E402

from build_mad_ig_deck import BLACK, CYAN, GREEN, MAGENTA, NAVY, WHITE, YELLOW, Canvas, set_notes  # noqa: E402
from footer import fix_footers                                     # noqa: E402
from story_slides import dark_header, set_bg                       # noqa: E402

SRC = os.path.join(DECK, "source", "Ascension Agentic Harness - final.pptx")
OUT = os.path.join(HERE, "Ascension Agentic Harness - Data Harness.pptx")
OUT_WHITE = os.path.join(HERE, "Ascension Agentic Harness - Data Harness (white).pptx")
LOGOS = os.path.join(DECK, "assets", "logos")
R = MSO_SHAPE.ROUNDED_RECTANGLE

CARD, CARD2, EDGE = "141414", "1E1E1E", "3A3A3A"
SOFT, DIM = "BFBFBF", "8C8C8C"
VIO = "8B7CD8"
GREY_BG, GREY_LN, GREY_TX = "0E0E0E", "3A3A3A", "6B6B6B"
TINT = {CYAN: "06262B", YELLOW: "2B2508", MAGENTA: "2B0716", VIO: "1A1630", GREEN: "0E2A08", NAVY: "0A1A24"}


# ------------------------------------------------------------------ helpers
def new_slide(prs, layout):
    s = prs.slides.add_slide(layout)
    set_bg(s, BLACK)
    return s, Canvas(s)


def logo(cv, name, x, y, size):
    cv.s.shapes.add_picture(os.path.join(LOGOS, name + ".png"), Inches(x), Inches(y), Inches(size), Inches(size))


def badge(cv, x, y, n, color=CYAN, size=0.22):
    cv.rect(x, y, size, size, color, None, shape=MSO_SHAPE.OVAL)
    cv.text(x, y, size, size, str(n), size=6.6, bold=True, color=BLACK, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)


def pill(cv, x, y, w, t, color, filled=False, h=0.18, size=5.6, tcolor=None):
    cv.rect(x, y, w, h, color if filled else None, color, 0.75, R, 0.5)
    cv.text(x, y, w, h, t, size=size, bold=True, color=tcolor or (BLACK if filled else color),
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, track=80)


def card(cv, x, y, w, h, title, body=None, line=EDGE, fill=CARD, n=None, ncolor=CYAN, tag=None, lw=1.0,
         dash=None, tsize=8, bsize=6.5, tcolor=WHITE, bcolor=SOFT, icon=None):
    cv.rect(x, y, w, h, fill, line, lw, R, 0.07, dash=dash)
    tx = x + 0.1
    if n is not None:
        badge(cv, x + 0.08, y + 0.07, n, ncolor)
        tx = x + 0.36
    if icon:
        logo(cv, icon, tx, y + 0.07, 0.2)
        tx += 0.26
    right = 0.08
    if tag:
        tw = 0.18 + 0.055 * len(tag[0])
        pill(cv, x + w - tw - 0.08, y + 0.08, tw, tag[0], tag[1], tag[2] if len(tag) > 2 else False)
        right = tw + 0.14
    cv.text(tx, y + 0.06, x + w - tx - right, 0.22, title, size=tsize, bold=True, color=tcolor,
            anchor=MSO_ANCHOR.MIDDLE)
    if body:
        if isinstance(body, str):
            body = [body]
        cv.text(x + 0.1, y + 0.31, w - 0.2, h - 0.34, [[(b, {})] if isinstance(b, str) else b for b in body],
                size=bsize, color=bcolor, after=1)


def bar(cv, x, y, w, t, sub, color, tcolor=BLACK, h=0.3):
    cv.rect(x, y, w, h, color, None)
    cv.text(x + 0.1, y, w - 0.2, h, [[(t, {"bold": True, "track": 150}), ("   " + sub, {"color": tcolor})]
                                     if sub else [(t, {"bold": True, "track": 150})]],
            size=7, color=tcolor, anchor=MSO_ANCHOR.MIDDLE)


def path(cv, pts, color, w=1.5, dash=None):
    for k, (p, q) in enumerate(zip(pts, pts[1:])):
        cv.arrow(p[0], p[1], q[0], q[1], color, w, dash=dash, head=(k == len(pts) - 2))


# ================================================================== 1 title
TITLE_NOTES = ("Open: the same Agentic Harness now does two jobs. For Health Cloud it takes a Jira story to a tested, "
               "approved change. For Marketing Cloud it plans and prepares campaigns on governed Google Cloud data, "
               "where agents only ever see masked data and people release the right patients.")


def title_slide(prs):
    s = prs.slides[0]
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("An Agentic Harness"):
            r = sh.text_frame.paragraphs[0].runs
            r[0].text = ("One harness, two jobs: Health Cloud delivery from Codex, and Marketing Cloud campaigns "
                         "on governed, masked Google Cloud data.")
            for extra in r[1:]:
                extra.text = ""
        elif t == "AI-DRIVEN SALESFORCE DELIVERY":
            sh.text_frame.paragraphs[0].runs[0].text = "AGENTS FOR DELIVERY AND DATA"
    set_notes(s, TITLE_NOTES)


# ================================================================== 2 agent team
DELIVERY = [("Orchestrator", "Classifies the task, gathers context, routes the chain, enforces gates", "COORDINATES",
             CYAN),
            ("Requirements Agent", "User stories and functional decomposition from the Jira story", "HUMAN GATE",
             YELLOW),
            ("Architect Agent", [[("Impact and reuse from the graph; ", {}),
                                  ("Change Guardian", {"bold": True, "color": CYAN}),
                                  (" flags cross-pod risk", {})]], "HUMAN GATE", YELLOW),
            ("Developer Agent", "Builds to Ascension standards in DEV; journeys for Marketing Cloud", "AI-LED",
             GREEN),
            ("Review Agent", "Quality, security and Ascension coding standards", "HUMAN GATE", YELLOW),
            ("Test Agent", "Unit, regression and performance tests; Zephyr to Copado", "HUMAN GATE", YELLOW),
            ("Release Agent", "Packages and promotes through Copado to lower environments", "AI-LED", GREEN)]
DATA = [("Orchestrator Agent", "Coordinates opportunity analysis, approvals and execution steps", "COORDINATES",
         CYAN),
        ("Campaign Strategist", "Identifies campaigns from approved, aggregated insight (Gold A)", "HUMAN GATE",
         YELLOW),
        ("Audience Analytics", "Aggregate cohorts, overlap and reach; counts, never names", "AI-LED", GREEN),
        ("Compliance Guardian", "Investigates policy risk, prepares evidence; cannot override policy", "GUARD",
         MAGENTA),
        ("Campaign Design", "Platform-ready specs; journeys on Marketing Cloud's native tools",
         "HUMAN GATE", YELLOW),
        ("Activation Readiness", "Tests on test contacts; readiness report without raw PHI", "GUARD", MAGENTA),
        ("Measurement Agent", "Aggregated outcomes, forecasts and experiments; ROI report", "AI-LED", GREEN)]

TEAM_NOTES = """SLIDE: One agent team: seven delivery agents and seven campaign agents   (about 2 minutes)

KEY MESSAGE
The seven delivery agents are unchanged and serve every pod. Seven campaign agents are added for Marketing Cloud campaigns on governed Google Cloud data. They work on aggregated, masked data only; deterministic services, not agents, resolve identity, evaluate eligibility, enforce policy and release patient data. People approve at every gate.

WALK IT
- Left: agent sources. Agents read from and write to these systems, always through the tool gateway and never in production. The engineering sources (Jira, Confluence, Salesforce, Git, Copado, Drive) for every pod. For Marketing Cloud pods, governed Google Cloud data: Gold A aggregates and masked views only.
- Middle: the seven delivery agents, then the seven campaign agents:
  Orchestrator (coordinates analysis, approvals and execution), Campaign Strategist (campaign opportunities from Gold A), Audience Analytics (cohorts, overlap and reach as counts), Compliance Guardian (policy risk and evidence; cannot override policy), Campaign Design (platform-ready specs, journeys with Marketing Cloud's native automation), Activation Readiness (tests on test contacts, readiness without raw PHI), Measurement (aggregated outcomes, forecasts, ROI).
- Right: what changes per pod. Health Cloud pods use the delivery agents; Marketing Cloud pods use the delivery and the campaign agents.
- Bottom: two governed paths, with who does each step. Delivery: the Developer Agent builds, the Release Agent moves it through Copado, Review and Test Agents run the quality gates; business users sign UAT and the change board approves at CAB. Campaigns: the Orchestrator takes the prompt, Strategist and Analytics plan on Gold A, Design and Readiness build and test, the SME and tech lead review before send, and the policy service releases.
- The Change Guardian sits with the Architect Agent: it flags any change that affects another pod before anything is built."""


def team_slide(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "ASCENSION AGENTIC HARNESS · AGENT TEAM",
                "One agent team: seven delivery agents and seven campaign agents",
                "Delivery agents build and release changes for every pod. Campaign agents plan, design and test on "
                "masked data; deterministic services release patients. People approve at every gate.")
    top, bot = 1.28, 5.95
    # left: what agents read
    lx, lw = 0.45, 2.75
    bar(cv, lx, top, lw, "AGENT SOURCES", "read + write via the gateway", NAVY, WHITE)
    cv.rect(lx, top + 0.3, lw, bot - top - 0.3, CARD, EDGE, 0.75)
    srcs = [("jira", "Jira", "stories, criteria, history"), ("confluence", "Confluence", "designs, business rules"),
            ("salesforce", "Salesforce", "objects, Flows, Apex, journeys"), ("git", "Git", "Apex, LWC, metadata"),
            (None, "Copado", "releases, test runs"), ("googledrive", "Google Drive", "specs, mapping sheets")]
    y = top + 0.42
    for ic, n, d in srcs:
        if ic:
            logo(cv, ic, lx + 0.12, y + 0.06, 0.22)
        else:
            pill(cv, lx + 0.12, y + 0.08, 0.22, "Co", SOFT, h=0.2, size=5.2)
        cv.text(lx + 0.42, y, lw - 0.5, 0.36, [[(n, {"bold": True, "color": WHITE})], [(d, {"color": DIM,
                                                                                            "size": 6})]],
                size=7.4, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.42
    cv.text(lx + 0.12, y + 0.02, lw - 0.24, 0.2, "FOR MARKETING CLOUD PODS", size=5.8, bold=True, color=MAGENTA,
            track=120)
    gy = y + 0.25
    cv.rect(lx + 0.1, gy, lw - 0.2, 0.86, TINT[MAGENTA], MAGENTA, 1.25, R, 0.08)
    logo(cv, "googlecloud", lx + 0.2, gy + 0.08, 0.22)
    cv.text(lx + 0.5, gy + 0.05, lw - 0.7, 0.26, "Governed GCP data", size=7.8, bold=True, color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(lx + 0.2, gy + 0.33, lw - 0.4, 0.5, "Gold A only: protected aggregates and masked views. Raw PHI "
            "stays behind the gateway.", size=6.3, color=SOFT)
    cv.text(lx + 0.12, bot - 0.42, lw - 0.24, 0.36, "Every fact is cited with its source and SHA-256.", size=6.2,
            color=DIM, anchor=MSO_ANCHOR.MIDDLE)

    # middle: agents
    mx, mw = 3.35, 6.35
    bar(cv, mx, top, mw, "DELIVERY AGENTS", "the same on every pod", CYAN)
    rh = 0.276
    y = top + 0.36

    def row(i, n, name, desc, mode, col):
        cv.rect(mx, y, mw, rh - 0.04, CARD if i % 2 == 0 else BLACK, None)
        badge(cv, mx + 0.08, y + (rh - 0.04 - 0.2) / 2, n, col, 0.2)
        cv.text(mx + 0.36, y, 1.65, rh - 0.04, name, size=7.0, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(mx + 2.0, y, mw - 3.1, rh - 0.04, desc, size=6.0, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
        pill(cv, mx + mw - 1.0, y + (rh - 0.04 - 0.16) / 2, 0.9, mode, col, h=0.16)

    for i, (name, desc, mode, col) in enumerate(DELIVERY):
        row(i, i + 1, name, desc, mode, col)
        y += rh
    y += 0.06
    bar(cv, mx, y, mw, "CAMPAIGN AGENTS", "new · masked data only · services, not agents, release patients",
        MAGENTA, WHITE)
    y += 0.36
    for i, (name, desc, mode, col) in enumerate(DATA):
        row(i, i + 8, name, desc, mode, col)
        y += rh

    # right: pod skills
    rx, rw = 9.85, 3.03
    bar(cv, rx, top, rw, "POD SKILLS", "what changes per pod", MAGENTA, WHITE)
    cv.rect(rx, top + 0.3, rw, bot - top - 0.3, CARD, EDGE, 0.75)
    pods = [("Health Cloud pods", "Network Services, Clinical Transformation, Contact Center",
             "Apex, Flows, LWC, sharing, Health Cloud objects", [("DELIVERY AGENTS", CYAN)]),
            ("Marketing Cloud pods", "Nudge / Demand, VBC outreach",
             "Data Cloud, SFMC SQL, AMPscript, journeys", [("DELIVERY + CAMPAIGN AGENTS", MAGENTA),
                                                            ("GCP DATA", MAGENTA)]),
            ("QA, all pods", "Every release", "Zephyr test cases, Copado robotic and regression", [])]
    y = top + 0.42
    for t, sub, d, tags in pods:
        tx0 = rx + 0.12
        if "Cloud" in t:
            logo(cv, "salesforce", tx0, y + 0.01, 0.2)
            tx0 += 0.27
        cv.text(tx0, y, rw - 0.24, 0.22, t, size=8, bold=True, color=MAGENTA if "Marketing" in t else WHITE,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(rx + 0.12, y + 0.22, rw - 0.24, 0.2, sub, size=6, color=DIM)
        cv.text(rx + 0.12, y + 0.42, rw - 0.24, 0.36, d, size=6.6, color=WHITE)
        tx = rx + 0.12
        for tg, c in tags:
            w = 0.2 + 0.064 * len(tg)
            pill(cv, tx, y + 0.78, w, tg, c)
            tx += w + 0.06
        y += 1.12
    cv.rect(rx + 0.1, bot - 0.62, rw - 0.2, 0.5, TINT[MAGENTA], None, shape=R, radius=0.1)
    cv.text(rx + 0.2, bot - 0.62, rw - 0.4, 0.5, "A new pod adds its sources and skills, not new agents.",
            size=6.6, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

    # bottom: two governed paths
    by, ph = 6.02, 0.38
    paths = [("DELIVERY PATH", [("AI build", "Developer Agent"), ("Copado", "Release Agent"),
                                ("Quality gates", "Review + Test Agents"), ("DEV", "Release Agent"),
                                ("SIT", "Test Agent"), ("UAT", "business users"), ("CAB", "change board"),
                                ("PROD", "owning team")], CYAN),
             ("CAMPAIGN PATH", [("Prompt in Codex", "Orchestrator"), ("Gold A insights", "Strategist + Analytics"),
                                ("Build + test", "Design + Readiness"), ("Pre-send review", "SME + tech lead"),
                                ("Gateway release", "policy service"), ("Journey runs", "Marketing Cloud")], MAGENTA)]
    for k, (lab, steps, col) in enumerate(paths):
        yy = by + k * (ph + 0.06)
        cv.text(0.45, yy, 1.3, ph, lab, size=6.4, bold=True, color=col, track=120, anchor=MSO_ANCHOR.MIDDLE)
        x = 1.8
        sw = (12.88 - 1.8 - (len(steps) - 1) * 0.08) / len(steps)
        for j, (st, who) in enumerate(steps):
            gate = st in ("CAB", "UAT", "Pre-send review")
            cv.rect(x, yy, sw, ph, TINT[YELLOW] if gate else CARD, YELLOW if gate else col, 0.9, R, 0.3)
            cv.text(x, yy, sw, ph, [[(st, {"bold": True, "color": YELLOW if gate else WHITE})],
                                    [(who, {"color": YELLOW if gate else SOFT})]], size=7,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=0.9)
            if j < len(steps) - 1:
                cv.arrow(x + sw, yy + ph / 2, x + sw + 0.08, yy + ph / 2, DIM, 0.75)
            x += sw + 0.08
    set_notes(s, TEAM_NOTES)
    return s


# ================================================================== 3 platform
PLATFORM_NOTES = """SLIDE: The platform, now with the data harness   (about 2 minutes)

KEY MESSAGE
Codex is the engine; the Agentic Harness is everything around it. For campaigns it adds a masked data context (Gold A), Collibra masking, a Compliance Guardian with a deterministic policy engine, and a masked pre-send review. Agents get the minimum data they need, never raw PHI.

WALK IT
- Surface: where people work. Codex (a Jira number for a delivery story, or a campaign prompt), the Salesforce LWC console, the approval page, masked Tableau dashboards and the pre-send report, and Jira.
- Runtime and Registry: unchanged. Durable runs, one agent task per step; tools, skills and templates, now including campaign skills and journey templates.
- Context: the knowledge graph and cited retrieval, plus the data context: Gold A aggregates and masked views.
- Govern: Change Guardian, Model Armor guardrails, identities, human gates, evals and the change passport, plus Collibra masking and DLP, and the Compliance Guardian with the policy engine. Patient data leaves only through the activation gateway.
- Bottom: it all runs on what Ascension already owns, including BigQuery, Collibra, DLP, Tableau and Spanner."""


def platform_slide(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "ASCENSION AGENTIC HARNESS · PLATFORM",
                "Agentic Harness platform, now with the data harness",
                "Codex is the engine; the harness is everything around it. New: a masked data context, Collibra "
                "masking, a Compliance Guardian with a policy engine, and a masked pre-send review.")
    chips = [("7 delivery agents", VIO), ("7 campaign agents", MAGENTA), ("9-step workflow", VIO),
             ("Campaign run", MAGENTA), ("Pod skill packs", VIO)]
    x = 1.6
    for t, c in chips:
        w = 0.4 + 0.075 * len(t)
        cv.rect(x, 1.27, w, 0.36, None, c, 1.25, R, 0.5)
        cv.text(x, 1.27, w, 0.36, t, size=8.5, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += w + 0.15
    fy, fh = 1.78, 3.98
    cv.rect(0.45, fy, 9.15, fh, None, YELLOW, 1.0, R, 0.03, dash="sysDot")

    def new_tag(x, y):
        pill(cv, x, y, 0.36, "NEW", CYAN, True, h=0.15, size=5)

    # surface
    sx, sw = 0.6, 1.85
    cv.rect(sx, fy + 0.14, sw, fh - 0.28, None, YELLOW, 1.25, R, 0.05)
    logo(cv, "harness", sx + 0.12, fy + 0.26, 0.28)
    cv.text(sx + 0.48, fy + 0.2, sw - 0.56, 0.42, [[("Harness", {})], [("Surface", {})]], size=10, color=WHITE,
            line=0.9)
    surf = [("openai-white", "Codex", "a Jira number, or a campaign ask"),
            ("salesforce", "Salesforce LWC", "the Agentic Harness console"),
            ("people", "Approval page", "SSO · H1-H4 · G1-G3"), ("tableau", "Tableau review", "masked dashboards, pre-send report"),
            ("jira", "Jira", "status of every step")]
    y = fy + 0.75
    for i, (ic, t, d) in enumerate(surf):
        logo(cv, ic, sx + 0.12, y + 0.01, 0.2)
        cv.text(sx + 0.4, y, sw - 0.48, 0.2, t, size=8, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(sx + 0.4, y + 0.2, sw - 0.48, 0.3, d, size=6.3, color=SOFT)
        if t == "Tableau review":
            new_tag(sx + sw - 0.46, y + 0.03)
        y += 0.6

    # runtime / registry / context
    mx, mw = 2.6, 4.25

    def block(y, h, title, sub, items, new=None):
        cv.rect(mx, y, mw, h, None, YELLOW, 1.25, R, 0.05)
        cv.text(mx + 0.12, y + 0.06, 1.3, 0.26, title, size=10, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(mx + 0.12, y + 0.32, 1.2, 0.5, sub, size=6, color=DIM)
        iw = (mw - 1.45 - 0.1) / len(items)
        for k, (t, d) in enumerate(items):
            ix = mx + 1.4 + k * iw
            cv.text(ix, y + 0.08, iw - 0.05, 0.2, t, size=7.4, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE)
            cv.text(ix, y + 0.3, iw - 0.05, 0.4, d, size=6, color=SOFT, align=PP_ALIGN.CENTER)
            if new and k in new:
                new_tag(ix + (iw - 0.41) / 2, y + h - 0.22)

    block(fy + 0.14, 0.9, "Runtime", "Runs every story or campaign as a durable workflow.",
          [("Orchestration", "Cloud Workflows"), ("Dispatcher", "one agent task per step"),
           ("State + memory", "Firestore checkpoints")])
    block(fy + 1.12, 0.95, "Registry", "What agents may use.",
          [("Tools", "MCP gateway tools"), ("Skills", "Health, Marketing, QA, campaign"),
           ("Templates", "briefs, designs, journeys")], new=[1, 2])
    cy, ch = fy + 2.15, 1.69
    cv.rect(mx, cy, mw, ch, None, YELLOW, 1.25, R, 0.05)
    cv.text(mx + 0.12, cy + 0.06, 1.3, 0.26, "Context", size=10, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(mx + 0.12, cy + 0.32, 1.2, 0.5, "What agents know, always cited.", size=6, color=DIM)
    citems = [("Knowledge graph", "Spanner Graph, GQL"), ("Cited retrieval", "Graph API over MCP"),
              ("Data context", "masked BigQuery aggregates")]
    iw = (mw - 1.45 - 0.1) / 3
    for k, (t, d) in enumerate(citems):
        ix = mx + 1.4 + k * iw
        cv.text(ix, cy + 0.08, iw - 0.05, 0.2, t, size=7.4, bold=True, color=MAGENTA if k == 2 else WHITE,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(ix, cy + 0.3, iw - 0.05, 0.36, d, size=6, color=SOFT, align=PP_ALIGN.CENTER)
    new_tag(mx + 1.4 + 2 * iw + (iw - 0.41) / 2, cy + 0.62)
    cv.text(mx + 0.12, cy + 0.88, 1.0, 0.16, "SOURCES", size=5.8, bold=True, color=YELLOW, track=150)
    srcs = ["Jira", "Confluence", "Salesforce", "Git", "Copado", "Google Drive", "GCP · BigQuery (masked)"]
    xx, yy = mx + 0.12, cy + 1.08
    for t in srcs:
        w = 0.16 + 0.058 * len(t)
        if xx + w > mx + mw - 0.1:
            xx, yy = mx + 0.12, yy + 0.26
        gcp = t.startswith("GCP")
        cv.rect(xx, yy, w, 0.2, TINT[MAGENTA] if gcp else CARD2, MAGENTA if gcp else EDGE, 0.75, R, 0.4)
        cv.text(xx, yy, w, 0.2, t, size=5.9, bold=gcp, color=MAGENTA if gcp else WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        xx += w + 0.06

    # govern
    gx, gw = 7.0, 2.45
    cv.rect(gx, fy + 0.14, gw, fh - 0.28, None, MAGENTA, 1.5, R, 0.05)
    cv.text(gx + 0.14, fy + 0.2, gw - 0.2, 0.28, "Harness Govern", size=10, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(gx + 0.14, fy + 0.47, gw - 0.2, 0.2, "Safe, controlled and measurable.", size=6, color=DIM)
    gov = [("Change Guardian", False), ("Guardrails in + out · Model Armor", False), ("Identity + access",
                                                                                        False),
           ("Collibra masking + DLP", True), ("Compliance Guardian + policy engine", True),
           ("Human gates H1-H4, CAB, G1-G3", False), ("Evaluation set", False),
           ("Observability + change passport", False)]
    y = fy + 0.78
    for t, nw in gov:
        key = t == "Change Guardian"
        cv.text(gx + 0.14, y, gw - 0.6, 0.34, t, size=8.6 if key else 7.4, bold=True,
                color=CYAN if key else (MAGENTA if nw else WHITE), anchor=MSO_ANCHOR.MIDDLE)
        if nw:
            new_tag(gx + gw - 0.48, y + 0.1)
        if key:
            pill(cv, gx + gw - 0.86, y + 0.09, 0.76, "WE BUILD IT", CYAN, True, h=0.17, size=5)
        y += 0.38

    # differentiators
    dx, dw = 9.75, 3.13
    cv.text(dx, 1.8, dw, 0.26, "Differentiators", size=11, bold=True, color=MAGENTA, anchor=MSO_ANCHOR.MIDDLE)
    diffs = [("Grounded in the org and its data", "Agents read a cited graph of the delivery estate, and masked "
              "aggregates for campaigns.", "Ask before build."),
             ("No pod breaks another", [[("The ", {}), ("Change Guardian", {"bold": True, "color": CYAN}),
                                         (" checks every change against who owns and uses it.", {})]],
              "Governed by design."),
             ("Minimum-data agent access", "Agents get aggregates and masked views, never raw PHI; patient "
              "data leaves only through the gateway.", "Private by design.")]
    y = 2.12
    for i, (t, d, k) in enumerate(diffs):
        cv.rect(dx, y, dw, 1.12, CARD, WHITE, 1.0, R, 0.1)
        badge(cv, dx + 0.12, y + 0.14, i + 1, MAGENTA, 0.3)
        cv.text(dx + 0.52, y + 0.1, dw - 0.6, 0.26, t, size=8.4, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(dx + 0.52, y + 0.38, dw - 0.6, 0.46, d, size=6.5, color=SOFT)
        cv.text(dx + 0.52, y + 0.84, dw - 0.6, 0.2, k, size=6.5, color=MAGENTA)
        y += 1.22

    # core components
    by = 5.9
    cv.rect(0.45, by, 12.43, 0.95, None, WHITE, 1.25, R, 0.4)
    cv.text(0.8, by + 0.08, 2.15, 0.2, "ASCENSION CORE COMPONENTS", size=6, bold=True, color=WHITE, track=60)
    cv.text(0.8, by + 0.36, 2.0, 0.5, [[("Infrastructure", {"size": 11})], [("Runs in Ascension's cloud", {"size":
                                                                                                       6.5,
                                                                                                       "color":
                                                                                                       SOFT})]],
            size=11, color=WHITE)
    comps = [("salesforce", "Salesforce", "Health, Marketing + Data Cloud"),
             ("git", "Delivery", "Copado, Git, Jira, Zephyr"),
             ("googlebigquery", "Data", "BigQuery, Collibra, DLP, Tableau, Spanner"),
             ("openai-white", "Models", "Codex Enterprise; model in config; tools via MCP"),
             ("googlecloud", "Cloud", "Google Cloud: Cloud Run, VPC-SC, CMEK")]
    cw_ = (12.88 - 3.0 - 0.2) / 5
    for i, (ic, t, d) in enumerate(comps):
        x = 3.0 + i * cw_
        logo(cv, ic, x, by + 0.14, 0.22)
        x += 0.3
        cv.text(x, by + 0.12, cw_ - 0.4, 0.26, t, size=9, bold=True, color=MAGENTA if t == "Data" else WHITE,
                anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x, by + 0.4, cw_ - 0.4, 0.45, d, size=6.6, color=SOFT)
    set_notes(s, PLATFORM_NOTES)
    return s


# ================================================================== 4 Health Cloud flow
HC_NOTES = """SLIDE: Health Cloud: a story from Codex to production   (about 3 minutes)

KEY MESSAGE
For Health Cloud the harness delivers a change, not a campaign, so it needs the engineering context, not patient data. The governed GCP data source is shown grey: it is not used for delivery stories.

WALK IT
1 A team member enters NS-1234 in Codex or the Salesforce console: an HFN validation rule for Texas CIN.
2 The Requirements Agent reads Jira and Confluence and the knowledge graph; the product owner approves the brief (H1).
3 The Architect Agent checks what already exists, its impact and its owners; the Change Guardian names affected pods; the architect and pod owners approve (H2).
4 The Developer Agent builds in the DEV sandbox on its own branch.
5 The Review and Test Agents check code, security and sharing, and run the tests; a failed check goes back to the Developer Agent, at most five times, then to a named owner; the tech lead approves (H3).
6 The Release Agent promotes through Copado to SIT and UAT; business signs UAT (H4); CAB approves; the owning team deploys to production. No agent holds a production credential.
Every step is scored by evals and recorded in the change passport.
BOTTOM ROW: the Change Guardian reviews every pod's metadata for the change and flags anything that would block deployment (a dependency, a conflict with another pod's work, a missing component) before release; people approve at H1-H4 and CAB, as today; the change passport records every step, approval and test."""


def health_slide(prs, layout):
    s, cv = new_slide(prs, layout)
    dark_header(cv, "ASCENSION AGENTIC HARNESS · HEALTH CLOUD FLOW",
                "Health Cloud: one story from Codex to production",
                "Example: NS-1234, an HFN validation rule for Texas CIN. Delivery needs the engineering context, not "
                "patient data, so the GCP data source stays out (grey).")
    # row 1: the run frame
    y1 = 1.3
    cv.rect(0.45, y1, 3.05, 0.78, BLACK, CYAN, 1.25, R, 0.1)
    badge(cv, 0.53, y1 + 0.08, 1)
    logo(cv, "openai-white", 0.82, y1 + 0.07, 0.22)
    cv.text(1.1, y1 + 0.05, 2.3, 0.26, "Start in Codex or the LWC", size=8, bold=True, color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(0.57, y1 + 0.36, 2.85, 0.4, "“Deliver NS-1234: HFN validation rule for Texas CIN”", size=6.8,
            color=SOFT, font="Courier New")
    frame = [("Run orchestration", "Cloud Workflows: 9 steps, checkpoints, resumable", CYAN),
             ("Tool gateway", "3 identities: read, develop, release; no production", MAGENTA),
             ("Evals on every step", "Scored before it moves; max 5 retries, then a named owner", YELLOW)]
    x = 3.65
    fw = (12.88 - 3.65 - 2 * 0.15) / 3
    for t, d, c in frame:
        card(cv, x, y1, fw, 0.78, t, d, line=c, fill=TINT[c], bsize=6.5)
        x += fw + 0.15
    cv.arrow(3.5, y1 + 0.39, 3.65, y1 + 0.39, CYAN, 1.5)

    # row 2: the chain
    y2, h2 = 2.4, 1.5
    chain = [("a", "Requirements", "User stories and functional decomposition from the Jira story; open questions "
              "flagged.", 2),
             ("g", "H1", "product owner"),
             ("a", "Architect", "Reuse and impact from the graph; Change Guardian names affected pods.", 3),
             ("g", "H2", "architect + pod owners"),
             ("a", "Developer", "Validation rule, Flow and tests in the DEV sandbox, on its own branch.", 4),
             ("a", "Review + Test", "Code, security and sharing checks; Apex tests and Zephyr cases.", 5),
             ("g", "H3", "tech lead + security"),
             ("a", "Release", "Copado promotion to SIT, then UAT; rollback plan and CAB pack.", 6),
             ("g", "H4 · CAB", "business UAT, then CAB"),
             ("p", "Production", "the owning team deploys; no agent credential")]
    widths = {"a": 1.6, "g": 0.62, "p": 1.0}
    total = sum(widths[c[0]] for c in chain) + (len(chain) - 1) * 0.1
    x = 0.45 + (12.43 - total) / 2
    xs = {}
    for c in chain:
        w = widths[c[0]]
        if c[0] == "a":
            card(cv, x, y2, w, h2 - 0.28, c[1] + " Agent" if c[1] != "Review + Test" else "Review + Test Agents",
                 c[2], line=CYAN, fill=CARD, n=c[3], tsize=7.4, bsize=6.3)
            pill(cv, x + (w - 0.5) / 2, y2 + h2 - 0.22, 0.5, "EVALS", YELLOW, True, h=0.17)
        elif c[0] == "g":
            cv.rect(x, y2, w, h2 - 0.28, YELLOW, None, shape=R, radius=0.2)
            logo(cv, "people-black", x + (w - 0.2) / 2, y2 + 0.1, 0.2)
            cv.text(x + 0.03, y2 + 0.3, w - 0.06, 0.3, c[1], size=8.5, bold=True, color=BLACK,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            cv.text(x + 0.03, y2 + 0.62, w - 0.06, 0.55, c[2], size=6, color=BLACK, align=PP_ALIGN.CENTER)
        else:
            cv.rect(x, y2, w, h2 - 0.28, None, SOFT, 1.0, R, 0.1, dash="dash")
            cv.text(x + 0.06, y2 + 0.1, w - 0.12, 0.26, c[1], size=8, bold=True, color=WHITE,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            cv.text(x + 0.06, y2 + 0.42, w - 0.12, 0.7, c[2], size=6.2, color=SOFT, align=PP_ALIGN.CENTER)
        xs[c[1]] = (x, w)
        x += w + 0.1
    keys = [c[1] for c in chain]
    for k1, k2 in zip(keys, keys[1:]):
        x1 = xs[k1][0] + xs[k1][1]
        cv.arrow(x1, y2 + (h2 - 0.28) / 2, x1 + 0.1, y2 + (h2 - 0.28) / 2, WHITE, 1.0)
    # 1 -> 2: the Codex ask starts the Requirements Agent
    rqx, rqw = xs["Requirements"]
    cv.arrow(rqx + rqw / 2, y1 + 0.78, rqx + rqw / 2, y2, CYAN, 1.75)
    # bounded retry: a failed check in 5 goes back to 4, over the top
    dvx, dvw = xs["Developer"]
    rtx, rtw = xs["Review + Test"]
    ly_ = y2 - 0.14
    path(cv, [(rtx + 0.35, y2), (rtx + 0.35, ly_), (dvx + dvw - 0.35, ly_), (dvx + dvw - 0.35, y2)], MAGENTA, 1.75)
    pill(cv, rtx + 0.4, ly_ - 0.09, 2.6, "5 → 4  FAIL → FIX · MAX 5, THEN OWNER", MAGENTA, True, h=0.18,
         tcolor=WHITE)

    # row 3: knowledge graph + systems
    y3 = 4.05
    kx, kw = 0.45, 7.55
    cv.rect(kx, y3, kw, 1.72, TINT[VIO], VIO, 1.25, R, 0.05)
    cv.text(kx + 0.14, y3 + 0.06, kw - 0.3, 0.24, [[("ENGINEERING CONTEXT KNOWLEDGE GRAPH", {"bold": True,
                                                                                         "track": 150}),
                                                  ("   Spanner Graph · every fact cited", {"color": SOFT})]],
            size=7, color=VIO, anchor=MSO_ANCHOR.MIDDLE)
    srcs = [("jira", "Jira", "stories, history"), ("confluence", "Confluence", "designs, rules"),
            ("salesforce", "Salesforce", "Health Cloud metadata"), ("git", "Git", "Apex, LWC, tests"),
            (None, "Copado", "releases, test runs"), ("googledrive", "Google Drive", "specs")]
    sw = (kw - 0.28 - 6 * 0.08) / 7
    for i, (ic, n, d) in enumerate(srcs):
        x = kx + 0.14 + i * (sw + 0.08)
        cv.rect(x, y3 + 0.38, sw, 0.62, CARD, VIO, 0.9, R, 0.1)
        if ic:
            logo(cv, ic, x + 0.08, y3 + 0.45, 0.2)
        else:
            pill(cv, x + 0.08, y3 + 0.46, 0.2, "Co", SOFT, h=0.18, size=5)
        cv.text(x + 0.33, y3 + 0.42, sw - 0.36, 0.24, n, size=7, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.08, y3 + 0.7, sw - 0.12, 0.28, d, size=5.8, color=SOFT)
    # the GCP source, grey: not needed for Health Cloud delivery
    x = kx + 0.14 + 6 * (sw + 0.08)
    cv.rect(x, y3 + 0.38, sw, 0.62, GREY_BG, GREY_LN, 0.9, R, 0.1, dash="dash")
    logo(cv, "googlecloud-grey", x + 0.08, y3 + 0.45, 0.2)
    cv.text(x + 0.33, y3 + 0.42, sw - 0.36, 0.24, "GCP data", size=7, bold=True, color=GREY_TX,
            anchor=MSO_ANCHOR.MIDDLE)
    cv.text(x + 0.08, y3 + 0.66, sw - 0.12, 0.32, "not needed for Health Cloud", size=5.6, color=GREY_TX)
    qs = ["graph.exists", "graph.find_reuse", "graph.blast_radius", "graph.owners", "graph.context_slice"]
    xx = kx + 0.14
    cv.text(xx, y3 + 1.1, 1.3, 0.2, "AGENTS ASK", size=5.8, bold=True, color=VIO, track=150,
            anchor=MSO_ANCHOR.MIDDLE)
    xx += 0.85
    for q in qs:
        w = 0.16 + 0.06 * len(q)
        cv.rect(xx, y3 + 1.1, w, 0.2, CARD2, EDGE, 0.75, R, 0.4)
        cv.text(xx, y3 + 1.1, w, 0.2, q, size=5.8, color=WHITE, font="Courier New", align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        xx += w + 0.06
    cv.text(kx + 0.14, y3 + 1.38, kw - 0.3, 0.26, "A bounded, cited context slice for each step, never a raw org "
            "dump. Agents never write queries.", size=6.4, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    # graph feeds requirements + architect
    for k in ("Requirements", "Architect"):
        x0, w0 = xs[k]
        cv.arrow(x0 + 0.3, y3, x0 + 0.3, y2 + h2 - 0.28, VIO, 1.5)

    # systems written
    sx, sw2 = 8.15, 4.73
    cv.rect(sx, y3, sw2, 1.72, CARD, EDGE, 0.75, R, 0.05)
    cv.text(sx + 0.14, y3 + 0.06, sw2 - 0.3, 0.24, [[("WHERE OUTPUT LANDS", {"bold": True, "track": 150}),
                                                    ("   through the tool gateway", {"color": SOFT})]],
            size=7, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    outs = [("jira", "Jira", "brief, status to Done"), ("confluence", "Confluence", "design draft, as-built"),
            ("git", "Git", "branch, pull request"), ("salesforce", "Salesforce DEV", "deploy, Apex tests"),
            (None, "Copado", "SIT, then UAT"), (None, "Zephyr Scale", "cases, results")]
    for i, (ic, n, d) in enumerate(outs):
        x = sx + 0.14 + (i % 2) * 2.25
        y = y3 + 0.36 + (i // 2) * 0.35
        if ic:
            logo(cv, ic, x, y + 0.05, 0.2)
        else:
            pill(cv, x, y + 0.06, 0.2, "Co" if n == "Copado" else "ZS", SOFT, h=0.18, size=5)
        cv.text(x + 0.27, y, 1.95, 0.32, [[(n, {"bold": True, "color": WHITE}), ("  " + d, {"color": SOFT,
                                                                                              "size": 6})]],
                size=7, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(sx + 0.14, y3 + 1.42, sw2 - 0.3, 0.22, "Production: the owning team's pipeline after CAB. No agent path.",
            size=6.4, bold=True, color=MAGENTA, anchor=MSO_ANCHOR.MIDDLE)

    # bottom line
    by = 5.95
    cells = [("Change Guardian", "reviews every pod's metadata; flags deployment blockers before release"),
             ("People approve", "H1, H2, H3, H4 and CAB, as today"),
             ("Change passport", "every step, approval and test recorded")]
    cw_ = (12.43 - 2 * 0.15) / 3
    for i, (t, d) in enumerate(cells):
        x = 0.45 + i * (cw_ + 0.15)
        cv.rect(x, by, cw_, 0.85, CARD, None)
        cv.rect(x, by, 0.06, 0.85, (CYAN, YELLOW, MAGENTA)[i], None)
        cv.text(x + 0.22, by + 0.08, cw_ - 0.3, 0.4, t, size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cv.text(x + 0.22, by + 0.48, cw_ - 0.3, 0.3, d, size=7, color=SOFT, anchor=MSO_ANCHOR.MIDDLE)
    set_notes(s, HC_NOTES)
    return s


# ================================================================== build
SECTION_NOTES = ("Section: the HIPAA-aligned agentic healthcare campaign platform. The same harness, now for "
                 "Marketing Cloud campaigns: the data is governed in Google Cloud first, agents plan and build on "
                 "masked data, and only approved patients are released. It is a proposed target state, piloted "
                 "first on the CHF post-discharge nudge.")


def section_slide(prs):
    """A section title on the IG colour title layout, styled like slide 1."""
    import copy
    first = prs.slides[0]
    s = prs.slides.add_slide(first.slide_layout)
    tree = s.shapes._spTree
    for sp in list(s.shapes):
        tree.remove(sp._element)
    for sp in first.shapes:
        tree.append(copy.deepcopy(sp._element))
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        paras = sh.text_frame.paragraphs
        if t.startswith("Ascension"):
            from pptx.util import Pt
            runs = paras[0].runs
            runs[0].text = "HIPAA-Aligned Agentic Healthcare Campaign Platform"
            runs[0].font.size = Pt(36)
            for r in runs[1:]:
                r.text = ""
        elif paras and paras[0].runs and (t.startswith("One harness") or t.startswith("An Agentic")):
            paras[0].runs[0].text = ("Governed data in Google Cloud, agents on masked data, approved patients "
                                     "released to Marketing Cloud. Proposed target state.")
            for r in paras[0].runs[1:]:
                r.text = ""
        elif t in ("AGENTS FOR DELIVERY AND DATA", "AI-DRIVEN SALESFORCE DELIVERY"):
            paras[0].runs[0].text = "MARKETING CLOUD CAMPAIGNS"
    set_notes(s, SECTION_NOTES)
    return s


def white_title(prs, layout, eyebrow, title, sub, notes):
    """A title slide on white for the white deck: IG gradient rule, cyan eyebrow, dark title."""
    import sierra_voice
    s = prs.slides.add_slide(layout)
    set_bg(s, "FFFFFF")
    cv = Canvas(s)
    sierra_voice.gbar(cv, 0.9, 2.15, 1.8, 0.08, 0)
    cv.text(0.9, 2.4, 11.5, 0.36, eyebrow, size=12, bold=True, color="0089A0", track=300, anchor=MSO_ANCHOR.MIDDLE)
    cv.text(0.9, 2.85, 11.5, 1.5, title, size=40, bold=True, color="1A1A1A", anchor=MSO_ANCHOR.TOP, line=0.95)
    cv.text(0.9, 4.45, 9.5, 0.8, sub, size=15, color="555555", anchor=MSO_ANCHOR.TOP)
    set_notes(s, notes)
    return s


def main(out=OUT, theme="dark"):
    import build_mad_ig_deck as ig
    import combined_flow
    import campaign_platform_deck as cpd
    import muted_theme as mt
    white = theme == "white"
    if white:
        global LOGOS
        combined_flow.light_logos()             # dark icon variants for white slides
        import build_data_harness_deck as bd    # this file may run as __main__: point both copies at them
        LOGOS = bd.LOGOS
        if out == OUT:
            out = OUT_WHITE
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst
    for el in list(lst)[1 if not white else 0:]:   # keep only the template title slide (dark deck)
        prs.part.drop_rel(el.rId)
        lst.remove(el)
    layout = next(l_ for l_ in prs.slide_layouts if l_.name == "DEFAULT")
    if white:
        white_title(prs, layout, "AGENTS FOR DELIVERY AND DATA", "Ascension Agentic Harness",
                    "One harness, two jobs: Health Cloud delivery from Codex, and Marketing Cloud campaigns on "
                    "governed, masked Google Cloud data.", TITLE_NOTES)
    else:
        title_slide(prs)
    team = team_slide(prs, layout)
    platform = platform_slide(prs, layout)
    health = health_slide(prs, layout)
    import hc_timeline
    hc_timeline.build(prs, layout)              # Health Cloud: ~2 weeks manual vs ~2 days on the harness
    if white:
        white_title(prs, layout, "MARKETING CLOUD CAMPAIGNS", "HIPAA-Aligned Agentic Healthcare Campaign Platform",
                    "Governed data in Google Cloud, agents on masked data, approved patients released to Marketing "
                    "Cloud. Proposed target state.", SECTION_NOTES)
    else:
        section_slide(prs)
    gov = cpd.phi_governance(prs, layout)       # dark slides are built with the white icon variants
    frame = ig.Canvas.frame                     # thinner gradient frame on the end-to-end slide
    ig.Canvas.frame = lambda cv, x, y, w, h, inner=WHITE, t=0.045: frame(cv, x, y, w, h, inner, min(t, 0.02))
    flow = combined_flow.build(prs, layout)
    ig.Canvas.frame = frame
    set_notes(flow, cpd.FLOW_NOTES)
    order = cpd.build_order(prs, layout)
    import campaign_timeline
    timeline = campaign_timeline.build(prs, layout)    # campaign: ~90 days manual vs ~10-12 on the data harness
    import sierra_voice
    sierra = sierra_voice.build(prs, layout)    # Sierra voice agent
    import platform_fit
    platform_fit.build(prs, layout, colorful=not white)   # platform fit: build, scale, accuracy, data, control
    import lean_team
    lean = lean_team.build(prs, layout)         # last slide: lean squads and savings
    w1 = (11.95 - 4 * 0.16) / 5
    xs1 = [0.6 + k * (w1 + 0.16) for k in range(5)]
    gh = (combined_flow.CH1 - 0.06) / 2
    metal_cards = [("bronze", xs1[1], 1.62, w1, combined_flow.CH1), ("silver", xs1[3], 1.62, w1, combined_flow.CH1),
                   ("gold", xs1[4], 1.62, w1, gh), ("gold", xs1[4], 1.62 + gh + 0.06, w1, gh)]
    if white:
        # muted styles: greys, one accent, thin borders, the IG gradient only on hooks
        mt.accent_dark(team, (CYAN, CYAN))
        mt.muted(platform, "dark", accent="yellow",
                 hooks=[("frame", lambda sh: mt.near(sh.left, 9.75) and mt.near(sh.width, 3.13))])
        mt.muted(health, "dark", yellow_if=lambda sh: sh.height > mt.EMU,
                 hooks=[("frame", lambda sh: mt.near(sh.top, 5.95) and mt.near(sh.height, 0.85)
                         and sh.width > 3 * mt.EMU),
                        ("bar", lambda sh: mt.near(sh.top, 5.95) and mt.near(sh.width, 0.06))])
        mt.muted(flow, "dark", hooks=[("bar", lambda sh: mt.near(sh.top, 6.36) and mt.near(sh.width, 0.05))])
        mt.metals(flow, metal_cards, theme="dark")
        mt.muted(order, "dark", accent="yellow", src={YELLOW},      # phase 5, first real patients, stands out
                 hooks=[("bar", lambda sh: mt.near(sh.top, 5.78) and mt.near(sh.width, 0.05))])
        mt.muted(lean, "dark")                  # the timelines keep their colours: they carry meaning
        mt.muted(gov, "dark", src={CYAN, YELLOW}, src_accent={YELLOW: "yellow"},
                 neutral_if=lambda sh: (mt.near(sh.top, 4.58) and mt.near(sh.height, 0.3))
                 or (4.58 * mt.EMU < sh.top < 6.26 * mt.EMU and sh.width < 0.1 * mt.EMU),
                 hooks=[("frame", lambda sh: mt.near(sh.top, 4.58) and mt.near(sh.height, 1.68))])
    else:
        # colourful black: the full IG colours, thin borders, metals on the data layers
        for sl in (team, platform, health, flow, order, gov):
            mt.vivid(sl)
        mt.metals(flow, metal_cards, theme="dark")
    # built as title, team, platform, health, health timeline, section, governance, flow, order, timeline, sierra,
    # platform fit, lean squads; shown with the timeline after the flow and governance before sierra
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for k in (0, 1, 2, 3, 4, 5, 7, 9, 8, 6, 10, 11, 12):
        lst.append(ids[k])
    if white:                                   # same muted design, recoloured for a white background
        for sl in prs.slides:
            set_bg(sl, "FFFFFF")
            mt.to_light(sl)
        sierra_voice.highlight(sierra)          # the Sierra slide is the key differentiator
    else:
        sierra_voice.highlight(sierra, dark=True)
    for sl in prs.slides:                       # readability: nothing below 7pt
        mt.min_font(sl, 7.0)
    for line in fix_footers(prs):
        print(line)
    prs.save(out)
    print("wrote", os.path.abspath(out), "-", len(prs.slides), "slides")


if __name__ == "__main__":
    main(*sys.argv[1:])
