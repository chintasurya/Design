"""Piedmont Healthcare: one solution slide, styled to sit after the client's
"Current State & The Ask" slide (dark ground, amber glow, IG fonts and footer).

    python3 build_solution_slide.py   ->  ../Piedmont Healthcare - Solution.pptx
"""
import copy
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Piedmont Healthcare - Solution.pptx")
BG = os.path.join(HERE, "assets", "bg-glow-amber.jpg")
MARK = os.path.join(HERE, "assets", "ig-mark-white.png")

# IG palette, as in deck/build_mad_ig_deck.py
CYAN, YELLOW, GREEN, ORANGE = "00D6F2", "FFD700", "56E13B", "FA8B35"
WHITE, BLACK, NAVY = "FFFFFF", "000000", "0B1A2B"
MUTED, GREY = "C9CED6", "8C939C"

TITLE_FONT = "Trenda IG Display Bold"
BODY = "Trenda IG Text"
SEMI = "Trenda IG Text Semibold"


# ---------------------------------------------------------------- helpers
def _alpha(fill_el, pct_opaque):
    clr = fill_el.find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(pct_opaque * 1000)))


def box(slide, name, x, y, w, h, fill=None, opaque=100, line=None, line_w=1.0,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, dash=False):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if opaque < 100:
            _alpha(sh._element.spPr.find(qn("a:solidFill")), opaque)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(line_w)
        if dash:
            ln = sh._element.spPr.find(qn("a:ln"))
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    sh.text_frame.text = ""
    return sh


def text(slide, name, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.0):
    """paras: list of dicts {runs:[(text, {size,color,font,italic})], bullet, align, after}."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(margin))
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        para.space_after = Pt(p.get("after", 0))
        para.line_spacing = 1.0
        if p.get("bullet"):
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.13)))
            pPr.set("indent", str(-Inches(0.13)))
            bc = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:srgbClr")).set("val", p.get("bullet_color", WHITE))
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        for t, o in p["runs"]:
            r = para.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", 10))
            f.color.rgb = RGBColor.from_string(o.get("color", WHITE))
            f.name = o.get("font", BODY)
            f.italic = o.get("italic", False)
    return tb


def P(t, size=10, color=WHITE, font=BODY, **kw):
    """One-run paragraph."""
    italic = kw.pop("italic", False)
    return dict(runs=[(t, dict(size=size, color=color, font=font, italic=italic))], **kw)


def bullets(items, size=9.5, color=WHITE, bullet_color=WHITE, after=3):
    return [P(t, size, color, bullet=True, bullet_color=bullet_color, after=after) for t in items]


def arrow(slide, name, x, y, w=0.3, h=0.24, color="E8D9A8"):
    box(slide, name, x, y, w, h, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)


def badge(slide, name, x, y, n, color, label):
    c = box(slide, name + " Badge", x, y, 0.36, 0.36, fill=color, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size, r.font.name = Pt(14), SEMI
    r.font.color.rgb = RGBColor.from_string(BLACK if color != GREY else WHITE)
    text(slide, name + " Heading", x + 0.48, y, 4.2, 0.36, [P(label, 15, color, SEMI)],
         anchor=MSO_ANCHOR.MIDDLE)


def set_background(slide, image_path):
    _, rid = slide.part.get_or_add_image_part(image_path)
    cSld = slide._element.find(qn("p:cSld"))
    bg = etree.SubElement(cSld, qn("p:bg"))
    cSld.remove(bg)
    cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, qn("p:bgPr"))
    fill = etree.SubElement(bgPr, qn("a:blipFill"), dpi="0", rotWithShape="1")
    etree.SubElement(fill, qn("a:blip")).set(qn("r:embed"), rid)
    etree.SubElement(fill, qn("a:srcRect"))
    etree.SubElement(etree.SubElement(fill, qn("a:stretch")), qn("a:fillRect"))
    etree.SubElement(bgPr, qn("a:effectLst"))


def slide_number(slide, x, y, w, h):
    tb = text(slide, "IG Footer Page Number", x, y, w, h, [P("", 8)])
    p = tb.text_frame.paragraphs[0]._p
    tb.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    for r in p.findall(qn("a:r")):
        p.remove(r)
    fld = etree.SubElement(p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}",
                           type="slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"), lang="en-US", sz="800")
    etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr")).set("val", WHITE)
    etree.SubElement(rpr, qn("a:latin")).set("typeface", BODY)
    etree.SubElement(fld, qn("a:t")).text = "3"
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        p.remove(end)
        p.append(end)


# ---------------------------------------------------------------- slide
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = [l for l in prs.slide_layouts if l.name == "Blank"][0]
    s = prs.slides.add_slide(blank)
    set_background(s, BG)

    # title
    text(s, "Title", 0.55, 0.22, 12.2, 0.55,
         [P("Piedmont Healthcare  |  Proposed Solution & Implementation", 26, WHITE, TITLE_FONT)],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, "Standfirst", 0.55, 0.78, 12.2, 0.32,
         [P("Fix the data, build a governed structure in Marketing Cloud, then support the team. "
            "Insight Global specialists work inside the Foundation Team.", 12, MUTED)])

    # ------------------------------------------------ 1. target architecture
    badge(s, "S1", 0.55, 1.18, 1, CYAN, "Target Data Architecture")

    # Epic and Health Cloud
    epic = box(s, "Epic", 0.55, 1.72, 1.75, 0.95, fill=NAVY, line="CBD5E1", line_w=1)
    text(s, "Epic Text", 0.55, 1.72, 1.75, 0.95,
         [P("Epic (EHR)", 13, WHITE, SEMI, align=PP_ALIGN.CENTER, after=2),
          P("Clinical source system", 9.5, MUTED, align=PP_ALIGN.CENTER)],
         anchor=MSO_ANCHOR.MIDDLE, margin=0.06)
    arrow(s, "Arrow Epic-HC", 2.37, 2.075)
    box(s, "Health Cloud", 2.74, 1.72, 1.93, 0.95, fill=CYAN)
    text(s, "Health Cloud Text", 2.74, 1.72, 1.93, 0.95,
         [P("Salesforce Health Cloud", 12.5, BLACK, SEMI, align=PP_ALIGN.CENTER, after=2),
          P("System of record", 9.5, "0B2A33", align=PP_ALIGN.CENTER)],
         anchor=MSO_ANCHOR.MIDDLE, margin=0.06)
    arrow(s, "Arrow HC-SFMC", 4.73, 2.075)

    # fix-at-source card under Epic and Health Cloud
    box(s, "Source Fixes", 0.55, 2.85, 4.12, 1.2, fill=BLACK, opaque=50, line=YELLOW, line_w=1)
    text(s, "Source Fixes Text", 0.55, 2.85, 4.12, 1.2,
         [P("Fix at the source", 10.5, YELLOW, SEMI, after=4)]
         + bullets(["Epic batch upserts on MRN (External ID): no new duplicates",
                    "Health Cloud duplicate and matching rules switched on",
                    "Subscriber Key = Contact ID; sync only the PHI fields needed"],
                   bullet_color=YELLOW, after=2),
         margin=0.12)

    # SFMC parent BU with its four layers
    box(s, "SFMC Foundation", 5.1, 1.55, 5.02, 2.5, fill=BLACK, opaque=50, line=CYAN, line_w=1.25)
    text(s, "SFMC Foundation Heading", 5.24, 1.6, 4.8, 0.32,
         [dict(runs=[("Salesforce Marketing Cloud  ", dict(size=11.5, color=CYAN, font=SEMI)),
                     ("parent BU data foundation", dict(size=10, color=MUTED))])],
         anchor=MSO_ANCHOR.MIDDLE)
    layers = [
        ("RAW", "Synced DEs from Health Cloud, never edited", "D0D4DA", "2E3338", GREY),
        ("CLEAN", "SQL automations standardize, validate, dedupe", CYAN, "07343C", CYAN),
        ("GOLDEN", "One patient master, one provider master, consent", YELLOW, "3B3100", YELLOW),
        ("SHARED", "Filtered, permissioned DEs for each BU", GREEN, "0E3510", GREEN),
    ]
    lx, lw, gap = 5.24, 1.06, 0.16
    for i, (label, desc, tcol, fcol, lcol) in enumerate(layers):
        x = lx + i * (lw + gap)
        box(s, "Layer " + label, x, 1.98, lw, 1.3, fill=fcol, line=lcol, line_w=1)
        text(s, "Layer %s Text" % label, x, 1.98, lw, 1.3,
             [P("%d  %s" % (i + 1, label), 11, tcol, SEMI, after=4),
              P(desc, 9, WHITE)], margin=0.07)
        if i < len(layers) - 1:
            box(s, "Layer Arrow %d" % (i + 1), x + lw + 0.035, 2.55, 0.09, 0.16,
                fill="E8D9A8", shape=MSO_SHAPE.ISOSCELES_TRIANGLE).rotation = 90
    box(s, "Quality Gate", 5.24, 3.4, 4.72, 0.52, fill=ORANGE, opaque=14, line=ORANGE,
        line_w=0.75, dash=True)
    text(s, "Quality Gate Text", 5.24, 3.4, 4.72, 0.52,
         [dict(runs=[("Quality gate  ", dict(size=10, color=ORANGE, font=SEMI)),
                     ("Bad records go to an exceptions DE with a reason code and are never sent. "
                      "A data-quality dashboard tracks them.", dict(size=9, color=WHITE))])],
         anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
    arrow(s, "Arrow SFMC-BUs", 10.18, 2.53)

    # child BUs
    text(s, "Child BU Heading", 10.55, 1.6, 2.25, 0.32,
         [P("Child BUs by department", 11, WHITE, SEMI)], anchor=MSO_ANCHOR.MIDDLE)
    bus = [("Patients · Stream A", "Active today", CYAN, BLACK, "0B2A33"),
           ("Patients · Stream B", "Active today", CYAN, BLACK, "0B2A33"),
           ("Providers", "Future state, ready on day one", GREY, WHITE, "EEF0F2")]
    for i, (name, sub, fcol, tcol, scol) in enumerate(bus):
        y = 1.98 + i * 0.62
        box(s, "BU " + name, 10.55, y, 2.25, 0.52, fill=fcol)
        text(s, "BU %s Text" % name, 10.55, y, 2.25, 0.52,
             [P(name, 11, tcol, SEMI, align=PP_ALIGN.CENTER),
              P(sub, 8.5, scol, italic=True, align=PP_ALIGN.CENTER)],
             anchor=MSO_ANCHOR.MIDDLE, margin=0.04)
    text(s, "Child BU Note", 10.55, 3.82, 2.25, 0.25,
         [P("Each team sees only its own data", 9, MUTED, italic=True, align=PP_ALIGN.CENTER)],
         anchor=MSO_ANCHOR.MIDDLE)

    # ------------------------------------------------ 2. roadmap
    badge(s, "S2", 0.55, 4.28, 2, YELLOW, "Implementation Roadmap")
    text(s, "Roadmap Duration", 5.0, 4.28, 4.6, 0.36,
         [P("About 9 months to a governed foundation, then ongoing support", 9.5, MUTED,
            italic=True, align=PP_ALIGN.RIGHT)], anchor=MSO_ANCHOR.MIDDLE)
    phases = [
        ("PHASE 0  ·  MONTHS 1–2", "Assess", CYAN,
         ["Profile the sync and every DE", "Find root causes of bad data",
          "Agree the BU access model"],
         "Data-quality baseline and target design"),
        ("PHASE 1  ·  MONTHS 3–5", "Stabilize", YELLOW,
         ["Subscriber Key to Contact ID", "Dedupe, standardize records",
          "Fix the Epic upsert at source"],
         "One clean patient master"),
        ("PHASE 2  ·  MONTHS 6–9", "Structure", GREEN,
         ["Build the 4-layer data model", "Set up BUs and shared DEs",
          "Align consent, add providers"],
         "Governed foundation, data dictionary, runbooks"),
        ("PHASE 3  ·  MONTH 10+", "Sustain", ORANGE,
         ["Monitor quality, exceptions", "Build new audiences",
          "Train, hand over to the team"],
         "Monthly scorecard; provider outreach live"),
    ]
    for i, (tag, name, col, items, out) in enumerate(phases):
        x, y, w, h = 0.55 + i * 2.3, 4.72, 2.15, 2.1
        box(s, "Phase " + name, x, y, w, h, fill=BLACK, opaque=50, line=col, line_w=1)
        text(s, "Phase %s Text" % name, x, y, w, 1.55,
             [P(tag, 8.5, col, SEMI, after=1), P(name, 15, WHITE, SEMI, after=4)]
             + bullets(items, size=9.5, bullet_color=col, after=2), margin=0.12)
        text(s, "Phase %s Output" % name, x, y + 1.52, w, 0.58,
             [dict(runs=[("Delivers  ", dict(size=9, color=col, font=SEMI)),
                         (out, dict(size=9, color=MUTED))])],
             anchor=MSO_ANCHOR.BOTTOM, margin=0.12)

    # ------------------------------------------------ 3. team and success
    badge(s, "S3", 9.85, 4.28, 3, GREEN, "Team & Success Measures")
    box(s, "Team Card", 9.85, 4.72, 2.95, 2.1, fill=BLACK, opaque=50, line=GREEN, line_w=1)
    text(s, "Team Card Text", 9.85, 4.72, 2.95, 2.1,
         [P("IG pod inside the Foundation Team", 10, GREEN, SEMI, after=3)]
         + bullets(["1 SFMC data architect / lead", "1–2 SFMC developers (SQL, AMPscript)",
                    "Health Cloud admin, part-time"], bullet_color=GREEN, after=1)
         + [P("How we measure success", 10, GREEN, SEMI, after=3)]
         + bullets(["Duplicate rate down", "Sendable, consented rate up",
                    "Sync errors down", "Audiences built in days, not weeks"],
                   bullet_color=GREEN, after=1),
         margin=0.12)
    # spacing before the second heading
    tf = s.shapes[-1].text_frame
    tf.paragraphs[4].space_before = Pt(6)

    # footer
    s.shapes.add_picture(MARK, Inches(0.55), Inches(7.03), height=Inches(0.167)).name = "IG Footer Mark"
    slide_number(s, 12.28, 7.0, 0.52, 0.22)

    s.notes_slide.notes_text_frame.text = (
        "Talk track: we fix bad data where it starts (the Epic upsert and Health Cloud duplicate "
        "rules), use one Subscriber Key, and layer SFMC data from raw to clean to golden to shared. "
        "Business units give each department only its own data, which is the separation they asked "
        "for. Four phases: assess, stabilize, structure, sustain. The IG pod works inside the "
        "Foundation Team, so knowledge stays with Piedmont. The plan allows about nine months to a "
        "governed foundation: time for Epic team dependencies, privacy and security review of PHI "
        "changes, and the Foundation Team's limited capacity. Confirm durations in discovery.")

    prs.save(OUT)
    print("wrote", os.path.relpath(OUT))


if __name__ == "__main__":
    build()
